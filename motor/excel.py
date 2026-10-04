"""Renderizador Excel (.xlsx) do POP de Elite — matrizes de qualidade e auditoria.

Portado do app.py da 1.0, sem caminho /tmp fixo (que não existe no Windows),
com nome de arquivo ASCII e linhas ajustadas ao número de colunas.
"""

import io
import os
import re

from .base import carregar_json, nome_de_saida, paleta, pasta_saida, texto_de


def _celulas(linha, n):
    """Valores de uma linha (lista, dict ou texto) ajustados para n colunas."""
    if isinstance(linha, dict):
        valores = list(linha.values())
    elif isinstance(linha, (list, tuple)):
        valores = list(linha)
    else:
        valores = [linha]
    valores = [texto_de(v) for v in valores]
    if n:
        valores = (valores + [""] * n)[:n]
    return valores


def gerar_excel(dados, pasta=None, cor=None, logo=None, overrides=None):
    """Gera planilha multi-aba a partir do JSON. Devolve o caminho do arquivo.
    Com logo (bytes), insere a imagem em A1 de cada aba."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage

    data = carregar_json(dados)
    meta = data.get("metadata_excel", data.get("metadata", {})) or {}
    planilhas = data.get("planilhas", []) or []
    if not planilhas:
        raise ValueError("JSON da planilha sem 'planilhas'.")

    pal = paleta(cor, overrides)
    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    usados = set()

    for pidx, planilha in enumerate(planilhas):
        if not isinstance(planilha, dict):
            continue
        nome_aba = re.sub(r'[\\/*?:\[\]]', '_', texto_de(planilha.get("nome_aba") or f"Planilha_{pidx+1}"))[:31]
        base_nome, k = nome_aba, 2
        while nome_aba.lower() in usados:  # GUARDA: o Excel não aceita abas repetidas
            sufixo = f"_{k}"
            nome_aba = base_nome[:31 - len(sufixo)] + sufixo
            k += 1
        usados.add(nome_aba.lower())
        ws = wb.create_sheet(title=nome_aba)

        colunas = [texto_de(c) for c in (planilha.get("colunas") or [])]
        linhas = planilha.get("linhas") or []
        titulo = texto_de(planilha.get("titulo") or planilha.get("nome_aba") or "")

        has_logo = False
        if logo:
            try:
                from PIL import Image as PILImage
                pil_img = PILImage.open(io.BytesIO(logo))
                target_h = 60
                target_w = max(1, int(target_h * pil_img.width / pil_img.height))
                pil_img = pil_img.resize((target_w, target_h), PILImage.LANCZOS)
                buf = io.BytesIO()
                pil_img.save(buf, "PNG")
                buf.seek(0)
                xl_img = XLImage(buf)
                xl_img.width = target_w
                xl_img.height = target_h
                ws.add_image(xl_img, "A1")
                ws.column_dimensions["A"].width = max(15, target_w / 7)
                ws.row_dimensions[1].height = 32
                ws.row_dimensions[2].height = 32
                has_logo = True
            except Exception:
                has_logo = False

        title_row = 3 if has_logo else 1
        header_row = title_row + 1 if titulo else title_row

        if titulo:
            cell = ws.cell(row=title_row, column=1, value=titulo)
            cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=pal["primary"])
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if len(colunas) > 1:
                ws.merge_cells(start_row=title_row, start_column=1,
                               end_row=title_row, end_column=len(colunas))
            ws.row_dimensions[title_row].height = 30

        for i, col_name in enumerate(colunas, 1):
            cell = ws.cell(row=header_row, column=i, value=col_name)
            cell.font = Font(name="Calibri", size=11, bold=True, color=pal["primary_dark"])
            cell.fill = PatternFill("solid", fgColor=pal["tertiary"])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        if colunas:
            ws.row_dimensions[header_row].height = 35

        for ri, row_data in enumerate(linhas):
            excel_row = header_row + 1 + ri
            zebra_fill = pal["quaternary"] if ri % 2 == 0 else "FFFFFF"
            for ci, val in enumerate(_celulas(row_data, len(colunas)), 1):
                cell = ws.cell(row=excel_row, column=ci, value=val)
                cell.font = Font(name="Calibri", size=10, color="1A1A1A")
                cell.fill = PatternFill("solid", fgColor=zebra_fill)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = border

        for i in range(1, len(colunas) + 1):
            max_len = len(colunas[i-1])
            for row_data in linhas:
                vals = _celulas(row_data, len(colunas))
                max_len = max(max_len, len(vals[i-1]))
            ws.column_dimensions[get_column_letter(i)].width = min(45, max(14, max_len + 2))

        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    if not wb.sheetnames:
        raise ValueError("Nenhuma aba válida no JSON da planilha.")
    nome = nome_de_saida(meta.get("codigo"), meta.get("arquivo") or meta.get("titulo"), "xlsx", "Planilha")
    fp = os.path.join(pasta_saida(pasta), nome)
    wb.save(fp)
    return fp


def gerar_excel_xlsx(json_str, primary_color=None, logo_bytes=None, palette_overrides=None):
    """Compatibilidade com a assinatura da 1.0 (app.py do Space)."""
    return gerar_excel(json_str, cor=primary_color, logo=logo_bytes, overrides=palette_overrides)
