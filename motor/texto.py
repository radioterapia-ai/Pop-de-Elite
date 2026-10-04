"""Layout "texto institucional" dos documentos de apoio.

TCLE, termo de ciência, checklist, formulário, informativo, nota técnica e código
de conduta saem com o mesmo cabeçalho controlado dos documentos normativos e um
corpo livre em blocos: título, parágrafo, lista, tabela, caixas de marcar e campos
para preencher. O rodapé segue o público: documento interno leva o bloco
Elaborado / Revisado / Aprovado / Vigência; documento para paciente leva só uma
linha com código, versão e vigência, e as assinaturas vão no fim do texto.
"""

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Pt, RGBColor

from .base import carregar_json, nome_arquivo_seguro, normalizar_metadata, paleta, pasta_saida, texto_de
from .i18n import rotulos
from .word import (FONTE, LARGURA_DXA, MARGEM_DIR, MARGEM_ESQ, MARGEM_INF, MARGEM_SUP, TAM_CORPO,
                   TAM_TINY, _ajustar_grades, _campo, add_table, build_footer, build_header, fmt, spacing)

LARGURA_UTIL = Cm(21 - 2.5 - 2.0)
TAM_TITULO = Pt(12)
CAIXA = "☐"


def _texto_rico(p, texto, size=TAM_CORPO, color="1A1A1A", bold=False):
    """Escreve o texto convertendo **negrito** do Markdown em negrito de verdade."""
    for i, parte in enumerate(re.split(r"\*\*(.+?)\*\*", texto_de(texto))):
        if parte:
            fmt(p, parte, bold=bold or i % 2 == 1, size=size, color=color)


def _titulo(doc, texto, pal):
    p = doc.add_paragraph(); spacing(p, 12, 4)
    p.paragraph_format.keep_with_next = True
    _texto_rico(p, texto, size=TAM_TITULO, color=pal["primary_dark"], bold=True)


def _paragrafo(doc, texto):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p, 2, 6, 1.15)
    _texto_rico(p, texto)


def _lista(doc, itens, pal):
    for item in itens or []:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        spacing(p, 1, 3, 1.15)
        p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.4)
        fmt(p, "●  ", size=Pt(9), color=pal["primary"], bold=True)
        _texto_rico(p, item)


def _caixas(doc, itens, pal):
    """Itens de marcar, um por linha; o crítico sai em negrito e vermelho, com ⚠."""
    for item in itens or []:
        critico = isinstance(item, dict) and bool(item.get("critico"))
        p = doc.add_paragraph(); spacing(p, 2, 4, 1.15)
        p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
        fmt(p, f"{CAIXA}  ", size=Pt(12), color=pal["primary_dark"])
        if critico:
            fmt(p, "⚠ ", bold=True, size=TAM_CORPO, color=pal["risk_red"])
        _texto_rico(p, item, color=pal["risk_red"] if critico else "1A1A1A", bold=critico)


def _linha_para_preencher(p):
    """Tabulação até a margem direita, com a linha de preenchimento como guia."""
    p.paragraph_format.tab_stops.add_tab_stop(LARGURA_UTIL, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
    p.add_run("\t")


def _campos(doc, itens, pal):
    """Campos em branco: rótulo e linha(s) para preencher à mão. O modelo nunca traz o dado."""
    for item in itens or []:
        rotulo = texto_de(item.get("rotulo") if isinstance(item, dict) else item).strip().rstrip(":")
        try:
            linhas = int(item.get("linhas", 1)) if isinstance(item, dict) else 1
        except (TypeError, ValueError):
            linhas = 1
        linhas = min(10, max(1, linhas))
        p = doc.add_paragraph(); spacing(p, 6, 2, 1.15)
        fmt(p, f"{rotulo}: ", bold=True, size=TAM_CORPO, color=pal["primary_dark"])
        if linhas == 1:
            _linha_para_preencher(p)
            continue
        p.paragraph_format.keep_with_next = True
        for _ in range(linhas):
            q = doc.add_paragraph(); spacing(q, 8, 0, 1.0)
            _linha_para_preencher(q)


def _tabela(doc, bloco, pal):
    colunas = [texto_de(c) for c in bloco.get("colunas") or []]
    if not colunas:
        return
    n = len(colunas)
    linhas = [[texto_de(v) for v in (list(linha) + [""] * n)[:n]]
              for linha in bloco.get("linhas") or [] if isinstance(linha, (list, tuple))]
    add_table(doc, colunas, linhas, [LARGURA_DXA // n] * n, pal)


def _bloco(doc, bloco, pal):
    if not isinstance(bloco, dict):
        _paragrafo(doc, bloco)
        return
    tipo = str(bloco.get("tipo", "")).strip().lower()
    if tipo == "titulo":
        _titulo(doc, bloco.get("texto"), pal)
    elif tipo == "lista":
        _lista(doc, bloco.get("itens"), pal)
    elif tipo == "tabela":
        _tabela(doc, bloco, pal)
    elif tipo == "caixas":
        _caixas(doc, bloco.get("itens"), pal)
    elif tipo == "campos":
        _campos(doc, bloco.get("itens"), pal)
    else:
        _paragrafo(doc, bloco.get("texto", bloco))


def _assinaturas(doc, assinaturas, pal, R):
    """Linha de assinatura e de data por signatário; o bloco não se parte entre páginas."""
    for a in assinaturas or []:
        papel = texto_de(a.get("papel") if isinstance(a, dict) else a)
        detalhe = texto_de(a.get("detalhe", "")) if isinstance(a, dict) else ""
        p = doc.add_paragraph(); spacing(p, 24, 0, 1.0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Cm(9.5), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
        p.add_run("\t")
        fmt(p, f"     {R['data_assinatura']}", size=TAM_CORPO, color="1A1A1A")
        q = doc.add_paragraph(); spacing(q, 2, 0, 1.0)
        q.paragraph_format.keep_with_next = bool(detalhe)
        fmt(q, f"{R['assinatura']} — ", size=Pt(10), color="666666")
        fmt(q, papel, bold=True, size=Pt(10), color=pal["primary_dark"])
        if detalhe:
            d = doc.add_paragraph(); spacing(d, 0, 0, 1.0)
            fmt(d, detalhe, size=Pt(9), color="666666")


def _rodape_discreto(section, meta, pal, R):
    """Uma linha: classificação · código · versão · vigência e "Página X de Y"."""
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.clear()
    pp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    spacing(pp, 2, 0)
    pp.paragraph_format.tab_stops.add_tab_stop(LARGURA_UTIL, WD_TAB_ALIGNMENT.RIGHT)
    classe = meta.get("classificacao", "publico")
    fmt(pp, R["class_" + classe].upper(), bold=True, size=TAM_TINY,
        color=pal["risk_red"] if classe == "confidencial" else "666666")
    fmt(pp, f"  ·  {meta['codigo']}  ·  v{meta['versao']}  ·  {R['vigencia']} {meta['data_elaboracao']} "
            f"{R['ate']} {meta['validade']}", size=TAM_TINY, color="888888")
    fmt(pp, f"\t{R['pagina']} ", size=TAM_TINY, color="888888")
    _campo(pp, "PAGE", TAM_TINY, "888888")
    fmt(pp, f" {R['de']} ", size=TAM_TINY, color="888888")
    _campo(pp, "NUMPAGES", TAM_TINY, "888888")


def gerar_texto(dados, pasta=None, cor=None, logo=None, overrides=None, hoje=None):
    """Gera o .docx de um documento de apoio. Devolve o caminho do arquivo."""
    data = carregar_json(dados)
    meta = normalizar_metadata(data.get("metadata", {}), hoje)
    R = rotulos(meta["idioma"])
    pal = paleta(cor, overrides)
    pal["_rot"] = R

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONTE; style.font.size = TAM_CORPO
    style.font.color.rgb = RGBColor.from_string("1A1A1A")
    style.paragraph_format.space_after = Pt(4); style.paragraph_format.line_spacing = 1.15

    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = MARGEM_SUP; sec.bottom_margin = MARGEM_INF
    sec.left_margin = MARGEM_ESQ; sec.right_margin = MARGEM_DIR
    build_header(sec, meta, pal, logo)
    if meta["publico"] == "paciente":
        _rodape_discreto(sec, meta, pal, R)
    else:
        build_footer(sec, meta, pal)

    if doc.paragraphs and doc.paragraphs[0].text == "":
        doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)

    for bloco in data.get("corpo") or []:
        _bloco(doc, bloco, pal)
    _assinaturas(doc, data.get("assinaturas"), pal, R)

    _ajustar_grades(doc)
    titulo = meta.get("titulo_processo") or meta["tipo_documento"]
    nome = f"{nome_arquivo_seguro(meta['codigo'])}_{nome_arquivo_seguro(titulo)}.docx"
    fp = os.path.join(pasta_saida(pasta), nome)
    doc.save(fp)
    return fp
