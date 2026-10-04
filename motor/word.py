"""Renderizador Word (.docx) do POP de Elite.

Portado do app.py da versão 1.0 (Hugging Face Space), com as regras da 2.0:
título por tipo documental, rodapé completo (Elaborado / Revisado / Aprovado /
Vigência, classificação de acesso e "Página X de Y"), validade de 2 anos
calculada em código, rótulos em português/inglês/espanhol, seção de riscos
nomeada pela natureza do processo e larguras de tabela coerentes no preview.
"""

import io
import os
import re
import shutil
import tempfile
import textwrap
import uuid

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

from . import diagramas as dg
from .base import (DEFAULT_PRIMARY, carregar_json, hex_to_rgb, lighten, nome_arquivo_seguro,
                   normalizar_metadata, paleta, pasta_saida, rgb_to_hex, sem_acento, texto_de)
from .i18n import rotulos, tipo_doc


def _rot(pal):
    """Rótulos do idioma do documento (guardados na paleta, que é por documento)."""
    return pal.get("_rot") or rotulos("pt")


def _rotulo_no(valor, largura=18, linhas=3):
    """Texto de um nó de diagrama: quebra em linhas em vez de cortar no meio."""
    t = " ".join(texto_de(valor).split())
    partes = textwrap.wrap(t, largura) or [""]
    if len(partes) > linhas:
        partes = partes[:linhas]
        partes[-1] = partes[-1][:largura - 1].rstrip() + "…"
    return "\n".join(partes)


FONTE = "Calibri"
TAM_CORPO = Pt(11)
TAM_H1 = Pt(12)
TAM_H2 = Pt(11)
TAM_SMALL = Pt(9)
TAM_HEADER = Pt(9)
TAM_TINY = Pt(8)

MARGEM_SUP = Cm(6.0)
MARGEM_INF = Cm(3.0)
MARGEM_ESQ = Cm(2.5)
MARGEM_DIR = Cm(2.0)
LARGURA_DXA = int((21 - 2.5 - 2.0) * 567)

IND_H2 = 0.5
IND_BODY = 0.5
IND_ITEM = 1.0
IND_SUBITEM = 1.5



def set_shading(cell, color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>'))

def set_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for side, val in [("top",top),("bottom",bottom),("left",left),("right",right)]:
        if val:
            c, s = val if isinstance(val, tuple) else (val, "4")
            borders.append(parse_xml(
                f'<w:{side} {nsdecls("w")} w:val="single" w:sz="{s}" w:space="0" w:color="{c}"/>'))
    tcPr.append(borders)

def set_valign(cell, v="center"):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="{v}"/>'))

def set_width(cell, w):
    # GUARDA: substitui a largura existente. O python-docx já cria um w:tcW por célula, e
    # um segundo elemento deixava o LibreOffice (preview) com colunas iguais.
    tcPr = cell._tc.get_or_add_tcPr()
    for antigo in tcPr.findall(qn("w:tcW")):
        tcPr.remove(antigo)
    tcPr.insert(0, parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>'))

def set_margins(cell, t=0, b=0, l=80, r=80):
    cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcMar {nsdecls("w")}><w:top w:w="{t}" w:type="dxa"/>'
        f'<w:left w:w="{l}" w:type="dxa"/><w:bottom w:w="{b}" w:type="dxa"/>'
        f'<w:right w:w="{r}" w:type="dxa"/></w:tcMar>'))

def fmt(p, text, bold=False, italic=False, size=None, color=None, font=FONTE, caps=False):
    run = p.add_run(texto_de(text))
    run.bold = bold; run.italic = italic
    if size: run.font.size = size
    if color: run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    if caps: run.font.all_caps = True
    return run


def _add_highlight_to_run(run, highlight_color="cyan"):
    """Aplica highlight (fundo colorido) a um run via XML.
    Aceita nomes Word (cyan, yellow...) ou hex 6-char para shading."""
    rPr = run._r.get_or_add_rPr()
    if len(highlight_color) == 6 and all(c in '0123456789ABCDEFabcdef' for c in highlight_color):
        rPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{highlight_color}" w:val="clear"/>'))
    else:
        rPr.append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="{highlight_color}"/>'))


def fmt_with_meta_badges(p, text, size=None, color="1A1A1A", font=FONTE, pal=None):
    """Renderiza texto com badges visuais para Metas OMS (Meta 1..Meta 6).
    Metas ficam em bold + highlight com cor terciária da paleta."""
    pattern = re.compile(r'((?:Meta|Goal|IPSG)\s+[1-6])', re.IGNORECASE)
    parts = pattern.split(texto_de(text))
    badge_color = pal["tertiary"] if pal else "cyan"
    badge_text_color = pal["primary_dark"] if pal else "1A3A6E"
    for part in parts:
        if pattern.match(part):
            run = p.add_run(f" {part} ")
            run.bold = True
            run.font.name = font
            if size: run.font.size = size
            run.font.color.rgb = RGBColor.from_string(badge_text_color)
            _add_highlight_to_run(run, badge_color)
        else:
            if part:
                run = p.add_run(part)
                run.font.name = font
                if size: run.font.size = size
                if color: run.font.color.rgb = RGBColor.from_string(color)

def spacing(p, before=0, after=0, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line

def p_shading(p, color):
    p._p.get_or_add_pPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>'))

def repeat_header(row):
    row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def build_header(section, meta, pal, logo_bytes=None):
    R = _rot(pal)
    titulo_tipo, _ = tipo_doc(meta.get("tipo_documento"), meta.get("idioma"))
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs: p.clear()

    tbl = header.add_table(rows=4, cols=3, width=Cm(16.5))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cw = [1700, 5800, 2900]
    brd = (pal["gray_border"], "4")

    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            set_width(cell, cw[i]); set_valign(cell)
            set_margins(cell, 30, 30, 80, 80)
            set_borders(cell, brd, brd, brd, brd)

    logo_cell = tbl.cell(0,0).merge(tbl.cell(3,0))
    logo_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo_bytes:
        run = logo_cell.paragraphs[0].add_run()
        run.add_picture(io.BytesIO(logo_bytes), height=Cm(1.8))
    else:
        fmt(logo_cell.paragraphs[0], R["logo"], bold=True, size=Pt(10), color="888888")

    tc = tbl.cell(0,1).merge(tbl.cell(1,1))
    tc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(tc.paragraphs[0], titulo_tipo, bold=True, size=Pt(13), color="000000")

    pc = tbl.cell(2,1).merge(tbl.cell(3,1))
    pc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(pc.paragraphs[0], texto_de(meta.get("titulo_processo","")).upper(),
        bold=True, size=Pt(12), color="000000")

    for i, (lbl, val) in enumerate([
        (R["codigo"], meta.get("codigo","")),
        (R["elaborado_em"], meta.get("data_elaboracao","")),
        (R["revisado_em"], meta.get("data_revisao","")),
        (R["valido_ate"], meta.get("validade","")),
    ]):
        c = tbl.cell(i, 2)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt(c.paragraphs[0], lbl, bold=True, size=TAM_HEADER, color="000000")
        fmt(c.paragraphs[0], texto_de(val) or "—", size=TAM_HEADER, color="000000")


def _campo(p, instrucao, tam, cor):
    """Campo do Word (PAGE / NUMPAGES), atualizado ao abrir e imprimir."""
    def _run():
        # GUARDA: tamanho em todos os runs do campo: o LibreOffice formata o número pelo run inicial
        r = p.add_run(); r.font.size = tam; r.font.color.rgb = RGBColor.from_string(cor)
        return r
    for x in [f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>',
              f'<w:instrText {nsdecls("w")} xml:space="preserve"> {instrucao} </w:instrText>',
              f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>']:
        _run()._r.append(parse_xml(x))
    _run().text = "1"
    _run()._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


def build_footer(section, meta, pal):
    """Rodapé completo: Elaborado / Revisado / Aprovado / Vigência, e na linha de
    baixo a classificação de acesso, o código, a versão e "Página X de Y"."""
    R = _rot(pal)
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs: p.clear()

    footer.add_paragraph()
    tbl = footer.add_table(rows=2, cols=4, width=Cm(16.5))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cw = [2450, 2450, 2450, 2005]
    brd = (pal["gray_border"], "4")

    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            set_width(cell, cw[i]); set_valign(cell)
            set_margins(cell, 30, 30, 60, 60)
            set_borders(cell, brd, brd, brd, brd)

    for i, h in enumerate([R["elaborado_por"], R["revisado_por"], R["aprovado_por"], R["vigencia"]]):
        c = tbl.cell(0, i)
        set_shading(c, pal["tertiary"])
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt(c.paragraphs[0], h, bold=True, size=TAM_SMALL, color=pal["primary_dark"])

    for i, key in enumerate(["elaborado_por", "revisado_por", "aprovado_por"]):
        c = tbl.cell(1, i)
        person = meta.get(key, {})
        nome = texto_de(person.get("nome","")) if isinstance(person, dict) else ""
        cargo = texto_de(person.get("cargo","")) if isinstance(person, dict) else ""
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if nome:
            fmt(c.paragraphs[0], nome, bold=True, size=TAM_SMALL, color=pal["text_body"])
            p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(p2, cargo, size=TAM_TINY, color="666666")
        else:
            fmt(c.paragraphs[0], "________________", size=TAM_SMALL, color="AAAAAA")

    cv = tbl.cell(1, 3)
    cv.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(cv.paragraphs[0], f"{meta.get('data_elaboracao','')} {R['ate']}", size=TAM_TINY, color=pal["text_body"])
    pv = cv.add_paragraph(); pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(pv, meta.get("validade", ""), bold=True, size=TAM_SMALL, color=pal["text_body"])

    pp = footer.add_paragraph()
    spacing(pp, 2, 0)
    pp.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)
    classe = meta.get("classificacao", "interno")
    fmt(pp, R["class_" + classe].upper(), bold=True, size=TAM_TINY,
        color=pal["risk_red"] if classe == "confidencial" else "666666")
    fmt(pp, f"  ·  {meta.get('codigo','')}  ·  v{meta.get('versao','01')}", size=TAM_TINY, color="888888")
    fmt(pp, f"\t{R['pagina']} ", size=TAM_TINY, color="888888")
    _campo(pp, "PAGE", TAM_TINY, "888888")
    fmt(pp, f" {R['de']} ", size=TAM_TINY, color="888888")
    _campo(pp, "NUMPAGES", TAM_TINY, "888888")


def _ajustar_grades(doc):
    """Faz a grade de cada tabela (w:gridCol) refletir a largura das células.

    O python-docx cria a grade com colunas iguais; o LibreOffice usa a grade e
    ignora a largura da célula — a coluna do ✓ ocupava meia página e o fluxo
    de 6 etapas saía cortado no preview."""
    partes = [doc.element.body]
    for s in doc.sections:
        partes += [s.header._element, s.footer._element]
    for parte in partes:
        for tbl in parte.iter(qn("w:tbl")):
            grade = tbl.find(qn("w:tblGrid"))
            linha = tbl.find(qn("w:tr"))
            if grade is None or linha is None:
                continue
            colunas = grade.findall(qn("w:gridCol"))
            celulas = linha.findall(qn("w:tc"))
            if len(colunas) != len(celulas):
                continue
            larguras = []
            for tc in celulas:
                tcW = tc.find(qn("w:tcPr") + "/" + qn("w:tcW"))
                if tcW is None or tcW.get(qn("w:type")) != "dxa":
                    larguras = None
                    break
                larguras.append(tcW.get(qn("w:w")))
            if larguras:
                for col, w in zip(colunas, larguras):
                    col.set(qn("w:w"), str(w))



def add_h1(doc, num, titulo, pal):
    p = doc.add_paragraph(); spacing(p, 16, 8)
    p_shading(p, pal["primary"])
    pf = p.paragraph_format; pf.left_indent = Cm(0.3); pf.right_indent = Cm(0.3)
    fmt(p, f"  {num}.  {titulo.upper()}", bold=True, size=TAM_H1, color=pal["text_on_primary"])
    pf.keep_with_next = True      # GUARDA: título nunca fica sozinho no pé da página
    return p

def add_h2(doc, num, titulo, pal):
    p = doc.add_paragraph(); spacing(p, 12, 6)
    p_shading(p, pal["secondary"])
    pf = p.paragraph_format; pf.left_indent = Cm(IND_H2 + 0.3); pf.right_indent = Cm(0.3)
    fmt(p, f"  {num}.  {titulo.upper()}", bold=True, size=TAM_H2, color=pal["text_on_secondary"])
    pf.keep_with_next = True

def add_body(doc, text, indent=IND_BODY):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p, 2, 5, 1.15); p.paragraph_format.left_indent = Cm(indent)
    fmt(p, text, size=TAM_CORPO, color="1A1A1A")

def add_bullet(doc, text, bold_prefix=None, indent=IND_ITEM, pal=None):
    c = pal["primary"] if pal else DEFAULT_PRIMARY
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p, 1, 3, 1.15)
    p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.first_line_indent = Cm(-0.4)
    fmt(p, "●  ", size=Pt(9), color=c, bold=True)
    if bold_prefix:
        fmt(p, bold_prefix, bold=True, size=TAM_CORPO, color=c)
        fmt(p, text, size=TAM_CORPO, color="1A1A1A")
    else:
        fmt(p, text, size=TAM_CORPO, color="1A1A1A")

def add_num(doc, n, text, indent=IND_SUBITEM, pal=None, critico=False):
    c = pal["primary"] if pal else DEFAULT_PRIMARY
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p, 1, 4, 1.15)
    p.paragraph_format.left_indent = Cm(indent); p.paragraph_format.first_line_indent = Cm(-0.5)
    if critico:
        fmt(p, f"\u26a0  ", bold=True, size=TAM_CORPO, color="C05000")
        fmt(p, f"{n}. {_rot(pal)['passo_critico']}  ", bold=True, size=TAM_CORPO, color="C05000")
    else:
        fmt(p, f"{n}.  ", bold=True, size=TAM_CORPO, color=c)
    fmt_with_meta_badges(p, text, size=TAM_CORPO, pal=pal)

def add_def_item(doc, termo, defn, pal):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p, 2, 5, 1.15)
    p.paragraph_format.left_indent = Cm(IND_ITEM); p.paragraph_format.first_line_indent = Cm(-0.4)
    fmt(p, "●  ", size=Pt(9), color=pal["primary"], bold=True)
    fmt(p, f"{termo}:  ", bold=True, size=TAM_CORPO, color=pal["primary"])
    fmt(p, defn, size=TAM_CORPO, color="1A1A1A")

def add_risk(doc, risco, barreira, pal):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p, 4, 3, 1.15)
    p.paragraph_format.left_indent = Cm(IND_ITEM); p.paragraph_format.first_line_indent = Cm(-0.5)
    fmt(p, "⚠  ", size=TAM_CORPO, color=pal["risk_red"], bold=True)
    fmt(p, _rot(pal)["risco"], bold=True, size=TAM_CORPO, color=pal["risk_red"])
    fmt(p, risco, size=TAM_CORPO, color="1A1A1A")
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spacing(p2, 0, 8, 1.15); p2.paragraph_format.left_indent = Cm(IND_SUBITEM)
    fmt(p2, _rot(pal)["barreira"], bold=True, size=TAM_CORPO, color=pal["barrier_green"])
    fmt(p2, barreira, size=TAM_CORPO, color="1A1A1A")




def add_table(doc, headers, rows, col_widths=None, pal=None):
    nc = len(headers)
    if not col_widths: col_widths = [LARGURA_DXA // nc] * nc
    tbl = doc.add_table(rows=1+len(rows), cols=nc)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    brd = (pal["border_light"], "4") if pal else ("3B6AA0", "4")
    brd2 = (pal["gray_border"], "2") if pal else ("7F8C9A", "2")

    for i, h in enumerate(headers):
        c = tbl.cell(0, i); set_shading(c, pal["primary"] if pal else DEFAULT_PRIMARY)
        set_width(c, col_widths[i]); set_valign(c)
        set_margins(c, 50, 50, 80, 80); set_borders(c, brd, brd, brd, brd)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt(c.paragraphs[0], h, bold=True, size=TAM_SMALL, color="FFFFFF")
    repeat_header(tbl.rows[0])

    for ri, rd in enumerate(rows):
        for ci, val in enumerate(rd):
            c = tbl.cell(ri+1, ci); set_width(c, col_widths[ci]); set_valign(c)
            set_margins(c, 40, 40, 80, 80); set_borders(c, brd2, brd2, brd2, brd2)
            a = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 and nc > 2 else WD_ALIGN_PARAGRAPH.CENTER
            c.paragraphs[0].alignment = a
            fmt(c.paragraphs[0], str(val), size=TAM_SMALL, color="1A1A1A")
        if ri % 2 == 0:
            for ci in range(nc):
                set_shading(tbl.cell(ri+1, ci), pal["tertiary"] if pal else "E8EFF7")
    doc.add_paragraph()




def add_process(doc, titulo, etapas, pal):
    if not etapas: return
    p = doc.add_paragraph(); spacing(p, 8, 6)
    p.paragraph_format.left_indent = Cm(IND_BODY)
    fmt(p, f"▸  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])

    n = len(etapas); total = n*2-1; aw = 400
    bw = (LARGURA_DXA - aw*(n-1)) // n
    tbl = doc.add_table(rows=1, cols=total)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    colors = [pal["primary"], pal["primary"], pal["primary"], pal["primary"]]

    for ci in range(total):
        c = tbl.cell(0, ci)
        if ci % 2 == 0:
            ei = ci // 2
            set_width(c, bw); set_shading(c, pal["primary"] if ei % 2 == 0 else lighten(pal["primary"], 0.15))
            bg = pal["primary"] if ei % 2 == 0 else lighten(pal["primary"], 0.15)
            set_borders(c, (bg,"6"), (bg,"6"), (bg,"6"), (bg,"6"))
            set_valign(c); set_margins(c, 60, 60, 80, 80)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(c.paragraphs[0], f"{_rot(pal)['etapa']} {ei+1}", bold=True, size=TAM_TINY, color="FFFFFF")
            p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(p2, etapas[ei], size=TAM_SMALL, color="FFFFFF")
        else:
            set_width(c, aw)
            set_borders(c,("FFFFFF","0"),("FFFFFF","0"),("FFFFFF","0"),("FFFFFF","0"))
            set_valign(c); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(c.paragraphs[0], "→", bold=True, size=Pt(16), color=pal["primary"])
    doc.add_paragraph()


def add_checklist(doc, titulo, itens, pal):
    if not itens: return
    p = doc.add_paragraph(); spacing(p, 8, 6)
    p.paragraph_format.left_indent = Cm(IND_BODY)
    fmt(p, f"▸  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])

    tbl = doc.add_table(rows=len(itens), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    chw = 600; tw = LARGURA_DXA - chw

    for ri, item in enumerate(itens):
        cc = tbl.cell(ri, 0); set_width(cc, chw); set_valign(cc)
        set_margins(cc, 40, 40, 40, 40); set_shading(cc, pal["primary"])
        set_borders(cc,("FFFFFF","2"),("FFFFFF","2"),(pal["primary"],"4"),("FFFFFF","2"))
        cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt(cc.paragraphs[0], "✓", bold=True, size=Pt(14), color="FFFFFF")

        ct = tbl.cell(ri, 1); set_width(ct, tw); set_valign(ct)
        set_margins(ct, 50, 50, 120, 80)
        bg = pal["quaternary"] if ri % 2 == 0 else "FFFFFF"
        set_shading(ct, bg)
        set_borders(ct,(pal["gray_border"],"2"),(pal["gray_border"],"2"),("FFFFFF","0"),(pal["gray_border"],"2"))
        ct.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt(ct.paragraphs[0], item, size=TAM_CORPO, color="1A1A1A")
    doc.add_paragraph()


def add_cycle(doc, titulo, etapas, pal):
    if not etapas: return
    p = doc.add_paragraph(); spacing(p, 8, 6)
    p.paragraph_format.left_indent = Cm(IND_BODY)
    fmt(p, f"▸  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])

    n = len(etapas); cpr = min(n, 4); rows_n = (n+cpr-1)//cpr
    bw = LARGURA_DXA // cpr
    tbl = doc.add_table(rows=rows_n, cols=cpr)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, et in enumerate(etapas):
        r, c_idx = idx//cpr, idx%cpr
        c = tbl.cell(r, c_idx); set_width(c, bw)
        bg = pal["primary"] if idx % 2 == 0 else lighten(pal["primary"], 0.15)
        set_shading(c, bg)
        set_borders(c,("FFFFFF","4"),("FFFFFF","4"),("FFFFFF","4"),("FFFFFF","4"))
        set_valign(c); set_margins(c, 60, 60, 80, 80)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        arrow = "  →" if idx < n-1 else "  ⟳"
        fmt(c.paragraphs[0], f"{idx+1}{arrow}", bold=True, size=TAM_TINY, color="FFFFFF")
        p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt(p2, et, size=TAM_SMALL, color="FFFFFF")

    for idx in range(n, rows_n*cpr):
        c = tbl.cell(idx//cpr, idx%cpr)
        set_borders(c,("FFFFFF","0"),("FFFFFF","0"),("FFFFFF","0"),("FFFFFF","0"))
    doc.add_paragraph()




ANNEX_BOX_HEIGHT_DXA = 8200

def _set_row_height(row, height_dxa):
    """Define altura mínima de uma linha de tabela."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(
        f'<w:trHeight {nsdecls("w")} w:val="{height_dxa}" w:hRule="atLeast"/>'))

def start_annex_border(doc, pal):
    """Cria container bordado com altura mínima full-page."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    brd = (pal["annex_border"], "6")
    set_borders(c, brd, brd, brd, brd)
    set_width(c, LARGURA_DXA)
    set_margins(c, 150, 150, 250, 250)
    _set_row_height(tbl.rows[0], ANNEX_BOX_HEIGHT_DXA)
    # GUARDA: sem cantSplit. O anexo já abre página; se o conteúdo não couber nela, a caixa inteira
    # pularia para a seguinte, deixando o título sozinho, e se partiria do mesmo jeito.
    return c, tbl




HAS_GRAPHVIZ = False


def gerar_diagrama_png(tipo, conteudo, titulo, pal, context="annex", raias=None):
    """Desenha o diagrama em PNG. Devolve (caminho, largura_cm, altura_cm) ou None.
    context: 'inline' (no corpo do procedimento, compacto) ou 'annex' (página de anexo)."""
    R = _rot(pal)
    pasta = pal.get("_tmp") or tempfile.gettempdir()
    rot = {"sim": R["sim"], "nao": R["nao"], "inicio": R.get("inicio", "Início"), "fim": R.get("fim", "Fim")}
    anexo = context != "inline"
    try:
        if tipo in ("processo", "fluxograma"):
            if not anexo:
                return dg.fluxo_png(conteudo, pal, pasta, titulo=titulo, rotulos=rot,
                                    largura_cm=14.0, altura_cm=6.5, com_piscina=False)
            return dg.fluxo_png(conteudo, pal, pasta, raias=raias, titulo=titulo, rotulos=rot,
                                largura_cm=15.4, altura_cm=15.2)
        if tipo == "ciclo":
            return dg.ciclo_png(conteudo, pal, pasta, largura_cm=13.5 if anexo else 11.5)
        if tipo == "piramide":
            return dg.piramide_png(conteudo, pal, pasta, largura_cm=13.5 if anexo else 11.0)
        if tipo == "hierarquia":
            return dg.hierarquia_png(conteudo, pal, pasta, largura_cm=15.0 if anexo else 13.0)
    except Exception:
        return None
    return None


def _inserir_diagrama_na_celula(cell, img_path):
    """Insere imagem PNG do diagrama dentro de uma célula do docx.
    Com (caminho, largura_cm, altura_cm), usa o tamanho calculado pelo desenho;
    com só o caminho, limita a 15 x 10 cm para caber na caixa do anexo."""
    if isinstance(img_path, tuple):
        caminho, w_cm, _ = img_path
        vazio = len(cell.paragraphs) == 1 and not cell.paragraphs[0].text and not cell.paragraphs[0].runs
        p = cell.paragraphs[0] if vazio else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(caminho, width=Cm(w_cm))
        return True
    if not img_path or not os.path.exists(img_path):
        return False
    try:
        from PIL import Image as PILImage
        img = PILImage.open(img_path)
        w_px, h_px = img.size
        max_w = Cm(15)
        max_h = Cm(10)
        ratio = w_px / h_px
        width = max_w
        height = int(width / ratio) if ratio > 0 else max_h
        if height > max_h:
            height = max_h
            width = int(height * ratio)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(img_path, width=width)
        return True
    except Exception:
        try:
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(img_path, width=Cm(12))
            return True
        except Exception:
            return False


def _inserir_diagrama_no_corpo(doc, img_path, legenda="", pal=None):
    """Insere imagem PNG de diagrama no corpo do documento (fora de tabela/célula).
    Centralizado, com legenda em itálico abaixo."""
    largura = Cm(14)
    if isinstance(img_path, tuple):
        img_path, w_cm, _ = img_path
        largura = Cm(w_cm)
    if not img_path or not os.path.exists(img_path):
        return
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=largura)
        if legenda:
            pl = doc.add_paragraph()
            pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pl.paragraph_format.space_before = Pt(2)
            pl.paragraph_format.space_after = Pt(8)
            fmt(pl, legenda, italic=True, size=Pt(9),
                color=pal["primary"] if pal else "333333")
    except Exception:
        pass


def _render_passo(doc, num, item, pal):
    """Renderiza um item de procedimento: string, dict com texto/prioridade, ou dict com diagrama.
    Formatos aceitos:
      - "texto simples"
      - {"texto": "...", "prioridade": "critica"}
      - {"diagrama": {"tipo":..., "titulo":..., "conteudo":...}}
      - {"tipo":..., "titulo":..., "conteudo":...}
    """
    if isinstance(item, str):
        add_num(doc, num, item, pal=pal)
    elif isinstance(item, dict):
        if "texto" in item and "diagrama" not in item and "tipo" not in item:
            texto = item.get("texto", "")
            critico = item.get("prioridade", "").lower() == "critica"
            add_num(doc, num, texto, pal=pal, critico=critico)
            return

        diag = item.get("diagrama", None)
        if diag is None and "tipo" in item:
            diag = item
        if diag is None:
            add_num(doc, num, str(item), pal=pal)
            return
        tipo = diag.get("tipo", "processo")
        titulo = diag.get("titulo", "Diagrama")
        conteudo = diag.get("conteudo", [])
        img_path = gerar_diagrama_png(tipo, conteudo, titulo, pal, context="inline")
        if img_path:
            _inserir_diagrama_no_corpo(doc, img_path, titulo, pal)
        else:
            items_txt = " → ".join(texto_de(c) for c in conteudo) if isinstance(conteudo, list) else texto_de(conteudo)
            add_num(doc, num, f"[{titulo}]: {items_txt}", pal=pal)



def render_annex(doc, anexo, num, pal):
    R = _rot(pal)
    titulo = texto_de(anexo.get("titulo", "")) if isinstance(anexo, dict) else str(anexo)
    # GUARDA: "quebra antes" no próprio título (e não um parágrafo com quebra): se o conteúdo anterior
    # terminar rente ao pé da página, não sobra uma página em branco antes do anexo
    cab = add_h1(doc, R["anexo_h1"].format(n=num), titulo.upper(), pal).paragraph_format
    cab.page_break_before = True
    cab.space_before = Pt(0)
    cab.keep_with_next = False

    if isinstance(anexo, str):
        add_body(doc, R["anexo_a_inserir"].format(t=titulo))
        return

    if anexo.get("descricao"):
        add_body(doc, anexo["descricao"])

    tipo = anexo.get("tipo", "texto")
    conteudo = anexo.get("conteudo", "")

    bc, border_tbl = start_annex_border(doc, pal)

    desenhado = False
    if tipo in ("processo", "fluxograma", "hierarquia", "piramide", "ciclo"):
        img_path = gerar_diagrama_png(tipo, conteudo, titulo, pal, raias=anexo.get("raias"))
        if img_path:
            _inserir_diagrama_na_celula(bc, img_path)
            desenhado = True

    if desenhado:
        pass
    elif tipo == "checklist":
        itens = conteudo if isinstance(conteudo, list) else [conteudo]
        pt = bc.paragraphs[0]; pt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt(pt, f"▸  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])
        inner = bc.add_table(rows=len(itens), cols=2)
        chw = 500; tw = LARGURA_DXA - 900
        for ri, item in enumerate(itens):
            cc = inner.cell(ri, 0); set_width(cc, chw); set_valign(cc)
            set_margins(cc, 30, 30, 30, 30); set_shading(cc, pal["primary"])
            set_borders(cc,("FFFFFF","2"),("FFFFFF","2"),(pal["primary"],"4"),("FFFFFF","2"))
            cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(cc.paragraphs[0], "✓", bold=True, size=Pt(12), color="FFFFFF")
            ct = inner.cell(ri, 1); set_width(ct, tw); set_valign(ct)
            set_margins(ct, 40, 40, 100, 60)
            set_shading(ct, pal["quaternary"] if ri % 2 == 0 else "FFFFFF")
            set_borders(ct,(pal["gray_border"],"1"),(pal["gray_border"],"1"),("FFFFFF","0"),(pal["gray_border"],"1"))
            fmt(ct.paragraphs[0], item, size=TAM_SMALL, color="1A1A1A")

    elif tipo in ("escala", "tabela", "referencia", "referência"):
        headers = anexo.get("colunas", [])
        rows_data = anexo.get("linhas", [])
        if not headers and isinstance(conteudo, list) and len(conteudo) > 0 and isinstance(conteudo[0], dict):
            tbl_obj = conteudo[0]
            headers = tbl_obj.get("colunas", [])
            rows_data = tbl_obj.get("linhas", [])
        if not rows_data and isinstance(conteudo, list) and len(conteudo) > 0 and isinstance(conteudo[0], list):
            rows_data = conteudo
        if headers and rows_data:
            bc.paragraphs[0].text = ""
            nc = len(headers); cw_each = (LARGURA_DXA - 900) // nc
            inner = bc.add_table(rows=1+len(rows_data), cols=nc)
            for i, h in enumerate(headers):
                c = inner.cell(0, i); set_shading(c, pal["primary"])
                set_width(c, cw_each); set_valign(c); set_margins(c, 40, 40, 60, 60)
                set_borders(c,(pal["primary"],"4"),(pal["primary"],"4"),(pal["primary"],"4"),(pal["primary"],"4"))
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt(c.paragraphs[0], h, bold=True, size=TAM_SMALL, color="FFFFFF")
            for ri, rd in enumerate(rows_data):
                for ci, val in enumerate(rd):
                    c = inner.cell(ri+1, ci); set_width(c, cw_each); set_valign(c)
                    set_margins(c, 30, 30, 60, 60)
                    set_borders(c,(pal["gray_border"],"1"),(pal["gray_border"],"1"),
                                (pal["gray_border"],"1"),(pal["gray_border"],"1"))
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fmt(c.paragraphs[0], str(val), size=TAM_SMALL, color="1A1A1A")
                if ri % 2 == 0:
                    for ci in range(nc): set_shading(inner.cell(ri+1, ci), pal["tertiary"])

    elif tipo in ("processo", "fluxograma"):
        etapas = conteudo if isinstance(conteudo, list) else [conteudo]
        etapas = etapas[:6] or [""]
        fmt(bc.paragraphs[0], f"▸  {R['fluxo_processo']}", bold=True, size=TAM_CORPO, color=pal["primary"])
        n = len(etapas); total = n*2-1; aw = 350
        bw = (LARGURA_DXA - 900 - aw*(n-1)) // n
        inner = bc.add_table(rows=1, cols=total)
        for ci in range(total):
            c = inner.cell(0, ci)
            if ci % 2 == 0:
                ei = ci // 2; set_width(c, bw)
                bg = pal["primary"] if ei % 2 == 0 else lighten(pal["primary"], 0.15)
                set_shading(c, bg); set_borders(c,(bg,"4"),(bg,"4"),(bg,"4"),(bg,"4"))
                set_valign(c); set_margins(c, 50, 50, 60, 60)
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt(c.paragraphs[0], f"{R['etapa']} {ei+1}", bold=True, size=TAM_TINY, color="FFFFFF")
                p2 = c.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt(p2, etapas[ei], size=Pt(8), color="FFFFFF")
            else:
                set_width(c, aw)
                set_borders(c,("FFFFFF","0"),("FFFFFF","0"),("FFFFFF","0"),("FFFFFF","0"))
                set_valign(c); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt(c.paragraphs[0], "\u2192", bold=True, size=Pt(14), color=pal["primary"])

    elif tipo == "hierarquia":
        items = conteudo if isinstance(conteudo, list) else [conteudo]
        fmt(bc.paragraphs[0], f"\u25b8  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])
        for item in items:
            if isinstance(item, dict):
                p1 = bc.add_paragraph()
                p1.paragraph_format.space_before = Pt(6)
                p1_shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{pal["primary"]}" w:val="clear"/>')
                p1._p.get_or_add_pPr().append(p1_shading)
                fmt(p1, f"  {item.get('titulo', '')}", bold=True, size=TAM_CORPO, color="FFFFFF")
                for sub in item.get("subitens", []):
                    p2 = bc.add_paragraph()
                    p2.paragraph_format.left_indent = Cm(1.0)
                    p2.paragraph_format.space_before = Pt(2)
                    p2_shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{pal["tertiary"]}" w:val="clear"/>')
                    p2._p.get_or_add_pPr().append(p2_shading)
                    fmt(p2, f"  \u25b9  {sub}", size=TAM_SMALL, color=pal["primary_dark"])
            elif isinstance(item, str):
                p1 = bc.add_paragraph()
                p1_shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{pal["secondary"]}" w:val="clear"/>')
                p1._p.get_or_add_pPr().append(p1_shading)
                fmt(p1, f"  {item}", bold=True, size=TAM_SMALL, color=pal["primary_dark"])

    elif tipo == "piramide":
        items = conteudo if isinstance(conteudo, list) else [conteudo]
        fmt(bc.paragraphs[0], f"\u25b8  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])
        n = len(items)
        max_w = LARGURA_DXA - 900
        for i, item in enumerate(items):
            level_w = int(max_w * (0.4 + 0.6 * i / max(n-1, 1)))
            margin_l = (max_w - level_w) // 2
            inner = bc.add_table(rows=1, cols=1)
            inner.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = inner.cell(0, 0)
            set_width(c, level_w)
            frac = i / max(n-1, 1)
            r1, g1, b1 = hex_to_rgb(pal["primary"])
            r2, g2, b2 = hex_to_rgb(pal["tertiary"])
            bg = rgb_to_hex(r1+(r2-r1)*frac, g1+(g2-g1)*frac, b1+(b2-b1)*frac)
            set_shading(c, bg)
            brd = (bg, "4")
            set_borders(c, brd, brd, brd, brd)
            set_valign(c); set_margins(c, 40, 40, 80, 80)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            txt_color = "FFFFFF" if frac < 0.5 else pal["primary_dark"]
            fmt(c.paragraphs[0], str(item), bold=True, size=TAM_SMALL, color=txt_color)

    elif tipo == "ciclo":
        etapas = (conteudo if isinstance(conteudo, list) else [conteudo]) or [""]
        fmt(bc.paragraphs[0], f"\u25b8  {titulo}", bold=True, size=TAM_CORPO, color=pal["primary"])
        n = len(etapas)
        if n == 4:
            inner = bc.add_table(rows=2, cols=2)
            inner.alignment = WD_TABLE_ALIGNMENT.CENTER
            positions = [(0,0),(0,1),(1,1),(1,0)]
            colors = [pal["primary"], lighten(pal["primary"],0.15),
                      lighten(pal["primary"],0.30), lighten(pal["primary"],0.45)]
            bw = (LARGURA_DXA - 900) // 2
            for idx, (r,c_idx) in enumerate(positions):
                c = inner.cell(r, c_idx); set_width(c, bw); set_valign(c)
                set_shading(c, colors[idx])
                brd = (colors[idx], "4")
                set_borders(c, brd, brd, brd, brd)
                set_margins(c, 50, 50, 80, 80)
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt(c.paragraphs[0], etapas[idx], bold=True, size=TAM_SMALL, color="FFFFFF")
            p_arrow = bc.add_paragraph()
            p_arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt(p_arrow, f"\u21bb  {R['ciclo_continuo']}", bold=True, size=TAM_SMALL, color=pal["primary"])
        else:
            cpr = min(n, 4); rows_n = (n+cpr-1)//cpr
            bw = (LARGURA_DXA - 900) // cpr
            inner = bc.add_table(rows=rows_n, cols=cpr)
            inner.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, et in enumerate(etapas):
                r, ci2 = divmod(idx, cpr)
                c = inner.cell(r, ci2); set_width(c, bw); set_valign(c)
                bg = pal["primary"] if idx % 2 == 0 else lighten(pal["primary"], 0.20)
                set_shading(c, bg)
                brd = (bg, "4")
                set_borders(c, brd, brd, brd, brd)
                set_margins(c, 40, 40, 60, 60)
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt(c.paragraphs[0], et, bold=True, size=TAM_TINY, color="FFFFFF")

    elif tipo == "misto":
        bc.paragraphs[0].text = ""
        elementos = conteudo if isinstance(conteudo, list) else [conteudo]
        for elem in elementos:
            if isinstance(elem, str):
                pe = bc.add_paragraph(); fmt(pe, elem, size=TAM_CORPO, color="1A1A1A")
            elif isinstance(elem, dict):
                st = elem.get("tipo", "texto")
                if st == "texto":
                    pe = bc.add_paragraph(); fmt(pe, elem.get("conteudo",""), size=TAM_CORPO, color="1A1A1A")
                elif st == "checklist":
                    pe = bc.add_paragraph()
                    fmt(pe, f"\n\u25b8  {elem.get('titulo','')}", bold=True, size=TAM_CORPO, color=pal["primary"])
                    for it in elem.get("itens", []):
                        pi = bc.add_paragraph(); pi.paragraph_format.left_indent = Cm(0.5)
                        fmt(pi, "\u2713  ", bold=True, size=TAM_SMALL, color=pal["primary"])
                        fmt(pi, it, size=TAM_SMALL, color="1A1A1A")
                elif st == "processo":
                    pe = bc.add_paragraph()
                    fmt(pe, f"\n\u25b8  {elem.get('titulo','')}", bold=True, size=TAM_CORPO, color=pal["primary"])
                    desenho = gerar_diagrama_png("processo", elem.get("etapas", elem.get("conteudo", [])),
                                                 elem.get("titulo", ""), pal, context="inline")
                    if desenho:
                        _inserir_diagrama_na_celula(bc, desenho)
                        continue
                    for i, et in enumerate(elem.get("etapas",[]), 1):
                        pi = bc.add_paragraph(); pi.paragraph_format.left_indent = Cm(0.5)
                        fmt(pi, f"{i}.  ", bold=True, size=TAM_SMALL, color=pal["primary"])
                        fmt(pi, et, size=TAM_SMALL, color="1A1A1A")
                elif st == "tabela":
                    pe = bc.add_paragraph()
                    fmt(pe, f"\n\u25b8  {elem.get('titulo','')}", bold=True, size=TAM_CORPO, color=pal["primary"])
                    hdrs = elem.get("colunas",[]); rws = elem.get("linhas",[])
                    if hdrs and rws:
                        nc2 = len(hdrs); cw2 = (LARGURA_DXA - 900) // nc2
                        it2 = bc.add_table(rows=1+len(rws), cols=nc2)
                        for i, h in enumerate(hdrs):
                            cc2 = it2.cell(0,i); set_shading(cc2, pal["primary"])
                            set_width(cc2,cw2); set_valign(cc2); set_margins(cc2,30,30,50,50)
                            set_borders(cc2,(pal["primary"],"3"),(pal["primary"],"3"),
                                       (pal["primary"],"3"),(pal["primary"],"3"))
                            cc2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            fmt(cc2.paragraphs[0], h, bold=True, size=TAM_TINY, color="FFFFFF")
                        for ri2, rd2 in enumerate(rws):
                            for ci2, v2 in enumerate(rd2):
                                cc2 = it2.cell(ri2+1,ci2); set_width(cc2,cw2); set_valign(cc2)
                                set_margins(cc2,25,25,50,50)
                                set_borders(cc2,(pal["gray_border"],"1"),(pal["gray_border"],"1"),
                                           (pal["gray_border"],"1"),(pal["gray_border"],"1"))
                                cc2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                                fmt(cc2.paragraphs[0], str(v2), size=TAM_TINY, color="1A1A1A")
                            if ri2 % 2 == 0:
                                for ci2 in range(nc2): set_shading(it2.cell(ri2+1,ci2), pal["tertiary"])

    elif tipo in ("glossario", "glossário", "simbolos", "símbolos", "glossario_visual"):
        bc.paragraphs[0].text = ""
        if isinstance(conteudo, list):
            if len(conteudo) > 0 and isinstance(conteudo[0], dict):
                tbl_obj = conteudo[0]
                if "colunas" in tbl_obj:
                    headers = tbl_obj.get("colunas", [])
                    rows_data = tbl_obj.get("linhas", [])
                    if headers and rows_data:
                        nc = len(headers); cw_each = (LARGURA_DXA - 900) // nc
                        inner = bc.add_table(rows=1+len(rows_data), cols=nc)
                        for i, h in enumerate(headers):
                            c = inner.cell(0, i); set_shading(c, pal["primary"])
                            set_width(c, cw_each); set_valign(c); set_margins(c, 40, 40, 60, 60)
                            set_borders(c,(pal["primary"],"4"),(pal["primary"],"4"),(pal["primary"],"4"),(pal["primary"],"4"))
                            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            fmt(c.paragraphs[0], h, bold=True, size=TAM_SMALL, color="FFFFFF")
                        for ri, rd in enumerate(rows_data):
                            for ci, val in enumerate(rd):
                                c = inner.cell(ri+1, ci); set_width(c, cw_each); set_valign(c)
                                set_margins(c, 30, 30, 60, 60)
                                set_borders(c,(pal["gray_border"],"1"),(pal["gray_border"],"1"),
                                            (pal["gray_border"],"1"),(pal["gray_border"],"1"))
                                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                                fmt(c.paragraphs[0], str(val), size=TAM_SMALL, color="1A1A1A")
                            if ri % 2 == 0:
                                for ci in range(nc): set_shading(inner.cell(ri+1, ci), pal["tertiary"])
                else:
                    for item in conteudo:
                        pi = bc.add_paragraph()
                        pi.paragraph_format.left_indent = Cm(0.5)
                        simb = item.get("simbolo", item.get("icone", "●"))
                        nome = item.get("nome", item.get("descricao", str(item)))
                        fmt(pi, f"{simb}  {nome}", size=TAM_CORPO, color="1A1A1A")
            else:
                for item in conteudo:
                    pi = bc.add_paragraph()
                    pi.paragraph_format.left_indent = Cm(0.5)
                    fmt(pi, str(item), size=TAM_CORPO, color="1A1A1A")

    else:
        if isinstance(conteudo, str):
            fmt(bc.paragraphs[0], conteudo, size=TAM_CORPO, color="1A1A1A")
        elif isinstance(conteudo, list):
            bc.paragraphs[0].text = ""
            for item in conteudo:
                pi = bc.add_paragraph(); fmt(pi, f"\u25cf  {item}", size=TAM_CORPO, color="1A1A1A")




_NATUREZAS = ("assistencial", "tecnico", "radiotecnico", "gestao", "misto")

# GUARDA: barreira assistencial em acao — nao basta a palavra "paciente", que aparece
# legitimamente num POP tecnico (agenda bloqueada, coorte tratada depois).
# "dois identificadores" ficou de fora de proposito: e tambem como se escreve o
# analogo tecnico de Meta 1 ("dois identificadores: numero de serie e sala").
_MARCA_ASSIST = ("higieniza", "meta 1", "meta 5", "pulseira", "consentimento",
                 "identificacao do paciente", "identificacao segura do paciente",
                 "acompanhante", "beira-leito", "apresentar-se ao paciente",
                 "prontuario do paciente")

_MARCA_TEC = ("loto", "bloqueio e etiquetagem", "disjuntor", "tensao zero", "multimetro",
              "intertravamento", "calibra", "numero de serie", "manutencao",
              "ordem de servico", "engenharia clinica", "dossie do equipamento",
              "energizacao", "equipamento")

_MARCA_RADIO = ("dosimetr", "acelerador", "bunker", "radioprote", "radiolog",
                "cnen", "braquiterapia", "area controlada", "radiometrico")


def _texto_do_documento(secoes):
    partes = [str(secoes.get("objetivo", "")), str(secoes.get("campo_aplicacao", ""))]
    proc = secoes.get("procedimento", {}) or {}
    for chave in ("acoes_iniciais", "execucao_tecnica", "acoes_finais"):
        for item in proc.get(chave, []) or []:
            if isinstance(item, str):
                partes.append(item)
            elif isinstance(item, dict):
                partes.append(str(item.get("texto", "")))
    for r in (secoes.get("riscos", {}) or {}).get("assistenciais", []) or []:
        if isinstance(r, dict):
            partes.append(str(r.get("risco", "")) + " " + str(r.get("barreira", "")))
        else:
            partes.append(str(r))
    return sem_acento(" ".join(partes))


def natureza_processo(meta, secoes):
    """Natureza declarada no JSON; se ausente, deduzida do texto do documento."""
    declarada = sem_acento(str(meta.get("natureza_processo", ""))).strip()
    declarada = {"tecnico-operacional": "tecnico", "tecnico operacional": "tecnico",
                 "radio-tecnico": "radiotecnico", "mista": "misto"}.get(declarada, declarada)
    if declarada in _NATUREZAS:
        return declarada

    txt = _texto_do_documento(secoes)
    assist = sum(1 for m in _MARCA_ASSIST if m in txt)
    tec = sum(1 for m in _MARCA_TEC if m in txt)

    if assist == 0 and tec >= 3:
        radio = any(m in txt for m in _MARCA_RADIO)
        return "radiotecnico" if radio else "tecnico"
    if assist and tec >= 6:
        return "misto"
    return "assistencial"


def rotulo_riscos(meta, secoes, R):
    """Titulo da secao 7.1, conforme a natureza real do processo."""
    return R["risco_" + natureza_processo(meta, secoes)]



def gerar_word(dados, pasta=None, cor=None, logo=None, overrides=None, hoje=None):
    """Gera o .docx a partir do JSON do documento. Devolve o caminho do arquivo.

    dados: dict ou texto JSON; pasta: destino (padrão $POP_SAIDA ou temp);
    cor: cor primária; logo: bytes da imagem; overrides: cores por chave da paleta.
    """
    data = carregar_json(dados)
    if _texto_institucional(data):
        from .texto import gerar_texto
        return gerar_texto(data, pasta=pasta, cor=cor, logo=logo, overrides=overrides, hoje=hoje)
    meta = normalizar_metadata(data.get("metadata", {}), hoje)
    secoes = data.get("secoes", {}) or {}
    R = rotulos(meta["idioma"])
    _, este_doc = tipo_doc(meta["tipo_documento"], meta["idioma"])

    pal = paleta(cor, overrides)
    pal["_rot"] = R
    pal["_tmp"] = tempfile.mkdtemp(prefix="pop_diag_")
    try:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = FONTE; style.font.size = TAM_CORPO
        style.font.color.rgb = RGBColor.from_string("1A1A1A")
        style.paragraph_format.space_after = Pt(4); style.paragraph_format.line_spacing = 1.15

        sec = doc.sections[0]
        sec.page_width = Cm(21); sec.page_height = Cm(29.7)
        sec.top_margin = MARGEM_SUP; sec.bottom_margin = MARGEM_INF
        sec.left_margin = MARGEM_ESQ; sec.right_margin = MARGEM_DIR

        build_header(sec, meta, pal, logo)
        build_footer(sec, meta, pal)

        if doc.paragraphs and doc.paragraphs[0].text == "":
            doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)

        sn = 1
        add_h1(doc, sn, R["s_objetivo"], pal)
        doc.paragraphs[0].paragraph_format.space_before = Pt(0)
        add_body(doc, secoes.get("objetivo") or R["sem_objetivo"])

        sn += 1; add_h1(doc, sn, R["s_campo"], pal)
        add_body(doc, secoes.get("campo_aplicacao") or R["sem_campo"])

        sn += 1; add_h1(doc, sn, R["s_conceitos"], pal)
        for it in secoes.get("conceitos", []) or []:
            if isinstance(it, dict): add_def_item(doc, it.get("termo",""), it.get("definicao",""), pal)
            else: add_bullet(doc, texto_de(it), pal=pal)

        sn += 1; add_h1(doc, sn, R["s_responsabilidades"], pal)
        for it in secoes.get("responsabilidades", []) or []:
            if isinstance(it, dict):
                add_bullet(doc, it.get("acoes",""), bold_prefix=f"{texto_de(it.get('papel',''))}:  ", pal=pal)
            else: add_bullet(doc, texto_de(it), pal=pal)

        sn += 1; add_h1(doc, sn, R["s_recursos"], pal)
        for it in secoes.get("recursos", []) or []:
            add_bullet(doc, texto_de(it), pal=pal)

        sn += 1; add_h1(doc, sn, R["s_procedimento"], pal)
        proc = secoes.get("procedimento", {}) or {}
        sub = 1
        add_h2(doc, f"{sn}.{sub}", R["h_iniciais"], pal)
        for i, p in enumerate(proc.get("acoes_iniciais",[]) or [], 1): _render_passo(doc, i, p, pal)
        sub += 1; add_h2(doc, f"{sn}.{sub}", R["h_execucao"], pal)
        for i, p in enumerate(proc.get("execucao_tecnica",[]) or [], 1): _render_passo(doc, i, p, pal)
        sub += 1; add_h2(doc, f"{sn}.{sub}", R["h_finais"], pal)
        for i, p in enumerate(proc.get("acoes_finais",[]) or [], 1): _render_passo(doc, i, p, pal)

        sn += 1; add_h1(doc, sn, R["s_riscos"], pal)
        riscos = secoes.get("riscos", {}) or {}
        add_h2(doc, f"{sn}.1", rotulo_riscos(meta, secoes, R), pal)
        riscos_list = riscos.get("assistenciais", []) or []
        def _risk_sort_key(it):
            r = sem_acento(it.get("risco","") if isinstance(it, dict) else it)
            if "crit" in r: return 0
            if "moder" in r: return 1
            return 2
        riscos_list = sorted(riscos_list, key=_risk_sort_key)
        for it in riscos_list:
            if isinstance(it, dict): add_risk(doc, it.get("risco",""), it.get("barreira",""), pal)
            else: add_bullet(doc, texto_de(it), pal=pal)
        add_h2(doc, f"{sn}.2", R["h_contingencia"], pal)
        _caixa_contingencia(doc, riscos.get("contingencia") or [], pal)

        sn += 1; add_h1(doc, sn, R["s_registros"], pal)
        add_body(doc, R["registros_intro"])
        for it in secoes.get("registros",[]) or []: add_bullet(doc, texto_de(it), pal=pal)

        sn += 1; add_h1(doc, sn, R["s_indicadores"], pal)
        add_body(doc, R["indicadores_intro"].format(doc=este_doc))
        inds = secoes.get("indicadores", []) or []
        if inds and isinstance(inds[0], dict) and "meta" in inds[0]:
            add_table(doc, R["cab_indicadores"],
                      [[i.get("nome",""),i.get("meta",""),i.get("periodicidade","")] for i in inds],
                      [5157,2400,1800], pal)
        else:
            for it in inds: add_bullet(doc, texto_de(it), pal=pal)

        sn += 1; add_h1(doc, sn, R["s_referencias"], pal)
        for i, ref in enumerate(secoes.get("referencias",[]) or [], 1):
            add_num(doc, i, texto_de(ref), indent=IND_ITEM, pal=pal)

        sn += 1; add_h1(doc, sn, R["s_anexos"], pal)
        anexos = secoes.get("anexos", []) or []
        if anexos:
            for idx, it in enumerate(anexos, 1):
                nome = texto_de(it.get("titulo", "")) if isinstance(it, dict) else str(it)
                add_bullet(doc, f"{R['anexo_item'].format(n=idx)}:  {nome}", pal=pal)
        else:
            add_body(doc, R["sem_anexos"].format(doc=este_doc))

        sn += 1; add_h1(doc, sn, R["s_historico"], pal)
        revs = secoes.get("historico_revisoes") or [{
            "versao": meta["versao"], "data": meta["data_elaboracao"],
            "descricao": R["elaboracao_inicial"], "responsavel": meta["elaborado_por"]["nome"]}]
        add_table(doc, R["cab_historico"],
                  [[r.get("versao",""),r.get("data",""),r.get("descricao",""),r.get("responsavel","")]
                   for r in revs if isinstance(r, dict)],
                  [1000,1200,5557,1600], pal)

        for idx, anx in enumerate(anexos, 1):
            render_annex(doc, anx, idx, pal)

        _ajustar_grades(doc)
        nome = f"{nome_arquivo_seguro(meta['codigo'])}_{nome_arquivo_seguro(meta.get('titulo_processo') or 'POP')}.docx"
        fp = os.path.join(pasta_saida(pasta), nome)
        doc.save(fp)
        return fp
    finally:
        shutil.rmtree(pal["_tmp"], ignore_errors=True)


# GUARDA: ordem dos filhos de w:pPr no esquema do Word (borda e sombreamento antes de espaçamento e recuo)
_ORDEM_PPR = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
              "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
              "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
              "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
              "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
              "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _ppr_inserir(p, el):
    nome = "w:" + el.tag.split("}")[1]
    p._p.get_or_add_pPr().insert_element_before(el, *_ORDEM_PPR[_ORDEM_PPR.index(nome) + 1:])


def _paragrafo_em_caixa(p, pal, fundo="FFF3E0"):
    """Borda e fundo no próprio parágrafo. Parágrafos seguidos com a mesma borda viram uma caixa só
    (no Word e no LibreOffice) e, ao contrário de uma tabela, quebram de página respeitando o
    controle de linhas viúvas e órfãs."""
    lado = pal["gray_border"]
    _ppr_inserir(p, parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="5" w:color="{lado}"/>'
        f'<w:left w:val="single" w:sz="24" w:space="9" w:color="{pal["risk_red"]}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="5" w:color="{lado}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="9" w:color="{lado}"/>'
        f'</w:pBdr>'))
    _ppr_inserir(p, parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fundo}"/>'))
    pf = p.paragraph_format
    pf.left_indent = Cm(0.5); pf.right_indent = Cm(0.5)
    pf.widow_control = True


def _caixa_contingencia(doc, cont_raw, pal):
    """Caixa de destaque com borda lateral vermelha e passos encadeados por →."""
    R = _rot(pal)
    if isinstance(cont_raw, str):
        cont_items = [cont_raw]
    elif isinstance(cont_raw, list):
        cont_items = cont_raw
    else:
        cont_items = [texto_de(cont_raw)]
    pt = doc.add_paragraph()
    _paragrafo_em_caixa(pt, pal)
    pt.paragraph_format.space_before = Pt(6)
    pt.paragraph_format.space_after = Pt(6)
    pt.paragraph_format.keep_with_next = True
    fmt(pt, R["contingencia_titulo"], bold=True, size=TAM_CORPO, color=pal["risk_red"])
    for idx, item_c in enumerate(cont_items, 1):
        pi = doc.add_paragraph()
        _paragrafo_em_caixa(pi, pal)
        pi.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pi.paragraph_format.line_spacing = 1.15
        pi.paragraph_format.space_before = Pt(3)
        pi.paragraph_format.space_after = Pt(3 if idx < len(cont_items) else 8)
        fmt(pi, f"{idx}. ", bold=True, size=TAM_SMALL, color=pal["risk_red"])
        txt = texto_de(item_c).strip()
        txt = re.sub(r'^\d+\.\s*', '', txt)
        txt = re.sub(r'^[\u2192→]\s*', '', txt)
        if ':' in txt:
            conceito, resto = txt.split(':', 1)
            conceito = conceito.replace('**', '').strip()
            fmt(pi, conceito + ":", bold=True, size=TAM_SMALL, color="1A1A1A")
            resto = re.sub(r'\s*[\u2192→]\s*', '  \u2192  ', resto.strip())
            fmt_with_meta_badges(pi, " " + resto, size=TAM_SMALL, pal=pal)
        else:
            fmt_with_meta_badges(pi, txt, size=TAM_SMALL, pal=pal)
    # GUARDA: parágrafo vazio fora da caixa: o que vem depois não "gruda" na borda
    spacing(doc.add_paragraph(), 0, 2)



def _texto_institucional(data):
    """Documento de apoio (motor/texto.py): traz 'corpo' em vez de 'secoes'."""
    return isinstance(data, dict) and "corpo" in data and "secoes" not in data


def validar_word(dados):
    """(ok, mensagem, dados) — exige metadata e o corpo; completa seções ausentes."""
    try:
        data = carregar_json(dados)
    except (ValueError, TypeError) as e:
        return False, f"JSON inválido: {e}", None
    erros = []
    texto = _texto_institucional(data)
    meta = data.get("metadata")
    if not isinstance(meta, dict):
        erros.append("bloco 'metadata' ausente")
    else:
        if not (meta.get("titulo_processo") or (texto and meta.get("titulo"))):
            erros.append("metadata.titulo_processo vazio")
        if not meta.get("codigo"):
            erros.append("metadata.codigo vazio")
    if texto:
        if not isinstance(data.get("corpo"), list) or not data["corpo"]:
            erros.append("'corpo' vazio")
    elif not isinstance(data.get("secoes"), dict):
        erros.append("bloco 'secoes' ausente")
    else:
        for s in ("objetivo", "procedimento"):
            if s not in data["secoes"]:
                erros.append(f"seção '{s}' ausente")
    if erros:
        return False, "; ".join(erros), None
    return True, "ok", data


def gerar_pop_docx(json_str, primary_color=None, logo_bytes=None, palette_overrides=None):
    """Compatibilidade com a assinatura da 1.0 (app.py do Space)."""
    return gerar_word(json_str, cor=primary_color, logo=logo_bytes, overrides=palette_overrides)
