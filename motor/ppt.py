"""Renderizador PowerPoint (.pptx) do POP de Elite — Fábrica de Slides 2.0.

Portado do app.py da 1.0, com os layouts que o prompt exigia e não existiam
(arvore_decisao_clinica, piramide_3d, funil_processos_3d, lista_dupla_circular),
alerta_seguranca em vermelho de verdade (antes caía no slide de citação e saía
vazio), texto sempre coerido para string e saída sem /tmp fixo.
"""

import contextvars
import functools
import hashlib
import io
import json
import math
import os
import re
import shutil
import tempfile
import time
import urllib.parse
import urllib.request

from .base import (DEFAULT_PRIMARY, carregar_json_multiplo, darken, hex_to_rgb, lighten,
                   nome_de_saida, paleta, parse_color_input, pasta_saida, rgb_to_hex,
                   texto_de)
from . import diagramas as _dg
from .i18n import rotulos

FONT_PRIMARY = "Calibri"   # GUARDA: Aptos falta no Office antigo e não tem substituta métrica

_FONTE_REAL = _dg.metrica_calibri()


@functools.lru_cache(maxsize=256)
def _fonte_px(negrito, px):
    return _dg.fonte("negrito" if negrito else "regular", px)


def _linhas_texto(texto, tam, largura_pt, negrito=False):
    """Linhas que o texto ocupa numa caixa de `largura_pt` (quebra por palavra, como o PowerPoint)."""
    fnt = _fonte_px(bool(negrito), max(8, int(round(tam * 4))))
    folga = 0.95 if "**" in texto else 0.98
    return sum(len(_dg.quebrar(par, fnt, largura_pt * 4 * folga)) for par in texto.replace("**", "").split("\n"))


def _largura_palavra_em(texto, negrito=False):
    """Largura da palavra mais longa, em em (múltiplos do tamanho da fonte)."""
    palavras = re.findall(r"\S+", texto.replace("**", ""))
    if not palavras:
        return 0.0
    fnt = _fonte_px(bool(negrito), 400)
    return max(fnt.getlength(w) for w in palavras) / 400


STYLES_AVAILABLE = ["cientifico", "minimalista"]

_WIKI_IMAGE_CACHE = {}
_PRAZO_IMAGENS = contextvars.ContextVar("prazo_imagens", default=None)


def _imagens_da_web_liberadas():
    if os.environ.get("POP_SEM_IMAGENS") == "1":
        return False
    prazo = _PRAZO_IMAGENS.get()
    return prazo is None or time.monotonic() < prazo

_WIKI_CACHE_DIR = os.path.join(tempfile.gettempdir(), "popdeelite_img_cache")
try:
    os.makedirs(_WIKI_CACHE_DIR, exist_ok=True)
except Exception:
    pass


def _query_wikimedia(query, max_results=3, timeout=8):
    """Provedor 1: Wikimedia Commons (CC, sem chave)."""
    if not query or not isinstance(query, str): return []
    cache_key = "wiki:" + hashlib.md5(query.encode("utf-8")).hexdigest()[:16]
    if cache_key in _WIKI_IMAGE_CACHE: return _WIKI_IMAGE_CACHE[cache_key]
    try:
        params = {
            "action": "query", "format": "json", "generator": "search",
            "gsrnamespace": "6", "gsrsearch": f'filetype:bitmap {query}',
            "gsrlimit": str(max_results * 2),
            "prop": "imageinfo", "iiprop": "url|size|extmetadata",
            "iiurlwidth": "1024",
        }
        url = "https://commons.wikimedia.org/w/api.php?" + urllib.parse.urlencode(params)
        req = urllib.request.Request(url, headers={
            "User-Agent": "POPdeElite/1.0 (radioterapia.ai)"
        })
        with urllib.request.urlopen(req, timeout=timeout) as r:
            data = json.loads(r.read().decode("utf-8"))
        pages = data.get("query", {}).get("pages", {})
        results = []
        for _pid, page in pages.items():
            ii = (page.get("imageinfo") or [{}])[0]
            thumb_url = ii.get("thumburl") or ii.get("url")
            mime = ii.get("mime", "")
            if not (mime.startswith("image/jpeg") or mime.startswith("image/png")):
                continue
            ext = ii.get("extmetadata", {}) or {}
            artist = (ext.get("Artist", {}) or {}).get("value", "")
            license_short = (ext.get("LicenseShortName", {}) or {}).get("value", "")
            artist_clean = re.sub(r"<[^>]+>", "", artist or "").strip()
            attribution = f"{artist_clean[:60]}" + (f" ({license_short})" if license_short else "")
            results.append({
                "url": thumb_url,
                "title": page.get("title", "").replace("File:", "").rsplit(".", 1)[0],
                "attribution": attribution.strip()[:120],
                "source": "Wikimedia Commons",
            })
            if len(results) >= max_results: break
        _WIKI_IMAGE_CACHE[cache_key] = results
        return results
    except Exception as e:
        print(f"[wiki] busca falhou para '{query}': {e}")
        _WIKI_IMAGE_CACHE[cache_key] = []
        return []


def _query_openverse(query, max_results=3, timeout=8):
    """Provedor 2: Openverse (CC agregado)."""
    if not query or not isinstance(query, str): return []
    cache_key = "ov:" + hashlib.md5(query.encode("utf-8")).hexdigest()[:16]
    if cache_key in _WIKI_IMAGE_CACHE: return _WIKI_IMAGE_CACHE[cache_key]
    try:
        params = {"q": query, "license": "cc0,by,by-sa",
                  "page_size": str(max_results), "format": "json"}
        url = "https://api.openverse.org/v1/images/?" + urllib.parse.urlencode(params)
        req = urllib.request.Request(url, headers={
            "User-Agent": "POPdeElite/1.0 (radioterapia.ai)",
            "Accept": "application/json"})
        with urllib.request.urlopen(req, timeout=timeout) as r:
            data = json.loads(r.read().decode("utf-8"))
        results = []
        for item in data.get("results", []):
            img_url = item.get("url") or item.get("thumbnail")
            if not img_url: continue
            mime = (item.get("filetype", "") or "").lower()
            if mime and mime not in ("jpg", "jpeg", "png"): continue
            creator = item.get("creator", "") or ""
            license_str = item.get("license", "")
            attribution = f"{creator[:60]} ({license_str})" if creator else f"({license_str})"
            results.append({
                "url": img_url, "title": item.get("title", "")[:80],
                "attribution": attribution.strip()[:120], "source": "Openverse"})
            if len(results) >= max_results: break
        _WIKI_IMAGE_CACHE[cache_key] = results
        return results
    except Exception as e:
        print(f"[openverse] busca falhou para '{query}': {e}")
        _WIKI_IMAGE_CACHE[cache_key] = []
        return []


def _download_image_to_cache(url, timeout=15):
    """Baixa imagem da URL e cacheia em disco. Retorna caminho local ou None."""
    if not url: return None
    cache_key = hashlib.md5(url.encode("utf-8")).hexdigest()[:24]
    ext = ".jpg"
    for cand in (".png", ".jpg", ".jpeg"):
        if url.lower().endswith(cand) or cand in url.lower()[-10:]:
            ext = cand; break
    local_path = os.path.join(_WIKI_CACHE_DIR, cache_key + ext)
    if os.path.exists(local_path) and os.path.getsize(local_path) > 2048:
        return local_path
    try:
        ua = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")
        req = urllib.request.Request(url, headers={"User-Agent": ua})
        with urllib.request.urlopen(req, timeout=timeout) as r:
            blob = r.read()
        if len(blob) < 2048: return None
        try:
            from PIL import Image as _PILImage
            from io import BytesIO as _BIO
            _PILImage.open(_BIO(blob)).verify()
        except Exception:
            return None
        with open(local_path, "wb") as f: f.write(blob)
        return local_path
    except Exception as e:
        print(f"[img-dl] falhou {url[:80]}: {e}")
        return None


def _generate_local_placeholder(query, primary_hex="283264"):
    """Estratégia OFFLINE final: gera placeholder PIL com texto."""
    cache_key = "local:" + hashlib.md5((query or "x").encode("utf-8")).hexdigest()[:16]
    local_path = os.path.join(_WIKI_CACHE_DIR, cache_key + ".png")
    if os.path.exists(local_path) and os.path.getsize(local_path) > 1024:
        return local_path
    try:
        from PIL import Image as PILImage, ImageDraw, ImageFont
        W, H = 1024, 768
        r1, g1, b1 = int(primary_hex[0:2], 16), int(primary_hex[2:4], 16), int(primary_hex[4:6], 16)
        r2, g2, b2 = int(r1 * 0.7), int(g1 * 0.7), int(b1 * 0.7)
        img = PILImage.new("RGB", (W, H), (r1, g1, b1))
        for y in range(H):
            t = y / H
            r = int(r1 * (1 - t) + r2 * t)
            g = int(g1 * (1 - t) + g2 * t)
            b = int(b1 * (1 - t) + b2 * t)
            for x in range(0, W, 2):
                img.putpixel((x, y), (r, g, b))
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("DejaVuSans-Oblique.ttf", 36)
            font_small = ImageFont.truetype("DejaVuSans.ttf", 22)
        except Exception:
            font = ImageFont.load_default(); font_small = font
        title = "Imagem ilustrativa"
        try:
            bbox = draw.textbbox((0, 0), title, font=font_small)
            tw = bbox[2] - bbox[0]
        except Exception:
            tw = 200
        draw.text(((W - tw) // 2, H // 2 - 80), title,
                   fill=(255, 255, 255, 200), font=font_small)
        q_text = f'"{query}"' if query else "(sem descrição)"
        if len(q_text) > 70: q_text = q_text[:67] + "...\""
        try:
            bbox = draw.textbbox((0, 0), q_text, font=font)
            tw = bbox[2] - bbox[0]
        except Exception:
            tw = 400
        draw.text(((W - tw) // 2, H // 2 - 30), q_text,
                   fill=(255, 255, 255), font=font)
        img.save(local_path, "PNG", optimize=True)
        return local_path
    except Exception as e:
        print(f"[local-placeholder] falhou: {e}")
        return None


def _resolve_web_image(queries=None, urls=None, fallback_query=None,
                        primary_hex="283264", force_local_fallback=True):
    """Pipeline cascata: anexo > URLs > queries (Wiki/Openverse) > placeholder.
    Sem foto aleatória: uma foto sem relação com o tema é pior que o placeholder."""
    if not _imagens_da_web_liberadas():
        urls = queries_web = None
    else:
        queries_web = queries
    if urls:
        if isinstance(urls, str): urls = [urls]
        for u in urls:
            if not u: continue
            local = _download_image_to_cache(u)
            if local:
                return {"path": local, "attribution": "Imagem fornecida",
                         "title": "", "source": "URL direta"}
    if isinstance(queries, str): queries = [queries]
    if queries_web:
        for q in queries:
            if not q: continue
            if not _imagens_da_web_liberadas(): break
            for r in _query_wikimedia(q, max_results=3):
                local = _download_image_to_cache(r["url"])
                if local:
                    return {"path": local, "attribution": r["attribution"],
                             "title": r["title"], "source": r["source"]}
            for r in _query_openverse(q, max_results=3):
                local = _download_image_to_cache(r["url"])
                if local:
                    return {"path": local, "attribution": r["attribution"],
                             "title": r["title"], "source": r["source"]}
    if force_local_fallback:
        text = (queries[0] if queries else fallback_query) or ""
        local = _generate_local_placeholder(text, primary_hex=primary_hex)
        if local:
            return {"path": local, "attribution": "",
                     "title": text, "source": "placeholder"}
    return None



def gerar_ppt(dados, pasta=None, cor=None, logo=None, overrides=None,
              anexos=None, estilo=None, idioma=None, imagens=None):
    """Gera a apresentação .pptx a partir do JSON. Devolve o caminho do arquivo.

    `imagens`: ilustrações geradas no site, {número do slide: bytes da imagem}; entram no slide
    no lugar da busca na web, com a legenda "Imagem gerada por IA".
    """
    _tmpdir = tempfile.mkdtemp(prefix="pop_ppt_")
    try:
        segundos = float(os.environ.get("POP_IMAGENS_SEGUNDOS", "45"))
    except ValueError:
        segundos = 45.0
    marca = _PRAZO_IMAGENS.set(time.monotonic() + segundos)
    try:
        return _gerar_ppt(dados, pasta, cor, logo, overrides, anexos, estilo, idioma, _tmpdir, imagens)
    finally:
        _PRAZO_IMAGENS.reset(marca)
        shutil.rmtree(_tmpdir, ignore_errors=True)


def gerar_ppt_pptx(json_str, primary_color=None, logo_bytes=None,
                   palette_overrides=None, template_bytes=None,
                   attachments=None, estilo_visual_override=None):
    """Compatibilidade com a assinatura da 1.0 (app.py do Space)."""
    return gerar_ppt(json_str, cor=primary_color, logo=logo_bytes, overrides=palette_overrides,
                     anexos=attachments, estilo=estilo_visual_override)


def _ilustracoes_como_anexos(slides_data, imagens, pasta_tmp):
    """{número do slide: bytes} → arquivos JPEG na pasta temporária, marcados no slide. Devolve os anexos."""
    from PIL import Image as PILImage
    anexos = {}
    for numero, blob in (imagens or {}).items():
        alvo = next((s for s in slides_data if str(s.get("numero")) == str(numero).strip()), None)
        if alvo is None or not blob:
            continue
        nome = f"ia_{int(alvo.get('numero'))}.jpg"
        caminho = os.path.join(pasta_tmp, nome)
        try:
            with PILImage.open(io.BytesIO(blob)) as im:
                im.load()
                im.convert("RGB").save(caminho, "JPEG", quality=88)
        except Exception:
            continue
        anexos[nome] = caminho
        alvo["imagem_arquivo"] = nome
        alvo["imagem_ia"] = True
    return anexos


def _gerar_ppt(json_str, pasta, primary_color, logo_bytes, palette_overrides,
               attachments, estilo_visual_override, idioma, _tmpdir, imagens=None):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor as PPTRGBColor
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
        from pptx.chart.data import CategoryChartData, XyChartData
        from pptx.oxml.ns import qn
        from lxml import etree
    except ImportError:
        raise RuntimeError("python-pptx/lxml não instalados.")

    data = carregar_json_multiplo(json_str)
    meta = data.get("metadata", data.get("metadata_ppt", {})) or {}
    slides_data = [s for s in (data.get("slides", []) or []) if isinstance(s, dict)]
    if not slides_data:
        raise ValueError("JSON PPT sem 'slides'.")
    R = rotulos(idioma or meta.get("idioma"))
    if imagens:
        attachments = {**(attachments if isinstance(attachments, dict) else {}),
                       **_ilustracoes_como_anexos(slides_data, imagens, _tmpdir)}

    TYPOGRAPHY = {"display": 72, "title": 40, "heading": 28, "subhead": 22,
                   "body": 18, "small": 15, "footer": 13, "caption": 12}
    FALLBACK_CHAIN = ["display", "title", "heading", "subhead",
                      "body", "small", "footer", "caption"]
    def _t(role): return TYPOGRAPHY.get(role, TYPOGRAPHY["body"])
    def _t_fallback(role, levels=1):
        try:
            i = FALLBACK_CHAIN.index(role)
            return FALLBACK_CHAIN[min(i + levels, len(FALLBACK_CHAIN) - 1)]
        except ValueError: return role

    def _estimate_text_height_emu(text, size_pt, width_emu, line_spacing=1.3, fator=0.55):
        if not text: return 0
        char_w_emu = max(1, int(size_pt * 12700 * fator))
        chars_per_line = max(8, width_emu // char_w_emu)
        lines = 0
        for paragraph_text in texto_de(text).split("\n"):
            words = paragraph_text.split()
            cur = 0; line_count = 1
            for w in words:
                wl = len(w) + 1
                if cur + wl > chars_per_line and cur > 0:
                    line_count += 1; cur = wl
                else:
                    cur += wl
            lines += max(1, line_count)
        line_h_emu = int(size_pt * 12700 * line_spacing)
        return lines * line_h_emu

    def _tam_palavra(texto, largura, maximo, minimo=10, fator=0.6):
        """Maior tamanho (pt) em que a palavra mais longa cabe numa linha: evita quebrar
        palavras no meio ("intertravamento" virando "intertravament/o"). Negrito: fator 0.66."""
        palavras = re.findall(r"\S+", texto_de(texto).replace("**", ""))
        if not palavras: return maximo
        if _FONTE_REAL:
            em = _largura_palavra_em(texto_de(texto), negrito=fator >= 0.66)
            return int(max(minimo, min(maximo, largura * 0.95 / max(1, em * 12700))))
        maior = max(len(x) for x in palavras)
        return int(max(minimo, min(maximo, largura / (maior * fator * 12700))))

    def _add_paragrafos(slide, left, top, width, height, paragrafos, vertical_center=False,
                        line_spacing=1.1, espaco_pt=3, alinh=None):
        """Caixa com vários parágrafos; cada um é uma lista de (texto, tam, negrito, itálico, cor)."""
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
        if vertical_center: tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, runs in enumerate(paragrafos):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.line_spacing = line_spacing; par.space_after = Pt(espaco_pt)
            if alinh == "centro":
                par.alignment = PP_ALIGN.CENTER
            for texto, tam, negrito, italico, cor in runs:
                for parte in re.split(r'(\*\*.+?\*\*)', texto_de(texto)):
                    if not parte: continue
                    run = par.add_run()
                    forte = parte.startswith("**") and parte.endswith("**")
                    run.text = parte[2:-2] if forte else parte
                    run.font.size = Pt(tam); run.font.bold = bool(negrito or forte)
                    run.font.italic = italico
                    run.font.color.rgb = _rgb(cor); _font(run)
        return tb

    if estilo_visual_override and estilo_visual_override in STYLES_AVAILABLE:
        estilo_visual = estilo_visual_override
    else:
        estilo_visual = meta.get("estilo_visual", "cientifico")
        if estilo_visual not in STYLES_AVAILABLE:
            estilo_visual = "cientifico"

    pri = parse_color_input(primary_color) if primary_color else DEFAULT_PRIMARY
    pal = paleta(pri)
    pal.setdefault("primary_darker", darken(pri, 0.30))
    pal.setdefault("text_muted", "6B7280")
    pal.setdefault("success_green", "1B5E20")
    pal.setdefault("warning_yellow", "B7791F")
    pal.setdefault("gray_soft", "F3F4F6")
    if palette_overrides:
        for k, v in palette_overrides.items():
            if v: pal[k] = parse_color_input(v)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    use_template_style = False
    sw = prs.slide_width
    sh = prs.slide_height
    total_slides = len(slides_data)
    blank_layout = prs.slide_layouts[6]

    def _rgb(hex_color):
        h = hex_color.lstrip("#")
        return PPTRGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

    def _add_bg_rect(slide, color_hex, x=0, y=0, width=None, height=None, height_frac=None):
        if height_frac is not None: height = int(sh * height_frac)
        if width is None: width = sw
        if height is None: height = sh
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, height)
        shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(color_hex)
        shape.line.fill.background()
        return shape

    def _add_gradient_rect(slide, color_a, color_b, x, y, width, height,
                           direction="horizontal"):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, height)
        shape.line.fill.background()
        sp = shape.fill._xPr
        for tag in ('a:solidFill', 'a:gradFill', 'a:noFill', 'a:pattFill', 'a:blipFill'):
            for el in sp.findall(qn(tag)): sp.remove(el)
        ang_map = {"horizontal": 0, "vertical": 5400000, "diagonal": 2700000}
        ang = ang_map.get(direction, 0)
        ca = color_a.lstrip("#").upper(); cb = color_b.lstrip("#").upper()
        grad_xml = (
            f'<a:gradFill rotWithShape="1" '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:gsLst>'
            f'<a:gs pos="0"><a:srgbClr val="{ca}"/></a:gs>'
            f'<a:gs pos="100000"><a:srgbClr val="{cb}"/></a:gs>'
            f'</a:gsLst><a:lin ang="{ang}" scaled="1"/></a:gradFill>'
        )
        sp.append(etree.fromstring(grad_xml))
        return shape

    def _font(run, name=FONT_PRIMARY):
        run.font.name = name
        try:
            rPr = run._r.get_or_add_rPr()
            for el in rPr.findall(qn('a:latin')): rPr.remove(el)
            latin_xml = (
                f'<a:latin xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                f'typeface="{name}"/>')
            rPr.append(etree.fromstring(latin_xml))
        except Exception: pass

    def _add_textbox_parsed(slide, left, top, width, height, text,
                             size=14, bold=False, color="1A1A1A", align="left",
                             italic=False, line_spacing=1.3, vertical_center=False):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        if vertical_center: tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        if align == "center": p.alignment = PP_ALIGN.CENTER
        elif align == "right": p.alignment = PP_ALIGN.RIGHT
        p.line_spacing = line_spacing
        parts = re.split(r'(\*\*.+?\*\*)', texto_de(text))
        for part in parts:
            if not part: continue
            run = p.add_run()
            if part.startswith("**") and part.endswith("**"):
                run.text = part[2:-2]; run.font.bold = True
            else:
                run.text = part; run.font.bold = bold
            run.font.size = Pt(size); run.font.italic = italic
            run.font.color.rgb = _rgb(color); _font(run)
        return tb

    def _add_textbox(slide, left, top, width, height, text, size=18, bold=False,
                     color="1A1A1A", align="left", italic=False,
                     line_spacing=1.15, vertical_center=False,
                     role=None, auto_shrink=False, min_size=10):
        if role is not None: size = _t(role)
        if auto_shrink and text and height > 0:
            current = size
            while current >= min_size:
                est_h = _estimate_text_height_emu(text, current, width, line_spacing)
                if est_h <= height: break
                current -= 2
            size = current
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
        if vertical_center: tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        if align == "center": p.alignment = PP_ALIGN.CENTER
        elif align == "right": p.alignment = PP_ALIGN.RIGHT
        p.line_spacing = line_spacing
        run = p.add_run(); run.text = texto_de(text)
        run.font.size = Pt(size); run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = _rgb(color); _font(run)
        return tb

    def _add_bullets(slide, left, top, width, height, items, size=16,
                     color="1A1A1A", bullet_char="•", bullet_color=None,
                     line_spacing=1.3, space_after_pt=10):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        bc = bullet_color or color
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = line_spacing
            p.space_after = Pt(space_after_pt)
            text = texto_de(item)
            parts = re.split(r'(\*\*.+?\*\*)', text)
            if bullet_char:
                marker = p.add_run(); marker.text = f"{bullet_char}  "
                marker.font.size = Pt(size); marker.font.color.rgb = _rgb(bc)
                marker.font.bold = True; _font(marker)
            for part in parts:
                if not part: continue
                run = p.add_run()
                if part.startswith("**") and part.endswith("**"):
                    run.text = part[2:-2]; run.font.bold = True
                else:
                    run.text = part
                run.font.size = Pt(size); run.font.color.rgb = _rgb(color)
                _font(run)
        return tb

    def _download_image(url, timeout=15):
        import urllib.request, urllib.error
        ua = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")
        headers = {'User-Agent': ua,
                   'Accept': 'image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8',
                   'Accept-Language': 'en-US,en;q=0.9,pt-BR;q=0.8,pt;q=0.7',
                   'Accept-Encoding': 'identity'}
        def _try_download(u):
            req = urllib.request.Request(u, headers=headers)
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read()
        data = None
        for tentativa in range(2):
            try:
                data = _try_download(url); break
            except Exception:
                continue
        if data is None:
            alt_url = None
            m = re.match(
                r'^(https?://upload\.wikimedia\.org/wikipedia/[^/]+)/thumb/([^/]+/[^/]+/[^/]+)/[^/]+$',
                url)
            if m:
                alt_url = f"{m.group(1)}/{m.group(2)}"
            if alt_url:
                try: data = _try_download(alt_url)
                except Exception: return None
            else: return None
        try:
            from PIL import Image as PILImage
            img = PILImage.open(io.BytesIO(data)); img.verify()
            ext = (img.format or "PNG").lower()
            tmp = tempfile.NamedTemporaryFile(suffix=f".{ext}", delete=False, dir=_tmpdir)
            tmp.write(data); tmp.close()
            return tmp.name
        except Exception: return None

    _attachments = attachments if isinstance(attachments, dict) else {}
    def _find_attachment(nome):
        if not nome or not _attachments: return None
        if nome in _attachments: return _attachments[nome]
        lower_map = {k.lower(): v for k, v in _attachments.items()}
        if nome.lower() in lower_map: return lower_map[nome.lower()]
        base_lower = os.path.basename(nome).lower()
        for k, v in _attachments.items():
            if os.path.basename(k).lower() == base_lower: return v
        return None

    def _extract_pdf_page_as_image(pdf_path, page_number=1):
        try:
            import pypdfium2 as pdfium
        except ImportError:
            return None
        try:
            pdf = pdfium.PdfDocument(pdf_path)
            n_pages = len(pdf)
            idx = max(0, min(page_number - 1, n_pages - 1))
            pil_img = pdf[idx].render(scale=2.0).to_pil()
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False, dir=_tmpdir)
            pil_img.save(tmp.name, "PNG"); tmp.close()
            return tmp.name
        except Exception: return None

    def _resolve_image_source(src, imagem_arquivo=None, imagem_pagina=None):
        if imagem_arquivo:
            att_path = _find_attachment(imagem_arquivo)
            if att_path and os.path.exists(att_path):
                if att_path.lower().endswith(".pdf"):
                    return _extract_pdf_page_as_image(att_path,
                                                     page_number=imagem_pagina or 1)
                return att_path
        if isinstance(src, (bytes, bytearray)):
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False, dir=_tmpdir)
            tmp.write(src); tmp.close(); return tmp.name
        if isinstance(src, str):
            if src.startswith(("http://", "https://")):
                return _download_image(src)
            if os.path.exists(src): return src
        return None

    def _image_dimensions(path):
        try:
            from PIL import Image as PILImage
            with PILImage.open(path) as img: return img.size
        except Exception: return None

    def _add_image_proportional(slide, src, container_left, container_top,
                                 container_width, container_height,
                                 align="center", show_placeholder_on_fail=True,
                                 imagem_arquivo=None, imagem_pagina=None,
                                 placeholder_text=None):
        path = _resolve_image_source(src, imagem_arquivo=imagem_arquivo,
                                     imagem_pagina=imagem_pagina)
        dims = _image_dimensions(path) if path else None
        if not path or not dims:
            if show_placeholder_on_fail:
                ph = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                             container_left, container_top,
                                             container_width, container_height)
                ph.fill.solid(); ph.fill.fore_color.rgb = _rgb(pal["quaternary"])
                ph.line.color.rgb = _rgb(pal.get("gray_border", "7F8C9A"))
                ph.line.width = Pt(0.5)
                msg = placeholder_text or (f"[anexo não encontrado: {imagem_arquivo}]"
                                            if imagem_arquivo else "[imagem indisponível]")
                try:
                    tf = ph.text_frame; tf.text = msg
                    for p in tf.paragraphs:
                        p.alignment = PP_ALIGN.CENTER
                        for r in p.runs:
                            r.font.size = Pt(11); r.font.italic = True
                            r.font.color.rgb = _rgb(pal["text_muted"]); _font(r)
                except Exception: pass
            return False
        try:
            img_w, img_h = dims
            scale = min(container_width / img_w, container_height / img_h)
            new_w = int(img_w * scale); new_h = int(img_h * scale)
            if align == "center":
                ox = (container_width - new_w) // 2
                oy = (container_height - new_h) // 2
            elif align == "top":
                ox = (container_width - new_w) // 2; oy = 0
            elif align == "left":
                ox = 0; oy = (container_height - new_h) // 2
            else:
                ox = oy = 0
            slide.shapes.add_picture(path, container_left + ox, container_top + oy,
                                      width=new_w, height=new_h)
            return True
        except Exception: return False

    ICON_MAP = {
        "heart": MSO_SHAPE.HEART, "star": MSO_SHAPE.STAR_5_POINT,
        "star_4": MSO_SHAPE.STAR_4_POINT, "star_6": MSO_SHAPE.STAR_6_POINT,
        "star_8": MSO_SHAPE.STAR_8_POINT, "sun": MSO_SHAPE.SUN,
        "moon": MSO_SHAPE.MOON, "bolt": MSO_SHAPE.LIGHTNING_BOLT,
        "gear": MSO_SHAPE.GEAR_6, "gear_9": MSO_SHAPE.GEAR_9,
        "cloud": MSO_SHAPE.CLOUD, "arrow_right": MSO_SHAPE.RIGHT_ARROW,
        "arrow_left": MSO_SHAPE.LEFT_ARROW, "arrow_up": MSO_SHAPE.UP_ARROW,
        "arrow_down": MSO_SHAPE.DOWN_ARROW,
        "shield": MSO_SHAPE.PENTAGON, "donut": MSO_SHAPE.DONUT,
        "diamond": MSO_SHAPE.DIAMOND, "hexagon": MSO_SHAPE.HEXAGON,
        "pentagon": MSO_SHAPE.PENTAGON, "octagon": MSO_SHAPE.OCTAGON,
        "explosion": MSO_SHAPE.EXPLOSION1, "smiley": MSO_SHAPE.SMILEY_FACE,
        "plus": MSO_SHAPE.MATH_PLUS, "minus": MSO_SHAPE.MATH_MINUS,
        "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
        "circle": MSO_SHAPE.OVAL, "square": MSO_SHAPE.RECTANGLE,
        "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
        "thermometer": MSO_SHAPE.UP_ARROW_CALLOUT,
        "search": MSO_SHAPE.OVAL, "email": MSO_SHAPE.OVAL,
        "phone": MSO_SHAPE.OVAL, "globe": MSO_SHAPE.OVAL,
        "website": MSO_SHAPE.OVAL,
    }
    CONTACT_GLYPH = {"email": "@", "instagram": "IG", "linkedin": "in",
                      "twitter": "X", "youtube": "▶", "phone": "☎",
                      "globe": "🌐", "website": "🌐", "whatsapp": "✆"}

    def _detect_contact_type(text):
        t = str(text).strip().lower()
        if "@" in t and "." in t and "/" not in t: return "email"
        if "instagram.com" in t or t.startswith("@"): return "instagram"
        if "linkedin.com" in t or "/in/" in t: return "linkedin"
        if "youtube.com" in t or "youtu.be" in t: return "youtube"
        if "whatsapp" in t or "wa.me" in t: return "whatsapp"
        if t.startswith(("http://", "https://", "www.")) or ".com" in t or ".br" in t:
            return "website"
        if any(c.isdigit() for c in t) and len([c for c in t if c.isdigit()]) >= 8:
            return "phone"
        return "generic"

    def _draw_icon(slide, icon_name, left, top, size, fill_color,
                    line_color=None, with_text=None, text_color="FFFFFF"):
        shape_enum = ICON_MAP.get(icon_name, MSO_SHAPE.OVAL)
        shp = slide.shapes.add_shape(shape_enum, left, top, size, size)
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fill_color)
        if line_color:
            shp.line.color.rgb = _rgb(line_color); shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        if with_text:
            tf = shp.text_frame
            tf.margin_left = Emu(0); tf.margin_right = Emu(0)
            tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            run = p.add_run(); run.text = str(with_text)
            run.font.size = Pt(max(10, int(size / 12000)))
            run.font.bold = True; run.font.color.rgb = _rgb(text_color); _font(run)
        return shp

    def _draw_icon_check(slide, left, top, size, color="FFFFFF", bg_color=None):
        if bg_color:
            circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
            circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(bg_color)
            circ.line.fill.background()
            _add_textbox(slide, left, top, size, size, "\u2713",
                         size=int(Emu(size).inches * 18), bold=True, color=color,
                         align="center", vertical_center=True)
        else:
            _add_textbox(slide, left, top, size, size, "\u2713",
                         size=int(Emu(size).inches * 18), bold=True, color=color,
                         align="center", vertical_center=True)

    def _work_area():
        x = Inches(0.5); w = sw - Inches(1.0)
        y_top = Inches(1.7) if estilo_visual != "minimalista" else Inches(1.7)
        y_bottom = sh - Inches(0.85)
        return (x, y_top, w, y_bottom - y_top)

    def _left_margin_for_style(estilo=None): return Inches(0.5)
    def _content_width_for_style(estilo=None): return sw - Inches(1.0)

    def _chrome_cientifico(slide, titulo, numero):
        header_h = int(sh * 0.15)
        _add_bg_rect(slide, pal["primary"], x=0, y=0, width=sw, height=header_h)
        if titulo:
            title_w = sw - Inches(0.5) - Inches(1.8)
            _add_textbox(slide, Inches(0.5), Inches(0.28), title_w, Inches(0.9),
                         titulo, size=26, bold=True, color="FFFFFF")
        if logo_bytes and not use_template_style:
            try:
                from PIL import Image as PILImage
                img = PILImage.open(io.BytesIO(logo_bytes))
                ratio = img.width / img.height
                logo_h_emu = int(Inches(0.9))
                logo_w_emu = int(logo_h_emu * ratio)
                max_w = int(Inches(1.6))
                if logo_w_emu > max_w:
                    logo_w_emu = max_w; logo_h_emu = int(logo_w_emu / ratio)
                logo_y = (header_h - logo_h_emu) // 2
                logo_x = sw - logo_w_emu - Inches(0.4)
                slide.shapes.add_picture(io.BytesIO(logo_bytes), logo_x, logo_y,
                                          width=logo_w_emu, height=logo_h_emu)
            except Exception: pass
        line_y = sh - Inches(0.6)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(0.5), line_y, sw - Inches(1.0), Emu(15000))
        accent.fill.solid(); accent.fill.fore_color.rgb = _rgb(pal["tertiary"])
        accent.line.fill.background()
        return Inches(1.55)

    def _chrome_minimalista(slide, titulo, numero):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                      Inches(0.5), Inches(0.4), Emu(40000), Inches(0.7))
        bar.fill.solid(); bar.fill.fore_color.rgb = _rgb(pal["primary"])
        bar.line.fill.background()
        if titulo:
            _add_textbox(slide, Inches(0.75), Inches(0.4),
                         sw - Inches(1.5), Inches(0.7),
                         titulo, size=24, bold=True,
                         color=pal.get("primary_darker", pal["primary_dark"]),
                         vertical_center=True)
        line_y = sh - Inches(0.55)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                         Inches(0.5), line_y, sw - Inches(1.0), Emu(8000))
        accent.fill.solid(); accent.fill.fore_color.rgb = _rgb(pal["tertiary"])
        accent.line.fill.background()
        return Inches(1.4)

    STYLE_DRAWERS = {"cientifico": _chrome_cientifico, "minimalista": _chrome_minimalista}

    def _draw_chrome(slide, titulo, numero, estilo=None):
        e = estilo or estilo_visual
        return STYLE_DRAWERS.get(e, _chrome_cientifico)(slide, titulo, numero)

    def _add_slide_number(slide, numero, total):
        _add_textbox(slide, sw - Inches(1.2), sh - Inches(0.45), Inches(1.0), Inches(0.3),
                     f"{numero} / {total}", size=12,
                     color=pal["text_muted"], align="right")

    def _author_year_fallback(nome_completo):
        import datetime
        ano = datetime.datetime.now().year
        nome = str(nome_completo).strip()
        for t in ("Dr.", "Dra.", "Prof.", "Profa.", "Dr ", "Dra "):
            if nome.startswith(t): nome = nome[len(t):].strip()
        parts = [p for p in nome.split() if p]
        if not parts: return f"— {ano}"
        sobrenome = parts[-1].upper()
        iniciais = ". ".join(p[0].upper() for p in parts[:-1])
        if iniciais: return f"{sobrenome}, {iniciais}. — {ano}"
        return f"{sobrenome} — {ano}"

    def _add_reference_footer(slide, ref_text, force_fallback_if_empty=True):
        if ref_text:
            t = str(ref_text).strip()
            if t.startswith("[") and t.endswith("]"): t = t[1:-1].strip()
            ref_text = t
        if not ref_text and force_fallback_if_empty:
            pal_name = meta.get("palestrante", "") or ""
            if pal_name.strip():
                ref_text = _author_year_fallback(pal_name)
        if not ref_text: return
        _add_textbox(slide, Inches(0.5), sh - Inches(0.5),
                     sw - Inches(2.0), Inches(0.35),
                     ref_text, size=15, color=pal["text_muted"],
                     italic=True, align="left", vertical_center=True)


    def render_capa(slide, sdata):
        titulo = sdata.get("titulo", "")
        subtitulo = sdata.get("subtitulo", "")
        palestrante = sdata.get("palestrante", meta.get("palestrante", ""))
        credenciais = sdata.get("credenciais", "")
        evento = sdata.get("evento", "")
        data_str = sdata.get("data", "")
        _add_bg_rect(slide, pal["primary"])
        _add_bg_rect(slide, pal["primary_dark"], y=sh - Inches(0.6), height=Inches(0.6))
        logo_fim_y = Inches(0.4)
        if logo_bytes and not use_template_style:
            try:
                from PIL import Image as PILImage
                img = PILImage.open(io.BytesIO(logo_bytes))
                ratio = img.width / img.height
                logo_h_in = 1.3
                logo_w_in = logo_h_in * ratio
                x = (sw - Inches(logo_w_in)) // 2
                slide.shapes.add_picture(io.BytesIO(logo_bytes), x, Inches(0.5),
                                          width=Inches(logo_w_in),
                                          height=Inches(logo_h_in))
                logo_fim_y = Inches(0.5) + Inches(logo_h_in)
            except Exception: pass
        title_color = "FFFFFF"
        def _fit_title_size(text, base_size=40, min_size=26, max_lines=2):
            if not text: return base_size, 0
            size = base_size
            while size >= min_size:
                chars_line = max(20, int(40 * 40 / size))
                words = str(text).split()
                lines, cur = 1, 0
                for w in words:
                    if cur + len(w) + 1 > chars_line:
                        lines += 1; cur = len(w) + 1
                    else:
                        cur += len(w) + 1
                if lines <= max_lines: return size, lines
                size -= 2
            return min_size, max_lines
        def _fit_single_line(text, base_size, min_size):
            if not text: return base_size
            size = base_size
            while size >= min_size:
                chars_line = max(15, int(40 * 40 / size))
                if len(str(text)) <= chars_line: return size
                size -= 2
            return min_size
        y_title = logo_fim_y + Inches(0.45)
        title_size, title_lines = _fit_title_size(titulo, 40, 26, 2)
        title_h = Inches(1.4 if title_lines <= 1 else 1.7)
        _add_textbox(slide, Inches(0.7), y_title, sw - Inches(1.4), title_h,
                      titulo, size=title_size, bold=True, color=title_color,
                      align="center", line_spacing=1.15)
        if subtitulo:
            sub_y = y_title + title_h + Inches(0.25)
            sub_size, _ = _fit_title_size(subtitulo, 32, 22, 2)
            _add_textbox(slide, Inches(0.7), sub_y, sw - Inches(1.4), Inches(0.95),
                          subtitulo, size=sub_size, color=title_color,
                          align="center", italic=True, line_spacing=1.15)
        pal_y = sh - Inches(1.85)
        if palestrante:
            pal_size = _fit_single_line(palestrante, 32, 20)
            _add_textbox(slide, Inches(0.7), pal_y, sw - Inches(1.4), Inches(0.6),
                          palestrante, size=pal_size, bold=True, color=title_color,
                          align="center")
        if credenciais:
            cred_size, _ = _fit_title_size(credenciais, 20, 14, 3)
            _add_textbox(slide, Inches(0.7), pal_y + Inches(0.55),
                          sw - Inches(1.4), Inches(1.0),
                          credenciais, size=cred_size, color=title_color,
                          italic=True, align="center", line_spacing=1.2)
        ev_parts = [p for p in (evento, data_str) if p]
        if ev_parts:
            _add_textbox(slide, Inches(0.5), sh - Inches(0.5),
                          sw - Inches(1), Inches(0.4),
                          " — ".join(ev_parts), size=16, color=title_color,
                          italic=True, align="right", vertical_center=True)

    def render_agenda(slide, sdata):
        titulo = sdata.get("titulo", "Agenda")
        topicos = sdata.get("topicos", [])
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        y = Inches(2.0)
        n_items = max(1, len(topicos))
        row_h = Inches(min(0.75, 4.8 / n_items))
        for i, t in enumerate(topicos[:7]):
            if isinstance(t, dict):
                num = t.get("numero", i + 1)
                tit = t.get("titulo", "")
                tempo = t.get("tempo_estimado", "")
            else:
                num = i + 1; tit = str(t); tempo = ""
            circle = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                             left, y, Inches(0.55), Inches(0.55))
            circle.fill.solid(); circle.fill.fore_color.rgb = _rgb(pal["primary"])
            circle.line.fill.background()
            _add_textbox(slide, left, y + Emu(40000), Inches(0.55), Inches(0.5),
                          str(num), size=16, bold=True, color="FFFFFF",
                          align="center")
            _add_textbox(slide, left + Inches(0.85), y + Emu(30000),
                          width - Inches(2.5), Inches(0.6),
                          tit, size=18, color=pal["text_body"], bold=True)
            if tempo:
                _add_textbox(slide, left + width - Inches(1.6), y + Emu(50000),
                              Inches(1.5), Inches(0.5),
                              f"~ {tempo}", size=12, color=pal["text_muted"],
                              italic=True, align="right")
            y += row_h

    def render_imagem_externa(slide, sdata):
        titulo = sdata.get("titulo", "")
        img_url = sdata.get("imagem_url", "")
        img_arquivo = sdata.get("imagem_arquivo", "")
        img_pagina = sdata.get("imagem_pagina", None)
        legenda = sdata.get("legenda", sdata.get("legenda_imagem", ""))
        credito = sdata.get("credito_imagem", "")
        bullets = sdata.get("bullets", sdata.get("bullet_points", []))
        descricao = sdata.get("descricao", sdata.get("conteudo", ""))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        top = Inches(2.0); avail_h = sh - top - Inches(0.8)
        has_text = bool(bullets) or bool(descricao)
        has_img = bool(img_url) or bool(img_arquivo)
        if has_text and has_img:
            text_w = int(width * 0.44); gutter = Inches(0.3)
            img_x = left + text_w + gutter; img_w = width - text_w - gutter
            if bullets:
                _add_bullets(slide, left, top, text_w, avail_h - Inches(0.3),
                              bullets, size=20, color=pal["text_body"],
                              bullet_color=pal["primary"], bullet_char="▸")
            elif descricao:
                _add_textbox(slide, left, top, text_w, avail_h - Inches(0.3),
                              descricao, size=20, color=pal["text_body"],
                              line_spacing=1.4)
            img_area_h = avail_h - (Inches(0.5) if legenda else 0)
            _add_image_proportional(slide, img_url, img_x, top, img_w, img_area_h,
                                    imagem_arquivo=img_arquivo, imagem_pagina=img_pagina)
            if legenda:
                _add_textbox(slide, img_x, top + img_area_h + Emu(50000),
                              img_w, Inches(0.4),
                              legenda, size=10, color=pal["text_muted"],
                              italic=True, align="center")
        elif has_img:
            img_area_h = avail_h - (Inches(0.5) if legenda else 0)
            _add_image_proportional(slide, img_url, left, top, width, img_area_h,
                                    imagem_arquivo=img_arquivo, imagem_pagina=img_pagina)
            if legenda:
                _add_textbox(slide, left, top + img_area_h + Emu(50000),
                              width, Inches(0.4),
                              legenda, size=11, color=pal["text_muted"],
                              italic=True, align="center")
        elif bullets or descricao:
            if bullets:
                _add_bullets(slide, left, top, width, avail_h,
                              bullets, size=20, color=pal["text_body"],
                              bullet_color=pal["primary"], bullet_char="▸")
            else:
                _add_textbox(slide, left, top, width, avail_h,
                              descricao, size=20, color=pal["text_body"],
                              line_spacing=1.4)
        credits_line = []
        if credito: credits_line.append(f"Imagem: {credito}")
        ref = sdata.get("referencia_rodape", "")
        if ref: credits_line.append(ref)
        _add_reference_footer(slide, " | ".join(credits_line))

    def _resolve_image_for_slide(sdata):
        img_arquivo = sdata.get("imagem_arquivo", "")
        if img_arquivo and _attachments:
            for nome, caminho in _attachments.items():
                if nome == img_arquivo or os.path.basename(caminho) == img_arquivo:
                    if sdata.get("imagem_ia"):
                        return (caminho, R["imagem_ia"], "ia")
                    return (caminho, sdata.get("legenda", ""), "anexo")
        urls = sdata.get("imagem_urls") or sdata.get("urls_candidatas")
        queries = sdata.get("imagem_queries") or sdata.get("queries_candidatas")
        if not queries:
            q_unica = sdata.get("imagem_query", "")
            if q_unica: queries = [q_unica]
        forcar = sdata.get("imagem_obrigatoria", True)
        fallback_q = (queries[0] if queries else None) if forcar else None
        if urls or queries or fallback_q:
            r = _resolve_web_image(queries=queries, urls=urls,
                                    fallback_query=fallback_q,
                                    primary_hex=pal["primary"],
                                    force_local_fallback=True)
            if r:
                source_label = r.get("source", "web").lower()
                attr = "" if source_label == "placeholder" else (r["attribution"] or r["title"] or r.get("source", ""))
                return (r["path"], attr, source_label)
        return (None, None, None)

    def _render_imagem_placeholder(slide, x, y, w, h, query_text):
        ph = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        ph.fill.solid(); ph.fill.fore_color.rgb = _rgb(pal["quaternary"])
        ph.line.color.rgb = _rgb(pal["primary"]); ph.line.width = Pt(1.5)
        try:
            ln = ph.line._get_or_add_ln()
            for d in ln.findall(qn('a:prstDash')): ln.remove(d)
            ln.append(etree.fromstring(
                '<a:prstDash xmlns:a="http://schemas.openxmlformats.org/'
                'drawingml/2006/main" val="dash"/>'))
        except Exception: pass
        _add_textbox(slide, x + Inches(0.3), y + Inches(0.4),
                      w - Inches(0.6), Inches(0.5),
                      R["imagem_sugerida"], size=18, bold=True,
                      color=pal["primary_dark"], align="center")
        _add_textbox(slide, x + Inches(0.3), y + Inches(1.5),
                      w - Inches(0.6), h - Inches(2.0),
                      f"\u201c{query_text}\u201d",
                      size=15, bold=True, color=pal["primary"],
                      align="center", vertical_center=True,
                      auto_shrink=True, min_size=11, line_spacing=1.4)

    def _render_imagem_real(slide, path, x, y, w, h, cobrir=False):
        try:
            from PIL import Image as PILImage
            img = PILImage.open(path)
            iw, ih = img.size
            if cobrir:
                pic = slide.shapes.add_picture(path, x, y, width=w, height=h)
                ar_img, ar_slot = iw / ih, w / h
                if ar_img > ar_slot:
                    sobra = (1 - ar_slot / ar_img) / 2
                    pic.crop_left = pic.crop_right = sobra
                else:
                    sobra = (1 - ar_img / ar_slot) / 2
                    pic.crop_top = pic.crop_bottom = sobra
                return True
            ratio_img = iw / ih; ratio_slot = w / h
            if ratio_img > ratio_slot:
                new_w = w; new_h = int(w / ratio_img)
            else:
                new_h = h; new_w = int(h * ratio_img)
            offset_x = (w - new_w) // 2
            offset_y = (h - new_h) // 2
            slot_area = w * h; img_area = new_w * new_h
            if img_area / slot_area < 0.65:
                zoom = min(1.3, (slot_area / img_area) ** 0.5 * 0.85)
                new_w = min(w, int(new_w * zoom))
                new_h = min(h, int(new_h * zoom))
                offset_x = (w - new_w) // 2
                offset_y = (h - new_h) // 2
            slide.shapes.add_picture(path, x + offset_x, y + offset_y,
                                      width=new_w, height=new_h)
            return True
        except Exception: return False

    def render_imagem_web_direita(slide, sdata):
        titulo = sdata.get("titulo", "")
        bullets = sdata.get("bullets", []); descricao = sdata.get("descricao", "")
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        top = Inches(1.7); avail_h = sh - top - Inches(0.85)
        col_w = (width - Inches(0.4)) // 2
        text_x = left; img_x = left + col_w + Inches(0.4)
        path, attribution, source = _resolve_image_for_slide(sdata)
        img_h_max = int(avail_h * 0.80)
        legenda_y = top + img_h_max + Inches(0.10)
        legenda_h = avail_h - img_h_max - Inches(0.10)
        if path:
            _render_imagem_real(slide, path, img_x, top, col_w, img_h_max)
        else:
            _render_imagem_placeholder(slide, img_x, top, col_w, img_h_max,
                                        sdata.get("imagem_query", "imagem"))
        leg = sdata.get("legenda", "")
        if attribution and source not in ("anexo", None):
            full_caption = (leg + "  ") if leg else ""
            full_caption += attribution
        else: full_caption = leg
        if full_caption:
            _add_textbox(slide, img_x, legenda_y, col_w, legenda_h,
                          full_caption, size=11, italic=True,
                          color=pal["text_muted"], align="center",
                          vertical_center=True, auto_shrink=True,
                          min_size=9, line_spacing=1.2)
        _texto_lateral(slide, text_x, top, col_w, img_h_max, bullets, descricao)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_imagem_web_esquerda(slide, sdata):
        titulo = sdata.get("titulo", "")
        bullets = sdata.get("bullets", []); descricao = sdata.get("descricao", "")
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        top = Inches(1.7); avail_h = sh - top - Inches(0.85)
        col_w = (width - Inches(0.4)) // 2
        img_x = left; text_x = left + col_w + Inches(0.4)
        path, attribution, source = _resolve_image_for_slide(sdata)
        img_h_max = int(avail_h * 0.80)
        legenda_y = top + img_h_max + Inches(0.10)
        legenda_h = avail_h - img_h_max - Inches(0.10)
        if path:
            _render_imagem_real(slide, path, img_x, top, col_w, img_h_max)
        else:
            _render_imagem_placeholder(slide, img_x, top, col_w, img_h_max,
                                        sdata.get("imagem_query", "imagem"))
        leg = sdata.get("legenda", "")
        if attribution and source not in ("anexo", None):
            full_caption = (leg + "  ") if leg else ""
            full_caption += attribution
        else: full_caption = leg
        if full_caption:
            _add_textbox(slide, img_x, legenda_y, col_w, legenda_h, full_caption,
                          size=11, italic=True, color=pal["text_muted"],
                          align="center", vertical_center=True,
                          auto_shrink=True, min_size=9, line_spacing=1.2)
        _texto_lateral(slide, text_x, top, col_w, img_h_max, bullets, descricao)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_imagem_web_panoramica(slide, sdata):
        titulo = sdata.get("titulo", "")
        bullets = sdata.get("bullets", []); descricao = sdata.get("descricao", "")
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        top = Inches(1.7); avail_h = sh - top - Inches(0.85)
        img_h = int(avail_h * 0.60)
        text_y = top + img_h + Inches(0.30)
        text_h = avail_h - img_h - Inches(0.30)
        path, attribution, source = _resolve_image_for_slide(sdata)
        if path:
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, img_h)
            bg.fill.solid()
            bg.fill.fore_color.rgb = _rgb(pal["primary_darker"])
            bg.line.fill.background()
            _render_imagem_real(slide, path, left, top, width, img_h, cobrir=(source == "ia"))
        else:
            _render_imagem_placeholder(slide, left, top, width, img_h,
                                        sdata.get("imagem_query", "imagem"))
        if attribution and source not in ("anexo", None):
            attr_y = top + img_h - Inches(0.32)
            _add_textbox(slide, left, attr_y, width - Inches(0.15), Inches(0.25),
                          attribution, size=9, italic=True, color="FFFFFF",
                          align="right", vertical_center=True,
                          auto_shrink=True, min_size=8)
        if bullets:
            _add_bullets(slide, left, text_y, width, text_h, bullets,
                          size=18, color=pal["text_body"],
                          bullet_color=pal["primary"], bullet_char="\u25b8",
                          line_spacing=1.35, space_after_pt=6)
        elif descricao:
            _add_textbox(slide, left, text_y, width, text_h, descricao,
                          size=18, color=pal["text_body"], line_spacing=1.4,
                          auto_shrink=True, align="center")
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_referencias(slide, sdata):
        titulo = sdata.get("titulo", R["referencias"])
        refs = sdata.get("referencias_completas", sdata.get("referencias", []))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        refs = [texto_de(r) for r in (refs or []) if texto_de(r)][:40]
        if not refs:
            _add_textbox(slide, wx, wy, ww, Inches(0.6), R["sem_referencias"], size=16, italic=True,
                         color=pal["text_muted"])
            return
        if len(refs) <= 8:
            d = Inches(0.46)
            larg = ww - d - Inches(0.3)
            tam = min(_tam_palavras(refs, larg, 22, 12),
                      _maior(lambda t: sum(max(_alt(r, t, larg, 1.2), d) + Inches(0.22) for r in refs), 22, 12, wh))
            alts = [max(_alt(r, tam, larg, 1.2), d) for r in refs]
            total = sum(alts) + Inches(0.22) * (len(refs) - 1)
            y = wy + max(0, (wh - total) // 2)
            for i, (r, a) in enumerate(zip(refs, alts)):
                _bolinha(slide, wx, y + (a - d) // 2, d, i + 1)
                _add_paragrafos(slide, wx + d + Inches(0.3), y, larg, a, [[(r, tam, False, False, pal["text_body"])]],
                                vertical_center=True, line_spacing=1.2)
                if i < len(refs) - 1:
                    _seg(slide, wx + d + Inches(0.3), y + a + Inches(0.11), wx + ww, y + a + Inches(0.11),
                         COR_CARD_ESCURO, larg=0.75)
                y += a + Inches(0.22)
            return
        n = len(refs)
        if n <= 15: font_sz, line_sp, space_after = 16, 1.3, 6
        elif n <= 22: font_sz, line_sp, space_after = 14, 1.25, 5
        else: font_sz, line_sp, space_after = 12, 1.2, 4
        refs_box = slide.shapes.add_textbox(wx, wy, ww, sh - wy - Inches(0.8))
        tf = refs_box.text_frame; tf.word_wrap = True
        tf.margin_left = Emu(0); tf.margin_right = Emu(0)
        for i, ref in enumerate(refs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(space_after); p.line_spacing = line_sp
            run = p.add_run(); run.text = f"{i+1}. {ref}"
            run.font.size = Pt(font_sz)
            run.font.color.rgb = _rgb(pal["text_body"]); _font(run)

    def render_agradecimento(slide, sdata):
        titulo = sdata.get("titulo", R["agradecimento"])
        palestrante = sdata.get("palestrante", meta.get("palestrante", ""))
        contatos = sdata.get("contatos", [])
        _add_bg_rect(slide, pal["primary"])
        if logo_bytes and not use_template_style:
            try:
                from PIL import Image as PILImage
                img = PILImage.open(io.BytesIO(logo_bytes))
                ratio = img.width / img.height
                logo_h_in = 1.2; logo_w_in = logo_h_in * ratio
                x = (sw - Inches(logo_w_in)) // 2
                slide.shapes.add_picture(io.BytesIO(logo_bytes), x, Inches(0.8),
                                          width=Inches(logo_w_in),
                                          height=Inches(logo_h_in))
            except Exception: pass
        larg_tit, alt_tit = sw - Inches(2), Inches(1.6)
        tam_tit = _tam_palavra(titulo, larg_tit, 66, 28, fator=0.66)
        while tam_tit > 28 and _estimate_text_height_emu(titulo, tam_tit, larg_tit, 1.1, fator=0.66) > alt_tit:
            tam_tit -= 2
        _add_textbox(slide, Inches(1), Inches(2.5), larg_tit, alt_tit, titulo, size=tam_tit,
                     bold=True, color="FFFFFF", align="center", vertical_center=True, line_spacing=1.05)
        y_info = Inches(4.4)
        if palestrante:
            _add_textbox(slide, Inches(1), y_info, sw - Inches(2), Inches(0.6),
                          palestrante, role="subhead", bold=True, color="FFFFFF",
                          align="center")
            y_info += Inches(0.75)
        if contatos:
            contatos_validos = [c for c in contatos if c]
            n_ct = len(contatos_validos)
            if n_ct > 0:
                icon_size = Inches(0.45)
                gap_icon_text = Inches(0.12); gap_between = Inches(0.5)
                font_pt = 14; char_w_in = font_pt * 0.0095
                max_text_chars = max(len(str(ct)) for ct in contatos_validos)
                slot_text_w = Inches(min(3.5, max(1.2, max_text_chars * char_w_in)))
                slot_w = icon_size + gap_icon_text + slot_text_w
                total_w = slot_w * n_ct + gap_between * (n_ct - 1)
                max_w = sw - Inches(1.0)
                if total_w > max_w and n_ct > 1:
                    excess = total_w - max_w
                    gap_between = max(Inches(0.15), gap_between - excess // (n_ct - 1))
                    total_w = slot_w * n_ct + gap_between * (n_ct - 1)
                start_x = (sw - total_w) // 2
                for i, ct in enumerate(contatos_validos):
                    ct_type = _detect_contact_type(ct)
                    glyph = CONTACT_GLYPH.get(ct_type, "•")
                    slot_x = start_x + i * (slot_w + gap_between)
                    icon_x = slot_x; icon_y = y_info + Inches(0.05)
                    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                                    icon_x, icon_y,
                                                    icon_size, icon_size)
                    icon.fill.solid()
                    icon.fill.fore_color.rgb = _rgb("FFFFFF")
                    icon.line.fill.background()
                    _add_textbox(slide, icon_x, icon_y, icon_size, icon_size,
                                  glyph, size=18, bold=True,
                                  color=pal["primary_darker"],
                                  align="center", vertical_center=True)
                    text_x = icon_x + icon_size + gap_icon_text
                    _add_textbox(slide, text_x, y_info, slot_text_w, Inches(0.55),
                                  str(ct), size=font_pt, color="FFFFFF", italic=True,
                                  vertical_center=True, align="left",
                                  auto_shrink=True, min_size=10)
        _add_textbox(slide, Inches(0.5), sh - Inches(0.55), sw - Inches(1), Inches(0.4),
                     "Powered by Radioterapia.AI", size=14,
                     color="FFFFFF", italic=True, align="right",
                     vertical_center=True)

    def render_transicao_tema(slide, sdata):
        numero_sec = sdata.get("numero_secao", sdata.get("secao", ""))
        titulo_sec = sdata.get("titulo_secao", sdata.get("titulo", ""))
        subtitulo = (sdata.get("subtitulo_secao", "") or
                     sdata.get("subtitulo", "") or
                     sdata.get("descricao", ""))
        _add_bg_rect(slide, pal["primary"])
        if str(numero_sec).strip():
            _add_textbox(slide, Inches(1), Inches(1.0), sw - Inches(2), Inches(2.5),
                         str(numero_sec), size=180, bold=True, color="FFFFFF",
                         align="center", vertical_center=True)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                       (sw - Inches(1.2)) // 2, Inches(3.7),
                                       Inches(1.2), Emu(50000))
        line.fill.solid(); line.fill.fore_color.rgb = _rgb("FFFFFF")
        line.line.fill.background()
        _add_textbox(slide, Inches(1), Inches(4.0), sw - Inches(2), Inches(1.6),
                     titulo_sec, role="title", bold=True, color="FFFFFF",
                     align="center", auto_shrink=True)
        if subtitulo:
            _add_textbox(slide, Inches(1), Inches(5.6), sw - Inches(2), Inches(1),
                         subtitulo, role="body", color="FFFFFF",
                         italic=True, align="center", line_spacing=1.4,
                         auto_shrink=True)

    def render_tabela_avancada(slide, sdata):
        titulo = sdata.get("titulo", "")
        cabecalho = sdata.get("colunas", sdata.get("cabecalho", sdata.get("header", [])))
        linhas = sdata.get("linhas", sdata.get("rows", []))
        destaques = sdata.get("destaques_celula", sdata.get("destaques_celulas", []))
        comentario = sdata.get("comentario", sdata.get("nota", ""))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        top = Inches(1.7)
        if not cabecalho and not linhas:
            _add_textbox(slide, left, top, width, Inches(4),
                         "[tabela vazia]", size=14, color=pal["text_muted"],
                         italic=True, align="center")
            return
        n_cols = len(cabecalho) if cabecalho else (len(linhas[0]) if linhas else 1)
        n_rows = (1 if cabecalho else 0) + len(linhas)
        reserva_baixo = Inches(0.55)
        if comentario: reserva_baixo += Inches(1.0)
        table_h = sh - top - reserva_baixo
        tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, table_h)
        tbl = tbl_shape.table
        per_col = width // n_cols
        for i in range(n_cols): tbl.columns[i].width = per_col
        font_size = 18
        if cabecalho:
            for ci, val in enumerate(cabecalho):
                cell = tbl.cell(0, ci)
                cell.fill.solid(); cell.fill.fore_color.rgb = _rgb(pal["secondary"])
                cell.margin_left = Emu(40000); cell.margin_right = Emu(40000)
                cell.margin_top = Emu(40000); cell.margin_bottom = Emu(40000)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame; tf.clear()
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                run = p.add_run(); run.text = str(val)
                run.font.size = Pt(font_size); run.font.bold = True
                run.font.color.rgb = _rgb(pal["primary_darker"]); _font(run)
        row_offset = 1 if cabecalho else 0
        for ri, row in enumerate(linhas):
            for ci in range(n_cols):
                val = row[ci] if ci < len(row) else ""
                cell = tbl.cell(ri + row_offset, ci)
                is_zebra = (ri % 2 == 1)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(pal["quaternary"] if is_zebra else "FFFFFF")
                cell.margin_left = Emu(40000); cell.margin_right = Emu(40000)
                cell.margin_top = Emu(30000); cell.margin_bottom = Emu(30000)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame; tf.clear()
                p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                run = p.add_run(); run.text = str(val)
                run.font.size = Pt(font_size)
                run.font.color.rgb = _rgb(pal["text_body"]); _font(run)
        if comentario:
            _add_textbox(slide, left, top + table_h + Emu(100000),
                         width, Inches(0.9),
                         comentario, size=18, color=pal["text_body"],
                         italic=True, line_spacing=1.35)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_jornada_caso_clinico(slide, sdata):
        titulo = sdata.get("titulo", "")
        eventos = sdata.get("eventos", sdata.get("etapas", sdata.get("steps", [])))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        left = _left_margin_for_style(estilo_visual)
        width = _content_width_for_style(estilo_visual)
        top = Inches(1.7); avail_h = sh - top - Inches(0.85)
        n = len(eventos)
        if n == 0:
            _add_textbox(slide, left, top, width, avail_h,
                         R["sem_eventos"], role="body", italic=True,
                         color=pal["text_muted"], align="center", vertical_center=True)
            return
        timeline_x = left + Inches(0.6)
        line_top = top + Inches(0.2); line_bottom = top + avail_h - Inches(0.2)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                       timeline_x - Emu(15000),
                                       line_top, Emu(30000),
                                       line_bottom - line_top)
        line.fill.solid(); line.fill.fore_color.rgb = _rgb(pal["secondary"])
        line.line.fill.background()
        eventos = eventos[:6]
        n = len(eventos)
        row_h = avail_h // max(1, n)
        content_x = timeline_x + Inches(0.6)
        content_w = width - (content_x - left) - Inches(0.2)
        bullet_size = Inches(0.55) if row_h >= Inches(0.9) else Inches(0.42)

        def _partes(i, ev):
            if not isinstance(ev, dict): ev = {"titulo": str(ev)}
            data_ev = texto_de(ev.get("data", ev.get("etapa", "")))
            tit = texto_de(ev.get("titulo", ""))
            desc = texto_de(ev.get("descricao", ev.get("texto", "")))
            subs = ev.get("subitens", ev.get("bullets", []))
            subs = " · ".join(texto_de(x) for x in subs) if isinstance(subs, list) else texto_de(subs)
            return data_ev, tit, desc, subs

        def _altura(partes, tt, td, ts):
            data_ev, tit, desc, subs = partes
            h = _estimate_text_height_emu(f"{data_ev}   {tit}", tt, content_w, 1.1, fator=0.62)
            if desc: h += _estimate_text_height_emu(desc, td, content_w, 1.1) + Pt(3)
            if subs: h += _estimate_text_height_emu("◦ " + subs, ts, content_w, 1.1) + Pt(3)
            return h

        partes = [_partes(i, ev) for i, ev in enumerate(eventos)]
        escala = 1.0
        while escala > 0.55 and any(_altura(p_, 22 * escala, 15 * escala, 12 * escala) > row_h - Inches(0.1)
                                    for p_ in partes):
            escala -= 0.05
        tt, td, ts = max(13, round(22 * escala)), max(10, round(15 * escala)), max(9, round(12 * escala))
        for i, (data_ev, ev_titulo, ev_desc, ev_subs) in enumerate(partes):
            row_y = top + i * row_h
            row_center = row_y + row_h // 2
            bullet_x = timeline_x - bullet_size // 2
            bullet_y = row_center - bullet_size // 2
            bullet = slide.shapes.add_shape(MSO_SHAPE.OVAL,
                                             bullet_x, bullet_y, bullet_size, bullet_size)
            bullet.fill.solid(); bullet.fill.fore_color.rgb = _rgb(pal["primary"])
            bullet.line.color.rgb = _rgb("FFFFFF"); bullet.line.width = Pt(2)
            _add_textbox(slide, bullet_x, bullet_y, bullet_size, bullet_size,
                          str(i + 1), size=16 if bullet_size > Inches(0.5) else 13, bold=True,
                          color="FFFFFF", align="center", vertical_center=True)
            linha_titulo = []
            if data_ev: linha_titulo.append((f"{data_ev}   ", max(9, ts), True, True, pal["primary_dark"]))
            linha_titulo.append((ev_titulo, tt, True, False, pal["primary_darker"]))
            paragrafos = [linha_titulo]
            if ev_desc: paragrafos.append([(ev_desc, td, False, False, pal["text_body"])])
            if ev_subs: paragrafos.append([("◦ " + ev_subs, ts, False, True, pal["text_muted"])])
            _add_paragrafos(slide, content_x, row_y + Inches(0.04), content_w, row_h - Inches(0.08),
                            paragrafos, vertical_center=True)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_citacao_destaque(slide, sdata):
        citacao = sdata.get("citacao", sdata.get("texto", ""))
        autor = sdata.get("autor", "")
        ref_completa = sdata.get("referencia", sdata.get("referencia_completa", ""))
        ref_rodape_extra = sdata.get("referencia_rodape", "")
        _add_bg_rect(slide, pal["quaternary"])
        _add_textbox(slide, Inches(0.3), Inches(0.5), Inches(3), Inches(2.5),
                     "\u201C", size=200, bold=True, color=pal["tertiary"], align="left")
        _add_textbox(slide, sw - Inches(3.3), sh - Inches(3.0), Inches(3), Inches(2.5),
                     "\u201D", size=200, bold=True, color=pal["tertiary"], align="right")
        citacao_txt = str(citacao).strip()
        while citacao_txt and citacao_txt[0] in ("\u201C", "\u201D", "\"", "'"):
            citacao_txt = citacao_txt[1:].lstrip()
        while citacao_txt and citacao_txt[-1] in ("\u201C", "\u201D", "\"", "'"):
            citacao_txt = citacao_txt[:-1].rstrip()
        _add_textbox(slide, Inches(1.2), Inches(2.2),
                     sw - Inches(2.4), Inches(3.3),
                     citacao_txt, size=28, italic=True, color=pal["primary_darker"],
                     align="center", vertical_center=True, line_spacing=1.4)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                       (sw - Inches(1.0)) // 2, Inches(5.75),
                                       Inches(1.0), Emu(25000))
        line.fill.solid(); line.fill.fore_color.rgb = _rgb(pal["primary"])
        line.line.fill.background()
        if autor:
            _add_textbox(slide, Inches(0.5), Inches(5.9),
                         sw - Inches(1), Inches(0.5),
                         f"— {autor}", size=14, bold=True, color=pal["primary_dark"],
                         align="center")
        ref_texto = ref_completa if ref_completa else ref_rodape_extra
        if ref_texto:
            _add_reference_footer(slide, ref_texto)

    def render_infografico_resumo(slide, sdata):
        """Constelação: núcleo no centro e 3 a 6 ideias em cards ligados a ele, com letra que ocupa o card."""
        titulo = sdata.get("titulo", "Resumo da apresentação")
        tema_central = texto_de(sdata.get("tema_central", sdata.get("nucleo", "")))
        subtitulo_central = texto_de(sdata.get("subtitulo_central", ""))
        ideias = sdata.get("ideias", sdata.get("itens", sdata.get("pontos", []))) or []
        ideias = [i if isinstance(i, dict) else {"titulo": texto_de(i)} for i in ideias][:6]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        n = len(ideias)
        if not n:
            return
        sat_w = Inches(3.5) if n <= 4 else Inches(3.2)
        h_max = Inches(1.75) if n <= 4 else Inches(1.45)
        nuc_w, nuc_h = (Inches(3.0), Inches(1.9)) if n <= 5 else (Inches(2.7), Inches(1.6))
        pad = Inches(0.16); d = Inches(0.4)
        inner = sat_w - 2 * pad
        tits = [texto_de(i.get("titulo", i.get("label", ""))) for i in ideias]
        descs = [texto_de(i.get("descricao", i.get("texto", ""))) for i in ideias]

        def altura(tt, td):
            return max(_alt(t, tt, inner, 1.1, True) + (_alt(x, td, inner, 1.15) + Inches(0.06) if x else 0)
                       for t, x in zip(tits, descs))
        tt, td = min(_tam_palavras(tits, inner, 22, 12, negrito=True), 22), min(_tam_palavras(descs, inner, 17, 11), 17)
        while (tt > 12 or td > 11) and altura(tt, td) > h_max - 2 * pad:
            tt = max(12, tt - 1); td = max(11, td - 1)
        sat_h = int(max(altura(tt, td) + 2 * pad, Inches(1.1)))
        esq, dir_ = wx, wx + ww - sat_w
        topo, base = wy, wy + wh - sat_h
        cx = wx + ww // 2
        if n == 1:
            pos = [(cx - sat_w // 2, topo)]
        elif n == 2:
            pos = [(esq, wy + (wh - sat_h) // 2), (dir_, wy + (wh - sat_h) // 2)]
        elif n == 3:
            pos = [(cx - sat_w // 2, topo), (esq, base), (dir_, base)]
        elif n == 4:
            pos = [(esq, topo), (dir_, topo), (dir_, base), (esq, base)]
        elif n == 5:
            pos = [(cx - sat_w // 2, topo), (dir_, topo), (dir_, base), (esq, base), (esq, topo)]
        else:
            pos = [(cx - sat_w // 2, topo), (dir_, topo), (dir_, base), (cx - sat_w // 2, base), (esq, base), (esq, topo)]
        if n in (1, 3, 5):
            cy = topo + sat_h + (wy + wh - topo - sat_h) // 2
        else:
            cy = wy + wh // 2
        nx, ny = cx - nuc_w // 2, cy - nuc_h // 2
        for (x, y) in pos:
            scx, scy = x + sat_w // 2, y + sat_h // 2
            ang = math.atan2(scy - cy, scx - cx)
            _seg(slide, cx + int(nuc_w / 2 * math.cos(ang)), cy + int(nuc_h / 2 * math.sin(ang)), scx, scy,
                 pal["secondary"], larg=1.5)
        nucleo = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(nx), int(ny), int(nuc_w), int(nuc_h))
        nucleo.fill.solid(); nucleo.fill.fore_color.rgb = _rgb(pal["primary"])
        nucleo.line.color.rgb = _rgb(pal["primary_darker"]); nucleo.line.width = Pt(2)
        _sem_sombra(nucleo)
        iw, ih = int(nuc_w * 0.72), int(nuc_h * 0.66)
        tn = min(_tam_palavras([tema_central], iw, 28, 14, negrito=True), 28)
        tsub = 17 if subtitulo_central else 0
        while tn > 14 and _alt(tema_central, tn, iw, 1.05, True) + (_alt(subtitulo_central, tsub, iw, 1.1) if tsub else 0) > ih:
            tn -= 1; tsub = max(12, tsub - 1) if tsub else 0
        pars = [[(tema_central, tn, True, False, "FFFFFF")]]
        if subtitulo_central:
            pars.append([(subtitulo_central, tsub, False, False, lighten(pal["primary"], 0.75))])
        _add_paragrafos(slide, nx + (nuc_w - iw) // 2, ny + (nuc_h - ih) // 2, iw, ih, pars,
                        vertical_center=True, line_spacing=1.05, alinh="centro")
        for k, (x, y) in enumerate(pos):
            _card(slide, x, y, sat_w, sat_h, fundo=COR_CARD, borda=COR_CARD_ESCURO)
            _bolinha(slide, x - d // 3, y - d // 3, d, k + 1)
            pars = [[(tits[k], tt, True, False, pal["primary_darker"])]]
            if descs[k]:
                pars.append([(descs[k], td, False, False, pal["text_body"])])
            _add_paragrafos(slide, x + pad, y + pad, inner, sat_h - 2 * pad, pars, vertical_center=True,
                            line_spacing=1.12, espaco_pt=4, alinh="centro")
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))


    def _mistura(c1, c2, t):
        r1, g1, b1 = hex_to_rgb(c1); r2, g2, b2 = hex_to_rgb(c2)
        return rgb_to_hex(r1 + (r2 - r1) * t, g1 + (g2 - g1) * t, b1 + (b2 - b1) * t)

    def _escuro(c):
        r, g, b = hex_to_rgb(c)
        return (0.299 * r + 0.587 * g + 0.114 * b) < 150

    def _tamanho_que_cabe(texto, largura, altura, maximo, minimo, negrito=False):
        tam = min(maximo, _tam_palavra(texto, largura, maximo, minimo, 0.66 if negrito else 0.6))
        while tam > minimo and _estimate_text_height_emu(texto, tam, largura, 1.15) > altura:
            tam -= 1
        return tam

    def _texto_na_forma(shp, texto, tam_max=18, tam_min=10, cor="FFFFFF", negrito=False,
                        margem_x=None, margem_y=None, ancora=None, area=None):
        """Texto centralizado dentro da forma. `area` = (largura, altura) úteis quando a
        geometria já reserva uma área interna menor (losango, triângulo)."""
        texto = texto_de(texto)
        w, h = shp.width, shp.height
        mx = Inches(0.12) if margem_x is None else margem_x
        my = Inches(0.06) if margem_y is None else margem_y
        tf = shp.text_frame; tf.word_wrap = True
        tf.margin_left = mx; tf.margin_right = mx
        tf.margin_top = my; tf.margin_bottom = my
        tf.vertical_anchor = ancora or MSO_ANCHOR.MIDDLE
        if area is None and shp.auto_shape_type == MSO_SHAPE.OVAL:
            area = (w * 0.7071 - 2 * mx, h * 0.7071 - 2 * my)  # GUARDA: na elipse, o texto só ocupa o retângulo inscrito
        uw, uh = area if area else (w - 2 * mx, h - 2 * my)
        tam = _tamanho_que_cabe(texto, max(1, int(uw)), max(1, int(uh)), tam_max, tam_min, negrito)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.05
        run = p.add_run(); run.text = texto
        run.font.size = Pt(tam); run.font.bold = negrito
        run.font.color.rgb = _rgb(cor); _font(run)
        return shp

    def _caixa(slide, x, y, w, h, texto, fundo, borda, cor_texto,
               forma=None, tam_max=18, tam_min=10, negrito=False):
        forma = forma or MSO_SHAPE.ROUNDED_RECTANGLE
        shp = slide.shapes.add_shape(forma, x, y, w, h)
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fundo)
        shp.line.color.rgb = _rgb(borda); shp.line.width = Pt(1.5)
        if forma == MSO_SHAPE.DIAMOND:
            return _texto_na_forma(shp, texto, tam_max, tam_min, cor_texto, negrito,
                                   margem_x=Inches(0.02), margem_y=Inches(0.02),
                                   area=(w * 0.5 + Inches(0.3), h * 0.5 + Inches(0.15)))
        return _texto_na_forma(shp, texto, tam_max, tam_min, cor_texto, negrito)

    def _seta(slide, x1, y1, x2, y2, cor, rotulo=None):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        c.line.color.rgb = _rgb(cor); c.line.width = Pt(2.25)
        try:
            ln = c.line._get_or_add_ln()
            ln.append(etree.fromstring(
                '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                'type="triangle" w="med" len="med"/>'))
        except Exception: pass
        if rotulo:
            mx = (x1 + x2) // 2; my = (y1 + y2) // 2
            _add_textbox(slide, mx - Inches(0.75), my - Inches(0.42), Inches(1.5), Inches(0.36),
                         rotulo, size=15, bold=True, color=cor, align="center",
                         vertical_center=True)

    def render_alerta_seguranca(slide, sdata):
        """Metas e riscos de dano grave: faixa vermelha, triângulo ⚠ e itens."""
        titulo = texto_de(sdata.get("titulo", "")) or R["alerta_seguranca"]
        itens = (sdata.get("bullet_points") or sdata.get("bullets") or sdata.get("itens")
                 or sdata.get("mensagens") or [])
        if isinstance(itens, str): itens = [itens]
        livre = texto_de(sdata.get("texto") or sdata.get("descricao") or sdata.get("citacao") or "")
        vermelho = pal["risk_red"]
        _add_bg_rect(slide, "FDF2F2")
        faixa = int(sh * 0.15)
        _add_bg_rect(slide, vermelho, x=0, y=0, width=sw, height=faixa)
        _add_bg_rect(slide, "F2C94C", x=0, y=faixa, width=sw, height=Emu(70000))
        _add_textbox(slide, Inches(0.5), Inches(0.28), sw - Inches(1.0), Inches(0.9),
                     titulo, size=26, bold=True, color="FFFFFF", auto_shrink=True, min_size=18)
        area_y = faixa + Emu(70000) + Inches(0.4)
        area_h = sh - Inches(0.85) - area_y
        box_x = Inches(3.5)
        box_w = sw - box_x - Inches(0.55)
        dentro_x = box_x + Inches(0.45); dentro_w = box_w - Inches(0.8)
        pad_v = Inches(0.35)
        itens = [texto_de(i) for i in itens[:9] if texto_de(i)]
        if itens:
            med_w = dentro_w - Inches(0.4)
            tam = min(_tam_palavras(itens, med_w, 30, 16),
                      _maior(lambda t: _alt_lista(itens, t, med_w, 1.25, 14), 30, 16, area_h - 2 * pad_v))
            alt = _alt_lista(itens, tam, med_w, 1.25, 14)
        elif livre:
            tam = min(_tam_palavras([livre], dentro_w, 30, 16),
                      _maior(lambda t: _alt(livre, t, dentro_w, 1.3), 30, 16, area_h - 2 * pad_v))
            alt = _alt(livre, tam, dentro_w, 1.3)
        else:
            tam, alt = 24, 0
        box_h = int(min(area_h, max(alt + 2 * pad_v, area_h * 0.5)))
        box_y = area_y + (area_h - box_h) // 2
        tri_w, tri_h = Inches(2.3), Inches(2.0)
        tri_x, tri_y = Inches(0.65), box_y + (box_h - tri_h) // 2
        tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, tri_x, tri_y, tri_w, tri_h)
        tri.fill.solid(); tri.fill.fore_color.rgb = _rgb(vermelho)
        tri.line.color.rgb = _rgb("FFFFFF"); tri.line.width = Pt(3)
        _add_textbox(slide, tri_x, tri_y + Inches(0.55), tri_w, Inches(1.3), "!",
                     size=80, bold=True, color="FFFFFF", align="center", vertical_center=True)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, box_x, box_y, box_w, box_h)
        box.fill.solid(); box.fill.fore_color.rgb = _rgb("FFFFFF")
        box.line.color.rgb = _rgb("E7B7B7"); box.line.width = Pt(1)
        _sem_sombra(box)
        _add_bg_rect(slide, vermelho, x=box_x, y=box_y, width=Emu(80000), height=box_h)
        if itens:
            _lista_em(slide, dentro_x, box_y + (box_h - alt) // 2, dentro_w, alt + Inches(0.1), itens, tam,
                      cor="3A1010", marcador="▸", ls=1.25, espaco_pt=14, cor_marcador=vermelho)
        elif livre:
            _add_textbox(slide, dentro_x, box_y + pad_v, dentro_w, box_h - 2 * pad_v,
                         livre, size=tam, color="3A1010", line_spacing=1.3, vertical_center=True)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def _poligono(slide, pontos, fundo, borda=None, larg=1.25):
        construtor = slide.shapes.build_freeform(int(pontos[0][0]), int(pontos[0][1]), scale=1.0)
        construtor.add_line_segments([(int(x), int(y)) for x, y in pontos[1:]], close=True)
        shp = construtor.convert_to_shape()
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fundo)
        if borda:
            shp.line.color.rgb = _rgb(borda); shp.line.width = Pt(larg)
        else:
            shp.line.fill.background()
        return _sem_sombra(shp)

    def render_piramide_3d(slide, sdata):
        """Pirâmide de verdade: um só traçado do vértice à base (o topo é um triângulo na mesma
        inclinação dos degraus) e face lateral que converge no vértice. O topo é um pouco mais alto
        que os degraus para o nome caber; se mesmo assim não couber, vai para a coluna da direita."""
        titulo = sdata.get("titulo", "")
        niveis = sdata.get("niveis") or sdata.get("camadas") or sdata.get("itens") or []
        niveis = [n if isinstance(n, dict) else {"titulo": texto_de(n)} for n in niveis][:6]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        n = len(niveis)
        if not n:
            return
        pir_w = int(ww * 0.52)
        cx = wx + Inches(0.05) + pir_w // 2
        y0 = wy + Inches(0.05)
        alt_total = wh - Inches(0.25)
        folga = Inches(0.06)
        u = alt_total / (n - 1 + 1.35)
        bordas = [y0] + [y0 + 1.35 * u + k * u for k in range(n)]
        prof_x, prof_y = Inches(0.34), Inches(0.12)

        def largura(y):
            return pir_w * (y - y0) / alt_total

        def face(y):
            t = (y - y0) / alt_total
            return prof_x * t, prof_y * t

        titulos = [texto_de(nv.get("titulo", "")) for nv in niveis]
        descs = [texto_de(nv.get("descricao") or nv.get("texto") or "") for nv in niveis]
        margem = Inches(0.1)

        def faixa(i):
            return bordas[i], bordas[i + 1] - folga

        tam = 24
        for i in range(1, n):
            yt, yb = faixa(i)
            ty0, ty1 = yt + (yb - yt) * 0.12, yb - (yb - yt) * 0.12
            w_ = largura(ty0) - 2 * margem
            if titulos[i]:
                tam = min(tam, _tam_palavras([titulos[i]], w_, 24, 10, negrito=True),
                          _maior(lambda t, a=titulos[i], w_=w_: _alt(a, t, w_, 1.05, True), 24, 10, ty1 - ty0))
        topo = None
        yt, yb = faixa(0)
        if titulos[0]:
            for t in range(int(tam if n > 1 else 24), 10, -1):
                linhas = 0
                for _ in range(2):
                    alt_b = max(1, linhas) * Pt(t) * 1.22
                    topo_b = yt + (yb - yt) * 0.70 - alt_b / 2
                    w_ = largura(topo_b) - 2 * margem
                    linhas = _linhas_texto(titulos[0], t, w_ / 12700, negrito=True)
                if (linhas <= 2 and _largura_palavra_em(titulos[0], negrito=True) * Pt(t) <= w_ * 0.95
                        and topo_b + alt_b <= yb):
                    topo = (t, topo_b, alt_b, w_)
                    break
        if topo:
            tam = min(tam, int(topo[0] * 1.35))
        for i, nv in enumerate(niveis):
            yt, yb = faixa(i)
            wt, wb = largura(yt), largura(yb)
            cor = _mistura(pal["primary_darker"], pal["secondary"], i / max(1, n - 1))
            (ftx, fty), (fbx, fby) = face(yt), face(yb)
            lado = [(cx + wt / 2, yt), (cx + wt / 2 + ftx, yt - fty), (cx + wb / 2 + fbx, yb - fby), (cx + wb / 2, yb)]
            if i == 0:
                lado = [lado[0], lado[2], lado[3]]
                frente = [(cx, yt), (cx + wb / 2, yb), (cx - wb / 2, yb)]
            else:
                frente = [(cx - wt / 2, yt), (cx + wt / 2, yt), (cx + wb / 2, yb), (cx - wb / 2, yb)]
            _poligono(slide, lado, darken(cor, 0.35))
            _poligono(slide, frente, cor, borda="FFFFFF", larg=1.25)
            cor_txt = "FFFFFF" if _escuro(cor) else pal["primary_darker"]
            if not titulos[i]:
                continue
            if i == 0:
                if topo:
                    t, topo_b, alt_b, w_ = topo
                    _add_paragrafos(slide, int(cx - w_ / 2), int(topo_b), int(w_), int(alt_b + Pt(t) * 0.3),
                                    [[(titulos[0], t, True, False, cor_txt)]], line_spacing=1.0, alinh="centro")
                continue
            ty0, ty1 = yt + (yb - yt) * 0.12, yb - (yb - yt) * 0.12
            w_ = largura(ty0) - 2 * margem
            _add_paragrafos(slide, int(cx - w_ / 2), int(ty0), int(w_), int(ty1 - ty0),
                            [[(titulos[i], tam, True, False, cor_txt)]], vertical_center=True,
                            line_spacing=1.0, alinh="centro")
        if titulos[0] and not topo:
            descs[0] = f"**{titulos[0]}**" + (f" — {descs[0]}" if descs[0] else "")
        desc_x = cx + pir_w // 2 + prof_x + Inches(0.55)
        desc_w = wx + ww - desc_x
        if any(descs):
            alt_d = u - folga
            tam_d = min(_tam_palavras([d for d in descs if d], desc_w, 22, 12),
                        _maior(lambda t: max(_alt(d, t, desc_w, 1.2) for d in descs if d), 22, 12, alt_d))
            for i, d in enumerate(descs):
                if not d:
                    continue
                yt, yb = faixa(i)
                meio = yt + (yb - yt) * (0.62 if i == 0 else 0.5)
                fx, fy = face(meio)
                _seg(slide, cx + largura(meio) / 2 + fx + Inches(0.08), meio - fy, desc_x - Inches(0.15), meio - fy,
                     pal["tertiary"], larg=1.25)
                _add_paragrafos(slide, int(desc_x), int(meio - fy - alt_d / 2), int(desc_w), int(alt_d),
                                [[(d, tam_d, False, False, pal["text_body"])]], vertical_center=True, line_spacing=1.2)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_funil_processos_3d(slide, sdata):
        """Funil do mais largo (topo) ao mais estreito, com valor de cada etapa."""
        titulo = sdata.get("titulo", "")
        etapas = sdata.get("etapas") or sdata.get("itens") or sdata.get("niveis") or []
        etapas = [e if isinstance(e, dict) else {"titulo": texto_de(e)} for e in etapas][:6]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        n = len(etapas)
        if not n:
            return
        fun_w = int(ww * 0.56); fun_x = wx + Inches(0.15)
        borda_topo = Inches(0.22)
        alt = (wh - borda_topo - Inches(0.1)) // n; folga = Inches(0.07)
        base_min = int(fun_w * 0.26)
        prof_x, prof_y = Inches(0.15), Inches(0.08)
        y0 = wy + borda_topo
        cor0 = pal["primary_darker"]
        aro = slide.shapes.add_shape(MSO_SHAPE.OVAL, fun_x, y0 - borda_topo // 2 - Inches(0.06),
                                     fun_w, borda_topo + Inches(0.06))
        aro.fill.solid(); aro.fill.fore_color.rgb = _rgb(lighten(cor0, 0.25))
        aro.line.color.rgb = _rgb("FFFFFF"); aro.line.width = Pt(1)
        val_x = fun_x + fun_w + Inches(0.45); badge = Inches(0.95)
        for i, et in enumerate(etapas):
            w_topo = fun_w - int((fun_w - base_min) * i / n)
            w_base = fun_w - int((fun_w - base_min) * (i + 1) / n)
            x = fun_x + (fun_w - w_topo) // 2
            y = y0 + i * alt; h = alt - folga
            cor = _mistura(cor0, pal["secondary"], i / max(1, n - 1))
            ajuste = ((w_topo - w_base) / 2) / max(1, min(w_topo, h))
            for desl, c in ((prof_x, darken(cor, 0.35)), (0, cor)):
                shp = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, x + desl, y + (prof_y if desl else 0), w_topo, h)
                shp.rotation = 180.0
                shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(c)
                if desl: shp.line.fill.background()
                else:
                    shp.line.color.rgb = _rgb("FFFFFF"); shp.line.width = Pt(1.25)
                try: shp.adjustments[0] = ajuste
                except Exception: pass
            larg_txt = max(Inches(1.0), (w_topo + w_base) // 2 - Inches(0.45))
            _add_textbox(slide, x + (w_topo - larg_txt) // 2, y, larg_txt, h,
                         texto_de(et.get("titulo", "")), size=16, bold=True,
                         color="FFFFFF" if _escuro(cor) else pal["primary_darker"],
                         align="center", vertical_center=True, auto_shrink=True, min_size=10)
            valor = texto_de(et.get("valor") or et.get("numero") or "")
            desc = texto_de(et.get("descricao") or et.get("texto") or "")
            by = y + (h - badge) // 2
            if valor:
                circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, val_x, by, badge, badge)
                circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(cor)
                circ.line.color.rgb = _rgb("FFFFFF"); circ.line.width = Pt(2)
                _texto_na_forma(circ, valor, 20, 10, "FFFFFF" if _escuro(cor) else pal["primary_darker"],
                                True, margem_x=Inches(0.04), margem_y=Inches(0.02))
            if desc:
                tx = val_x + (badge + Inches(0.25) if valor else 0)
                _add_textbox(slide, tx, y, wx + ww - tx, h, desc, size=15,
                             color=pal["text_body"], vertical_center=True,
                             auto_shrink=True, min_size=10, line_spacing=1.2)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    COR_CARD = lighten(pal["primary"], 0.93)
    COR_CARD_ESCURO = lighten(pal["primary"], 0.84)
    GAP = Inches(0.3)
    PAD = Inches(0.28)

    def _sem_sombra(shp):
        try:
            shp.shadow.inherit = False
        except Exception:
            pass
        return shp

    def _card(slide, x, y, w, h, fundo=None, borda=None, raio_in=0.14):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(x), int(y), int(w), int(h))
        try:
            shp.adjustments[0] = max(0.02, min(0.5, Inches(raio_in) / max(1, min(w, h))))
        except Exception:
            pass
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fundo or COR_CARD)
        if borda:
            shp.line.color.rgb = _rgb(borda); shp.line.width = Pt(1.25)
        else:
            shp.line.fill.background()
        return _sem_sombra(shp)

    def _bolinha(slide, x, y, d, texto, fundo=None, cor="FFFFFF"):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(x), int(y), int(d), int(d))
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fundo or pal["primary"])
        shp.line.fill.background()
        _sem_sombra(shp)
        _texto_na_forma(shp, str(texto), max(11, int(Emu(d).inches * 28)), 9, cor, True,
                        margem_x=Emu(0), margem_y=Emu(0))
        return shp

    def _alt(texto, tam, largura, ls=1.15, negrito=False):
        texto = texto_de(texto)
        if not texto:
            return 0
        if _FONTE_REAL:
            return int(_linhas_texto(texto, tam, largura / 12700, negrito) * tam * 1.22 * ls * 12700)
        return _estimate_text_height_emu(texto.replace("**", ""), tam, int(largura), ls,
                                         fator=0.6 if negrito else 0.55)

    def _alt_lista(itens, tam, largura, ls=1.15, espaco_pt=6):
        return sum(_alt(i, tam, largura, ls) + Pt(espaco_pt) for i in itens)

    def _maior(medir, maximo, minimo, limite):
        """Maior tamanho de fonte (pt) em que medir(tam) <= limite."""
        tam = maximo
        while tam > minimo and medir(tam) > limite:
            tam -= 1
        return tam

    def _tam_palavras(textos, largura, maximo, minimo, negrito=False):
        return min([maximo] + [_tam_palavra(t, largura, maximo, minimo, 0.66 if negrito else 0.6)
                               for t in textos if texto_de(t)])

    def _lista_em(slide, x, y, w, h, itens, tam, cor=None, marcador="•", ls=1.15, espaco_pt=6, ancora_meio=False,
                  cor_marcador=None):
        tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = tb.text_frame; tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Emu(0))
        if ancora_meio:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, item in enumerate(itens):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = ls; p.space_after = Pt(espaco_pt)
            if marcador:
                r = p.add_run(); r.text = f"{marcador}  "
                r.font.size = Pt(tam); r.font.bold = True; r.font.color.rgb = _rgb(cor_marcador or pal["primary"]); _font(r)
            for parte in re.split(r'(\*\*.+?\*\*)', texto_de(item)):
                if not parte:
                    continue
                r = p.add_run()
                forte = parte.startswith("**") and parte.endswith("**")
                r.text = parte[2:-2] if forte else parte
                r.font.size = Pt(tam); r.font.bold = forte
                r.font.color.rgb = _rgb(cor or pal["text_body"]); _font(r)
        return tb

    def _seg(slide, x1, y1, x2, y2, cor, larg=1.75, seta=False):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
        c.line.color.rgb = _rgb(cor); c.line.width = Pt(larg)
        if seta:
            try:
                c.line._get_or_add_ln().append(etree.fromstring(
                    '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                    'type="triangle" w="med" len="med"/>'))
            except Exception:
                pass
        return c

    def _polilinha(slide, pontos, cor, larg=1.6, seta=True):
        if len(pontos) < 2:
            return None
        construtor = slide.shapes.build_freeform(int(pontos[0][0]), int(pontos[0][1]), scale=1.0)
        construtor.add_line_segments([(int(x), int(y)) for x, y in pontos[1:]], close=False)
        shp = construtor.convert_to_shape()
        shp.fill.background()
        shp.line.color.rgb = _rgb(cor); shp.line.width = Pt(larg)
        _sem_sombra(shp)
        if seta:
            try:
                shp.line._get_or_add_ln().append(etree.fromstring(
                    '<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                    'type="triangle" w="med" len="med"/>'))
            except Exception:
                pass
        return shp

    def _pilula(slide, cx, cy, texto, cor, tam=12):
        w = Inches(0.25) + Pt(tam) * 0.62 * len(texto)
        h = Pt(tam) * 1.9
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(cx - w / 2), int(cy - h / 2), int(w), int(h))
        try:
            shp.adjustments[0] = 0.5
        except Exception:
            pass
        shp.fill.solid(); shp.fill.fore_color.rgb = _rgb("FFFFFF")
        shp.line.color.rgb = _rgb(cor); shp.line.width = Pt(1.25)
        _sem_sombra(shp)
        _texto_na_forma(shp, texto, tam, 9, cor, True, margem_x=Emu(0), margem_y=Emu(0))
        return shp

    def render_texto_simples(slide, sdata):
        titulo = sdata.get("titulo", "")
        sections = sdata.get("sections", sdata.get("secoes", []))
        bullets = [texto_de(b) for b in (sdata.get("bullets") or sdata.get("bullet_points") or []) if texto_de(b)]
        descricao = texto_de(sdata.get("descricao", sdata.get("conteudo", "")))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        if sections and isinstance(sections, list):
            colunas = [{"titulo": texto_de(s.get("subtitulo", s.get("titulo", ""))),
                        "bullets": s.get("bullets", s.get("bullet_points", []))}
                       for s in sections if isinstance(s, dict)][:3]
            _cards_em_colunas(slide, colunas, wx, wy, ww, wh, numerar=False)
        elif bullets and len(bullets) <= 4 and max(len(b) for b in bullets) <= 160:
            _grade_cards(slide, bullets, wx, wy, ww, wh)
        elif bullets:
            itens = bullets[:14]
            duas = len(itens) > 7
            col_w = (ww - GAP) // 2 if duas else ww - Inches(0.6)
            grupos = [itens[:(len(itens) + 1) // 2], itens[(len(itens) + 1) // 2:]] if duas else [itens]
            tam = min(_tam_palavras(itens, col_w, 26, 14),
                      _maior(lambda t: max(_alt_lista(g, t, col_w - Inches(0.35), espaco_pt=10) for g in grupos), 26, 14, wh - Inches(0.3)))
            alt = max(_alt_lista(g, tam, col_w - Inches(0.35), espaco_pt=10) for g in grupos)
            y0 = wy + max(0, (wh - alt) // 2)
            for k, g in enumerate(grupos):
                x0 = wx + (Inches(0.3) if not duas else k * (col_w + GAP))
                _lista_em(slide, x0, y0, col_w, alt + Inches(0.2), g, tam, espaco_pt=10)
        elif descricao:
            larg = ww - Inches(1.6)
            tam = min(_tam_palavras([descricao], larg, 28, 14),
                      _maior(lambda t: _alt(descricao, t, larg, 1.3), 28, 14, wh - Inches(0.8)))
            alt = _alt(descricao, tam, larg, 1.3) + 2 * PAD
            y0 = wy + max(0, (wh - alt) // 2)
            _card(slide, wx, y0, ww, alt)
            _add_textbox(slide, wx + Inches(0.8), y0 + PAD, larg, alt - 2 * PAD, descricao, size=tam,
                         color=pal["text_body"], line_spacing=1.3, vertical_center=True)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def _texto_lateral(slide, x, y, w, h, bullets, descricao):
        """Tópicos ao lado da imagem: letra ajustada ao espaço e bloco centralizado na altura da imagem."""
        itens = [texto_de(b) for b in (bullets or []) if texto_de(b)]
        if itens:
            larg = w - Inches(0.35)
            tam = min(_tam_palavras(itens, larg, 26, 14),
                      _maior(lambda t: _alt_lista(itens, t, larg, 1.25, 12), 26, 14, h))
            alt = _alt_lista(itens, tam, larg, 1.25, 12)
            _lista_em(slide, x, y + max(0, (h - alt) // 2), w, alt + Inches(0.2), itens, tam, marcador="▸",
                      ls=1.25, espaco_pt=12)
        elif texto_de(descricao):
            texto = texto_de(descricao)
            tam = min(_tam_palavras([texto], w, 26, 14), _maior(lambda t: _alt(texto, t, w, 1.35), 26, 14, h))
            _add_textbox(slide, x, y, w, h, texto, size=tam, color=pal["text_body"], line_spacing=1.35,
                         vertical_center=True)

    def _grade_cards(slide, itens, wx, wy, ww, wh):
        """1-4 frases em cards numerados: 1×n (até 3) ou 2×2."""
        n = len(itens)
        cols = 2 if n in (2, 4) else n
        linhas = 2 if n == 4 else 1
        cw = (ww - GAP * (cols - 1)) // cols
        d = Inches(0.62)
        vertical = cols >= 3
        larg_txt = cw - 2 * PAD - (0 if vertical else d + Inches(0.25))
        alt_max = (wh - GAP * (linhas - 1)) // linhas - 2 * PAD - (d + Inches(0.2) if vertical else 0)
        tam = min(_tam_palavras(itens, larg_txt, 28, 14),
                  _maior(lambda t: max(_alt(i, t, larg_txt, 1.2) for i in itens), 28, 14, alt_max))
        conteudo = max(_alt(i, tam, larg_txt, 1.2) for i in itens)
        ch = int(max(conteudo + 2 * PAD + (d + Inches(0.2) if vertical else 0), d + 2 * PAD,
                     (wh - GAP * (linhas - 1)) / linhas * 0.62))
        total = linhas * ch + (linhas - 1) * GAP
        y0 = wy + max(0, (wh - total) // 2)
        for k, txt in enumerate(itens):
            c, l = k % cols, k // cols
            x = wx + c * (cw + GAP); y = y0 + l * (ch + GAP)
            _card(slide, x, y, cw, ch)
            if vertical:
                _bolinha(slide, x + PAD, y + PAD, d, k + 1)
                _add_paragrafos(slide, x + PAD, y + PAD + d + Inches(0.2), larg_txt, ch - 2 * PAD - d - Inches(0.2),
                                [[(txt, tam, False, False, pal["text_body"])]], line_spacing=1.2)
            else:
                _bolinha(slide, x + PAD, y + (ch - d) // 2, d, k + 1)
                _add_paragrafos(slide, x + PAD + d + Inches(0.25), y + PAD, larg_txt, ch - 2 * PAD,
                                [[(txt, tam, False, False, pal["text_body"])]], vertical_center=True, line_spacing=1.2)

    def _cards_em_colunas(slide, colunas, wx, wy, ww, wh, numerar=True, icones=False):
        """Colunas como cards: [número/ícone] + título + itens; altura pelo conteúdo, grupo centralizado."""
        n = max(1, len(colunas))
        cw = (ww - GAP * (n - 1)) // n
        larg = cw - 2 * PAD
        d = (Inches(0.82) if icones else Inches(0.66)) if (numerar or icones) else 0
        titulos = [texto_de(c.get("titulo", "")) for c in colunas]
        corpos = []
        for c in colunas:
            itens = c.get("bullets") or c.get("itens") or []
            itens = [texto_de(i) for i in (itens if isinstance(itens, list) else [itens]) if texto_de(i)]
            if c.get("texto"):
                itens = [texto_de(c["texto"])] + itens
            corpos.append(itens)
        tam_t = min(_tam_palavras(titulos, larg, 26, 13, negrito=True), 26)
        alt_t = max([_alt(t, tam_t, larg, 1.1, True) for t in titulos if t] + [0])
        topo = d + (Inches(0.18) if d else 0) + (alt_t + Inches(0.15) if alt_t else 0)
        lista = n <= 2 or any(len(c) > 1 for c in corpos)
        larg_c = larg - (Inches(0.3) if lista else 0)
        todos = [i for c in corpos for i in c]
        tam = min(_tam_palavras(todos, larg_c, 24, 13),
                  _maior(lambda t: max([_alt_lista(c, t, larg_c) for c in corpos] + [0]), 24, 13, wh - 2 * PAD - topo))
        alt_c = max([_alt_lista(c, tam, larg_c) for c in corpos] + [0])
        bloco = topo + alt_c
        ch = int(min(wh, max(bloco + 2 * PAD, wh * 0.6)))
        y0 = wy + max(0, (wh - ch) // 2)
        for k, c in enumerate(colunas):
            x = wx + k * (cw + GAP)
            _card(slide, x, y0, cw, ch)
            yy = y0 + max(PAD, (ch - bloco) // 2)
            if icones and c.get("icone") in ICON_MAP:
                circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(x + (cw - d) // 2), int(yy), int(d), int(d))
                circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(pal["primary"]); circ.line.fill.background()
                _sem_sombra(circ)
                _draw_icon(slide, c["icone"], int(x + (cw - d * 0.56) // 2), int(yy + d * 0.22), int(d * 0.56), "FFFFFF")
                yy += d + Inches(0.18)
            elif d:
                _bolinha(slide, x + (cw - d) // 2 if icones else x + PAD, yy, d, k + 1)
                yy += d + Inches(0.18)
            alinh = "centro" if icones else None
            if titulos[k]:
                _add_paragrafos(slide, x + PAD, yy, larg, alt_t + Inches(0.05),
                                [[(titulos[k], tam_t, True, False, pal["primary_darker"])]], line_spacing=1.1, alinh=alinh)
                yy += alt_t + Inches(0.15)
            if corpos[k]:
                if lista:
                    _lista_em(slide, x + PAD, yy, larg, max(Inches(0.3), y0 + ch - PAD - yy), corpos[k], tam, marcador="•")
                else:
                    _add_paragrafos(slide, x + PAD, yy, larg, max(Inches(0.3), y0 + ch - PAD - yy),
                                    [[(corpos[k][0], tam, False, False, pal["text_body"])]], line_spacing=1.2, alinh=alinh)

    def render_texto_duas_colunas(slide, sdata):
        titulo = sdata.get("titulo", "")
        col_e = sdata.get("coluna_esq", {}) or {}
        col_d = sdata.get("coluna_dir", {}) or {}
        if not col_e and "bullet_points_col1" in sdata:
            col_e = {"titulo": "", "bullets": sdata["bullet_points_col1"]}
        if not col_d and "bullet_points_col2" in sdata:
            col_d = {"titulo": "", "bullets": sdata["bullet_points_col2"]}
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        cols = [c if isinstance(c, dict) else {"bullets": [texto_de(c)]} for c in (col_e, col_d)]
        _cards_em_colunas(slide, cols, wx, wy, ww, wh, numerar=False)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_tres_colunas(slide, sdata):
        titulo = sdata.get("titulo", "")
        colunas = list(sdata.get("colunas", []) or [])
        if not colunas:
            for grupo in (("coluna_1", "coluna_esq", "esquerda"), ("coluna_2", "coluna_centro", "centro"),
                          ("coluna_3", "coluna_dir", "direita")):
                for k in grupo:
                    if k in sdata:
                        colunas.append(sdata[k]); break
        colunas = [c if isinstance(c, dict) else {"titulo": "", "bullets": [texto_de(c)]} for c in colunas][:3]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        if colunas:
            _cards_em_colunas(slide, colunas, wx, wy, ww, wh, numerar=True)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def _render_quatro_colunas(slide, sdata):
        titulo = sdata.get("titulo", "")
        colunas = [c if isinstance(c, dict) else {"titulo": "", "texto": texto_de(c)}
                   for c in (sdata.get("colunas", sdata.get("cards", [])) or [])][:4]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        if colunas:
            _cards_em_colunas(slide, colunas, wx, wy, ww, wh, numerar=not any(c.get("icone") for c in colunas),
                              icones=any(c.get("icone") for c in colunas))
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_take_home(slide, sdata):
        titulo = sdata.get("titulo", "Take-Home Messages")
        msgs = [texto_de(m) for m in (sdata.get("mensagens") or sdata.get("bullet_points") or sdata.get("bullets") or [])
                if texto_de(m)][:5]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        if not msgs:
            return
        n = len(msgs)
        d = Inches(0.66)
        larg = ww - 2 * PAD - d - Inches(0.3)
        alt_max = (wh - GAP * (n - 1)) // n - Inches(0.3)
        tam = min(_tam_palavras(msgs, larg, 30, 15),
                  _maior(lambda t: max(_alt(m, t, larg, 1.2) for m in msgs), 30, 15, alt_max))
        rh = int(min(alt_max + Inches(0.3), max(max(_alt(m, tam, larg, 1.2) for m in msgs) + Inches(0.4),
                                                 d + Inches(0.3), (alt_max + Inches(0.3)) * 0.55)))
        total = n * rh + (n - 1) * Inches(0.2)
        y0 = wy + max(0, (wh - total) // 2)
        for i, m in enumerate(msgs):
            y = y0 + i * (rh + Inches(0.2))
            _card(slide, wx, y, ww, rh, fundo=COR_CARD if i % 2 == 0 else "FFFFFF", borda=None if i % 2 == 0 else COR_CARD_ESCURO)
            _bolinha(slide, wx + PAD, y + (rh - d) // 2, d, i + 1)
            _add_paragrafos(slide, wx + PAD + d + Inches(0.3), y, larg, rh,
                            [[(m, tam, i == 0, False, pal["primary_darker"] if i == 0 else pal["text_body"])]],
                            vertical_center=True, line_spacing=1.2)

    def render_lista_dupla_circular(slide, sdata):
        titulo = sdata.get("titulo", "")
        itens = sdata.get("itens") or sdata.get("regras") or sdata.get("bullets") or []
        itens = [it if isinstance(it, dict) else {"titulo": texto_de(it)} for it in itens][:8]
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        n = len(itens)
        if not n:
            return
        cols = 1 if n <= 2 else 2
        por_col = (n + cols - 1) // cols
        cw = (ww - GAP * (cols - 1)) // cols
        d = Inches(0.6)
        larg = cw - 2 * PAD - d - Inches(0.25)
        tits = [texto_de(it.get("titulo", "")) for it in itens]
        txts = [texto_de(it.get("texto") or it.get("descricao") or "") for it in itens]
        alt_max = (wh - GAP * (por_col - 1)) // por_col - Inches(0.3)

        def medir(t):
            return max(_alt(a, t + 2, larg, 1.1, True) + (_alt(b, t, larg, 1.15) if b else 0) for a, b in zip(tits, txts))
        tam = min(_tam_palavras(tits, larg, 26, 14, negrito=True) - 2,  # GUARDA: o título sai em negrito, 2 pt acima
                  _tam_palavras(txts, larg, 24, 12), _maior(medir, 24, 12, alt_max))
        rh = int(min(alt_max + Inches(0.3), max(medir(tam) + Inches(0.4), d + Inches(0.3), (alt_max + Inches(0.3)) * 0.6)))
        total = por_col * rh + (por_col - 1) * GAP
        y0 = wy + max(0, (wh - total) // 2)
        for k in range(n):
            c, l = k // por_col, k % por_col
            x = wx + c * (cw + GAP); y = y0 + l * (rh + GAP)
            _card(slide, x, y, cw, rh)
            _bolinha(slide, x + PAD, y + (rh - d) // 2, d, k + 1)
            pars = [[(tits[k], tam + 2, True, False, pal["primary_darker"])]]
            if txts[k]:
                pars.append([(txts[k], tam, False, False, pal["text_body"])])
            _add_paragrafos(slide, x + PAD + d + Inches(0.25), y, larg, rh, pars, vertical_center=True, line_spacing=1.12)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_disclaimer_creditos(slide, sdata):
        titulo = sdata.get("titulo", "Declaração de Conflitos de Interesse")
        declaracao = texto_de(sdata.get("declaracao",
                              "Declaro que não possuo conflitos de interesse relacionados a esta apresentação."))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        d = Inches(0.9)
        larg = ww - 2 * PAD - d - Inches(0.6)
        tam = min(_tam_palavras([declaracao], larg, 30, 16), _maior(lambda t: _alt(declaracao, t, larg, 1.3), 30, 16, wh - Inches(1)))
        ch = int(max(_alt(declaracao, tam, larg, 1.3) + 2 * PAD + Inches(0.3), d + 2 * PAD))
        y0 = wy + max(0, (wh - ch) // 2)
        _card(slide, wx, y0, ww, ch)
        _bolinha(slide, wx + PAD + Inches(0.1), y0 + (ch - d) // 2, d, "i")
        _add_textbox(slide, wx + PAD + d + Inches(0.6), y0 + PAD, larg, ch - 2 * PAD, declaracao, size=tam,
                     color=pal["primary_darker"], line_spacing=1.3, vertical_center=True)

    def render_estatisticas_duplas(slide, sdata):
        titulo = sdata.get("titulo", "")
        est = [sdata.get("estat_esq", sdata.get("stat_esq", {})), sdata.get("estat_dir", sdata.get("stat_dir", {}))]
        bullets = [texto_de(b) for b in (sdata.get("bullets") or sdata.get("bullet_points") or []) if texto_de(b)]
        conteudo = texto_de(sdata.get("conteudo", sdata.get("descricao", "")))
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        cw = (ww - GAP) // 2
        larg = cw - 2 * PAD
        nums = [texto_de(e.get("numero", e.get("valor", "—")) if isinstance(e, dict) else e) for e in est]
        descs = [texto_de(e.get("descricao", e.get("legenda", "")) if isinstance(e, dict) else "") for e in est]
        tam_n = min(80, min(_tam_palavra(n, larg, 80, 36, 0.62) for n in nums))
        tam_d = min(_tam_palavras(descs, larg, 22, 14), 22)
        alt_d = max([_alt(x, tam_d, larg, 1.2) for x in descs] + [0])
        ch = int(Pt(tam_n) * 1.15 + Inches(0.15) + alt_d + 2 * PAD)
        extra = []
        if bullets:
            tam_b = min(_tam_palavras(bullets, ww - Inches(0.6), 20, 13), 20)
            alt_b = _alt_lista(bullets, tam_b, ww - Inches(0.9), espaco_pt=6)
            extra = [("lista", alt_b + Inches(0.1), tam_b)]
        elif conteudo:
            tam_b = 20
            extra = [("texto", _alt(conteudo, tam_b, ww - Inches(0.6), 1.3) + Inches(0.1), tam_b)]
        total = ch + (GAP + extra[0][1] if extra else 0)
        y0 = wy + max(0, (wh - total) // 2)
        cores = [pal["primary"], pal.get("primary_dark", pal["primary"])]
        for i in range(2):
            x = wx + i * (cw + GAP)
            _card(slide, x, y0, cw, ch, fundo=COR_CARD)
            _add_textbox(slide, x + PAD, y0 + PAD, larg, int(Pt(tam_n) * 1.15), nums[i], size=tam_n, bold=True,
                         color=cores[i], align="center", vertical_center=True, line_spacing=1.0)
            if descs[i]:
                _add_paragrafos(slide, x + PAD, y0 + PAD + int(Pt(tam_n) * 1.15) + Inches(0.15), larg, alt_d + Inches(0.05),
                                [[(descs[i], tam_d, False, False, pal["text_body"])]], line_spacing=1.2, alinh="centro")
        if extra:
            tipo, alt, tam_b = extra[0]
            yb = y0 + ch + GAP
            if tipo == "lista":
                _lista_em(slide, wx + Inches(0.3), yb, ww - Inches(0.6), alt, bullets, tam_b, espaco_pt=6)
            else:
                _add_textbox(slide, wx + Inches(0.3), yb, ww - Inches(0.6), alt, conteudo, size=tam_b,
                             color=pal["text_body"], line_spacing=1.3)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    _jornada_vertical = render_jornada_caso_clinico

    def render_jornada_caso_clinico(slide, sdata):
        eventos = sdata.get("eventos", sdata.get("etapas", sdata.get("steps", []))) or []
        if not (2 <= len(eventos) <= 5):
            return _jornada_vertical(slide, sdata)
        titulo = sdata.get("titulo", "")
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        n = len(eventos)
        cw = (ww - GAP * (n - 1)) // n
        larg = cw - 2 * Inches(0.18)
        d = Inches(0.62)
        evs = []
        for i, ev in enumerate(eventos):
            ev = ev if isinstance(ev, dict) else {"titulo": texto_de(ev)}
            subs = ev.get("subitens", ev.get("bullets", []))
            evs.append((texto_de(ev.get("data", ev.get("etapa", ""))), texto_de(ev.get("titulo", "")),
                        texto_de(ev.get("descricao", ev.get("texto", ""))),
                        [texto_de(s) for s in (subs if isinstance(subs, list) else [subs]) if texto_de(s)]))
        alt_max = wh - d - Inches(0.55)

        def medir(t):
            return max(Pt(t) * 1.3 + _alt(b, t + 4, larg, 1.1, True) + _alt(c, t, larg, 1.2) +
                       (_alt_lista(s, t - 1, larg - Inches(0.2), 1.15, 3) if s else 0) + Inches(0.25)
                       for a, b, c, s in evs)
        tam = min(_tam_palavras([e[1] for e in evs], larg, 25, 15, negrito=True) - 4,  # GUARDA: o título sai em negrito, 4 pt acima
                  _tam_palavras([e[2] for e in evs] + [x for e in evs for x in e[3]], larg, 21, 11),
                  _maior(medir, 21, 11, alt_max))
        bloco = medir(tam)
        total = d + Inches(0.3) + bloco
        y0 = wy + max(0, (wh - total) // 2)
        y_linha = y0 + d // 2
        _seg(slide, wx + cw // 2, y_linha, wx + (n - 1) * (cw + GAP) + cw // 2, y_linha, pal["secondary"], larg=3)
        for i, (data, tit, desc, subs) in enumerate(evs):
            x = wx + i * (cw + GAP)
            _bolinha(slide, x + (cw - d) // 2, y0, d, i + 1)
            yy = y0 + d + Inches(0.3)
            card = _card(slide, x, yy - Inches(0.12), cw, bloco + Inches(0.12))
            pars = []
            if data:
                pars.append([(data.upper(), max(9, tam - 3), True, False, pal["primary"])])
            pars.append([(tit, tam + 4, True, False, pal["primary_darker"])])
            if desc:
                pars.append([(desc, tam, False, False, pal["text_body"])])
            for s in subs:
                pars.append([("• " + s, tam - 1, False, False, pal["text_muted"])])
            _add_paragrafos(slide, x + Inches(0.18), yy, larg, bloco - Inches(0.1), pars, line_spacing=1.12, espaco_pt=4)
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_arvore_decisao_clinica(slide, sdata):
        titulo = sdata.get("titulo", "")
        raiz = sdata.get("nodo_raiz") or sdata.get("raiz") or sdata.get("arvore") or {}
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        if not isinstance(raiz, dict) or not raiz:
            _add_textbox(slide, wx, wy, ww, wh, texto_de(raiz) or "—", size=22, color=pal["text_body"],
                         align="center", vertical_center=True)
            return

        def no(d, sim=None, nivel=0):
            if not isinstance(d, dict):
                d = {"acao": texto_de(d)}
            filhos = []
            if d.get("pergunta") and nivel < 2:
                for chave, s in (("ramo_sim", True), ("ramo_nao", False)):
                    ramo = d.get(chave) or d.get("sim" if s else "nao")
                    if ramo:
                        filhos.append(no(ramo, s, nivel + 1))
            return {"texto": texto_de(d.get("pergunta") or d.get("acao") or d.get("texto") or ""),
                    "pergunta": bool(d.get("pergunta")), "sim": sim, "filhos": filhos,
                    "destaque": str(d.get("destaque", "")).lower() in ("risk_red", "vermelho", "red", "critico", "critica")}
        arvore = no(raiz)

        def folhas(n):
            return 1 if not n["filhos"] else sum(folhas(f) for f in n["filhos"])
        def prof(n):
            return 1 + max([prof(f) for f in n["filhos"]] + [0])
        niveis = prof(arvore)
        L = folhas(arvore)
        gap_x = Inches(0.4)
        col = (ww - gap_x * (L - 1)) // L
        gap_y = Inches(0.8)
        alt_q = Inches(1.15)
        acoes = []

        def coletar(n):
            if not n["filhos"] and not n["pergunta"]:
                acoes.append(n["texto"])
            for f in n["filhos"]:
                coletar(f)
        coletar(arvore)
        prof_folha = niveis - 1
        alt_folha_max = wh - prof_folha * (alt_q + gap_y)
        larg_folha = min(col, Inches(5.0)) - 2 * PAD
        tam_f = min(_tam_palavras(acoes, larg_folha, 26, 12),
                    _maior(lambda t: max([_alt(a, t, larg_folha, 1.2) for a in acoes] + [0]), 26, 12, alt_folha_max - 2 * PAD))
        alt_f = int(min(alt_folha_max, max(max([_alt(a, tam_f, larg_folha, 1.2) for a in acoes] + [0]) + 2 * PAD,
                                           alt_folha_max * 0.7)))
        total = (niveis - 1) * (alt_q + gap_y) + alt_f
        y0 = wy + max(0, (wh - total) // 2)
        posicoes = []

        def colocar(n, x_ini, nivel):
            larg_n = folhas(n) * col + (folhas(n) - 1) * gap_x
            cx = x_ini + larg_n // 2
            y = y0 + nivel * (alt_q + gap_y)
            n["cx"], n["y"] = cx, y
            n["h"] = alt_q if (n["pergunta"] and n["filhos"]) else alt_f
            posicoes.append(n)
            x = x_ini
            for f in n["filhos"]:
                larg_f = folhas(f) * col + (folhas(f) - 1) * gap_x
                colocar(f, x, nivel + 1)
                x += larg_f + gap_x
        colocar(arvore, wx, 0)
        verde, vermelho = pal["barrier_green"], pal["risk_red"]
        for n in posicoes:
            if not n["filhos"]:
                continue
            y_base = n["y"] + n["h"]
            y_barra = y_base + gap_y // 2 - Inches(0.05)
            _seg(slide, n["cx"], y_base, n["cx"], y_barra, pal["primary_dark"])
            xs = [f["cx"] for f in n["filhos"]]
            if len(xs) > 1:
                _seg(slide, min(xs), y_barra, max(xs), y_barra, pal["primary_dark"])
            for f in n["filhos"]:
                cor = verde if f["sim"] else vermelho
                _seg(slide, f["cx"], y_barra, f["cx"], f["y"], pal["primary_dark"], seta=True)
                _pilula(slide, f["cx"], y_barra + (f["y"] - y_barra) // 2 - Inches(0.02),
                        R["sim"] if f["sim"] else R["nao"], cor, tam=14)
        for n in posicoes:
            larg_n = min(col if not n["filhos"] else max(col, Inches(4.6)), ww)
            if n["pergunta"] and n["filhos"]:
                larg_n = min(max(col, Inches(5.0)), Inches(7.5), ww)
                shp = _card(slide, n["cx"] - larg_n // 2, n["y"], larg_n, n["h"], fundo=pal["primary"])
                _texto_na_forma(shp, n["texto"], 26, 12, "FFFFFF", True, margem_x=Inches(0.3))
            else:
                larg_n = min(col, Inches(5.0))
                if n["destaque"] or n["sim"] is False:
                    fundo, borda, cor_t = "FDF0F0", vermelho, "5A1010"
                else:
                    fundo, borda, cor_t = "EEF6EF", verde, "12391A"
                shp = _card(slide, n["cx"] - larg_n // 2, n["y"], larg_n, n["h"], fundo=fundo, borda=borda)
                _texto_na_forma(shp, n["texto"], tam_f, 11, cor_t, False, margem_x=PAD, margem_y=Inches(0.12))
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    def render_fluxograma(slide, sdata):
        """Fluxograma BPMN (estilo Bizagi) com formas nativas, editáveis no PowerPoint."""
        from . import diagramas as dg
        titulo = sdata.get("titulo", "")
        etapas = sdata.get("etapas") or sdata.get("conteudo") or sdata.get("nos") or []
        _draw_chrome(slide, titulo, sdata.get("numero", 0), estilo_visual)
        wx, wy, ww, wh = _work_area()
        rot = {"sim": R["sim"], "nao": R["nao"], "inicio": R.get("inicio", "Início"), "fim": R.get("fim", "Fim")}
        f = dg.montar_fluxo(etapas, sdata.get("raias"), rot, texto_de(sdata.get("processo") or titulo))
        if not f.nos:
            return
        f.med = dict(dg.MED_SLIDE, cab_piscina=0)
        L = dg.ajustar_ao_quadro(f, ww, wh, vao_min=Inches(0.36))
        M = f.med
        H = L.horizontal
        larg_u, alt_u = dg._medidas(L)
        k = min(ww / larg_u, wh / alt_u)
        ox = wx + (ww - larg_u * k) / 2
        oy = wy + (wh - alt_u * k) / 2
        cor = dg.cores_bpmn(pal)

        def XY(m, c):
            return (ox + m * k, oy + c * k) if H else (ox + c * k, oy + m * k)

        def pt(u):
            return max(8.0, math.floor(u * k / 12700 * 2) / 2)

        def medir(tam, negrito=False):
            fnt = dg.fonte("negrito" if negrito else "regular", tam * 4)
            return lambda s: fnt.getlength(s) / 4 * 12700

        def linhas_em(texto, tam, largura):
            fnt = dg.fonte("regular", tam * 4)
            return dg.quebrar(texto, fnt, largura / 12700 * 4)

        def retangulo(x0, y0, x1, y1, fundo, borda=None, larg=1.0):
            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(min(x0, x1)), int(min(y0, y1)),
                                         int(abs(x1 - x0)), int(abs(y1 - y0)))
            if fundo:
                shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fundo)
            else:
                shp.fill.background()
            if borda:
                shp.line.color.rgb = _rgb(borda); shp.line.width = Pt(larg)
            else:
                shp.line.fill.background()
            return _sem_sombra(shp)

        def escrever(tf, linhas, tam, cor_t, negrito=False, alinh=PP_ALIGN.CENTER):
            """Linhas já quebradas aqui (mesma métrica do Calibri): o PowerPoint não corta palavra."""
            p = tf.paragraphs[0]
            p.alignment = alinh
            p.line_spacing = 1.0
            for i, ln in enumerate(linhas):
                if i:
                    p.add_line_break()
                r = p.add_run(); r.text = ln
                r.font.size = Pt(tam); r.font.bold = negrito
                r.font.color.rgb = _rgb(cor_t); _font(r)

        def rotulo(x0, y0, x1, y1, linhas, tam, cor_t, negrito=False, vertical=False, alinh="center", quebra=False):
            folga = 0 if (vertical or alinh == "left") else Inches(0.3)
            tb = slide.shapes.add_textbox(int(min(x0, x1) - folga), int(min(y0, y1)),
                                          int(abs(x1 - x0) + 2 * folga), int(abs(y1 - y0)))
            tf = tb.text_frame
            tf.word_wrap = quebra
            for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                setattr(tf, m, Emu(0))
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            if vertical:
                tf._txBody.bodyPr.set("vert", "vert270")
            escrever(tf, linhas, tam, cor_t, negrito, PP_ALIGN.LEFT if alinh == "left" else PP_ALIGN.CENTER)
            return tb

        mx, my = Inches(0.07), Inches(0.04)
        hm_t, hc_t = dg._forma(dg.No(id="_", tipo="tarefa", texto=""), H, M)
        larg_t = 2 * (hm_t if H else hc_t) * k - 2 * mx
        alt_t = 2 * (hc_t if H else hm_t) * k - 2 * my
        tam_t = pt(M["fonte_tarefa"])
        tarefas = [n for n in f.nos if n.tipo == "tarefa"]
        while tam_t > 9 and any(len(linhas_em(n.texto, tam_t, larg_t * 0.96)) * Pt(tam_t) * 1.22 > alt_t
                                or any(medir(tam_t)(w) > larg_t * 0.96 for w in n.texto.split())
                                for n in tarefas):
            tam_t -= 0.5
        tam_r = min(pt(M["fonte_rotulo"]), tam_t)
        tam_cab = min(pt(M["fonte_cab"]), tam_t)
        tam_s = min(pt(M["fonte_seta"]), tam_t)

        if L.cab_raia:
            for i, (nome, c0, c1) in enumerate(L.raias):
                a = XY(L.cab, c0); b = XY(L.m1, c1)
                retangulo(a[0], a[1], b[0], b[1], cor["raia_fundo"] if i % 2 == 0 else cor["raia_fundo_alt"])
            for nome, c0, c1 in L.raias:
                a = XY(L.cab_piscina, c0); b = XY(L.cab, c1)
                retangulo(a[0], a[1], b[0], b[1], cor["raia_cab"], cor["raia_linha"], 0.75)
                comprimento = abs((b[1] - a[1]) if H else (b[0] - a[0])) - Inches(0.12)
                fnt_cab = dg.fonte("negrito", tam_cab * 4)
                linhas = dg.quebrar(nome, fnt_cab, comprimento / 12700 * 4, max_linhas=2)
                rotulo(a[0], a[1], b[0], b[1], linhas, tam_cab, cor["raia_txt"], True, vertical=H)
            a = XY(L.cab_piscina, L.raias[0][1]); b = XY(L.m1, L.raias[-1][2])
            retangulo(a[0], a[1], b[0], b[1], None, cor["piscina_borda"], 1.25)
            for nome, c0, c1 in L.raias[1:]:
                a = XY(L.cab_piscina, c0); b = XY(L.m1, c0)
                _seg(slide, a[0], a[1], b[0], b[1], cor["raia_linha"], larg=0.75)
        for a in f.arestas:
            _polilinha(slide, [XY(m, c) for m, c in a.pontos], cor["fluxo"], larg=1.5)
        for n in f.nos:
            cx, cy = XY(n.m, n.c)
            hm, hc = dg._forma(n, H, M)
            hw, hh = ((hm, hc) if H else (hc, hm))
            hw, hh = hw * k, hh * k
            if n.tipo == "tarefa":
                shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(cx - hw), int(cy - hh), int(2 * hw), int(2 * hh))
                try:
                    shp.adjustments[0] = 0.12
                except Exception:
                    pass
                shp.fill.gradient(); shp.fill.gradient_angle = 270
                pontos = shp.fill.gradient_stops
                pontos[0].color.rgb = _rgb(cor["tarefa_topo"]); pontos[1].color.rgb = _rgb(cor["tarefa_base"])
                shp.line.color.rgb = _rgb(cor["critico"] if n.critico else cor["tarefa_borda"])
                shp.line.width = Pt(2.25 if n.critico else 1.25)
                tf = shp.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = mx
                tf.margin_top = tf.margin_bottom = my
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                escrever(tf, linhas_em(n.texto, tam_t, larg_t * 0.96), tam_t, cor["tarefa_txt"])
                if n.critico:
                    dd = Inches(0.24)
                    _bolinha(slide, cx + hw - dd * 0.65, cy - hh - dd * 0.35, dd, "!", fundo=cor["critico"])
            elif n.tipo == "decisao":
                shp = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, int(cx - hw), int(cy - hh), int(2 * hw), int(2 * hh))
                shp.fill.gradient(); shp.fill.gradient_angle = 270
                pontos = shp.fill.gradient_stops
                pontos[0].color.rgb = _rgb(cor["gate_topo"]); pontos[1].color.rgb = _rgb(cor["gate_base"])
                shp.line.color.rgb = _rgb(cor["gate_borda"]); shp.line.width = Pt(1.25)
                s = hw * 0.32
                _seg(slide, cx - s, cy - s, cx + s, cy + s, cor["gate_x"], larg=2.25)
                _seg(slide, cx - s, cy + s, cx + s, cy - s, cor["gate_x"], larg=2.25)
            else:
                shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(cx - hw), int(cy - hh), int(2 * hw), int(2 * hh))
                fundo, borda, larg = {"inicio": (cor["inicio_fundo"], cor["inicio_borda"], 1.75),
                                      "fim": (cor["fim_fundo"], cor["fim_borda"], 3.5)}.get(
                    n.tipo, (cor["inter_fundo"], cor["inter_borda"], 1.5))
                shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fundo)
                shp.line.color.rgb = _rgb(borda); shp.line.width = Pt(larg)
            _sem_sombra(shp)
        for nid, (m0, c0, m1, c1, linhas, tam, alinh) in L.rotulos.items():
            a = XY(m0, c0); b = XY(m1, c1)
            rotulo(a[0], a[1], b[0], b[1], linhas, tam_r, cor["rotulo"],
                   alinh="left" if alinh == "esquerda" else "center")
        for a in f.arestas:
            if a.sentido not in ("sim", "nao") or len(a.pontos) < 2:
                continue
            (x0, y0), (x1, y1) = XY(*a.pontos[0]), XY(*a.pontos[1])
            cor_r = cor["sim"] if a.sentido == "sim" else cor["nao"]
            larg_r = medir(tam_s, True)(a.rotulo) + Inches(0.08)
            alt_r = Pt(tam_s) * 1.3
            if abs(y1 - y0) < 1:
                tx = x0 + Inches(0.07) if x1 > x0 else x0 - Inches(0.07) - larg_r
                rotulo(tx, y0 - alt_r - Inches(0.02), tx + larg_r, y0 - Inches(0.02), [a.rotulo], tam_s, cor_r, True, alinh="left")
            else:
                ty = y0 + Inches(0.05) if y1 > y0 else y0 - Inches(0.05) - alt_r
                rotulo(x0 + Inches(0.07), ty, x0 + Inches(0.07) + larg_r, ty + alt_r, [a.rotulo], tam_s, cor_r, True, alinh="left")
        _add_reference_footer(slide, sdata.get("referencia_rodape", ""))

    RENDERERS = {
        "capa": render_capa,
        "disclaimer_creditos": render_disclaimer_creditos,
        "disclaimer": render_disclaimer_creditos,
        "agenda": render_agenda,
        "transicao_tema": render_transicao_tema,
        "secao": render_transicao_tema,
        "transicao": render_transicao_tema,
        "texto_simples": render_texto_simples,
        "texto_duas_colunas": render_texto_duas_colunas,
        "tres_colunas": render_tres_colunas,
        "3_colunas": render_tres_colunas,
        "quatro_colunas": _render_quatro_colunas,
        "4_colunas": _render_quatro_colunas,
        "imagem_externa": render_imagem_externa,
        "imagem_ilustrativa": render_imagem_externa,
        "texto_imagem": render_imagem_externa,
        "imagem_web_direita": render_imagem_web_direita,
        "imagem_direita": render_imagem_web_direita,
        "imagem_web_esquerda": render_imagem_web_esquerda,
        "imagem_esquerda": render_imagem_web_esquerda,
        "imagem_web_panoramica": render_imagem_web_panoramica,
        "imagem_panoramica": render_imagem_web_panoramica,
        "imagem_central": render_imagem_web_panoramica,
        "jornada_caso_clinico": render_jornada_caso_clinico,
        "jornada": render_jornada_caso_clinico,
        "timeline_vertical": render_jornada_caso_clinico,
        "tabela_avancada": render_tabela_avancada,
        "tabela": render_tabela_avancada,
        "estatisticas_duplas": render_estatisticas_duplas,
        "big_numbers": render_estatisticas_duplas,
        "infografico_resumo": render_infografico_resumo,
        "infografico": render_infografico_resumo,
        "constelacao_resumo": render_infografico_resumo,
        "citacao_destaque": render_citacao_destaque,
        "citacao": render_citacao_destaque,
        "alerta_seguranca": render_alerta_seguranca,
        "alerta": render_alerta_seguranca,
        "arvore_decisao_clinica": render_arvore_decisao_clinica,
        "arvore_decisao": render_arvore_decisao_clinica,
        "piramide_3d": render_piramide_3d,
        "piramide": render_piramide_3d,
        "funil_processos_3d": render_funil_processos_3d,
        "funil": render_funil_processos_3d,
        "lista_dupla_circular": render_lista_dupla_circular,
        "fluxograma": render_fluxograma,
        "fluxograma_bpmn": render_fluxograma,
        "fluxo_processo": render_fluxograma,
        "processo_bpmn": render_fluxograma,
        "bpmn": render_fluxograma,
        "matriz_tabela": render_tabela_avancada,
        "resumo_one_page": render_take_home,
        "take_home": render_take_home,
        "take_home_messages": render_take_home,
        "resumo": render_take_home,
        "referencias": render_referencias,
        "referencias_completas": render_referencias,
        "agradecimento": render_agradecimento,
        "agradecimentos": render_agradecimento,
        "obrigado": render_agradecimento,
    }

    for sdata in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        tipo = sdata.get("tipo", "texto_simples")
        numero = sdata.get("numero", slides_data.index(sdata) + 1)
        sdata["numero"] = numero
        renderer = RENDERERS.get(tipo, render_texto_simples)
        try:
            renderer(slide, sdata)
        except Exception as e:
            import traceback; traceback.print_exc()
            _draw_chrome(slide, sdata.get("titulo", f"Slide {numero}"),
                          numero, estilo_visual)
            _add_textbox(slide, Inches(0.5), Inches(2.5), sw - Inches(1), Inches(3),
                          f"[erro ao renderizar tipo '{tipo}': {e}]",
                          size=14, color=pal["risk_red"], italic=True, align="center")
        if tipo not in ("capa", "agradecimentos", "agradecimento", "obrigado",
                         "transicao_tema", "transicao", "secao",
                         "citacao_destaque", "citacao"):
            _add_slide_number(slide, numero, total_slides)

    nome = nome_de_saida(meta.get("codigo"), meta.get("codigo_arquivo") or meta.get("apresentacao"),
                         "pptx", "Apresentacao")
    fp = os.path.join(pasta_saida(pasta), nome)
    prs.save(fp)
    return fp
