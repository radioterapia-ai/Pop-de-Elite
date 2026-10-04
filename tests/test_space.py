"""Testes da API do motor (Space do Hugging Face) — rode com:  python -m pytest"""

import base64
import importlib.util
import io
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

RAIZ = Path(__file__).resolve().parent.parent
EXEMPLOS = RAIZ / "exemplos"


def imagem_b64(formato):
    saida = io.BytesIO()
    Image.new("RGB", (120, 60), (15, 118, 110)).save(saida, formato)
    return base64.b64encode(saida.getvalue()).decode()


def carregar(nome):
    return json.loads((EXEMPLOS / nome).read_text(encoding="utf-8"))


def app_do_space(monkeypatch, token=""):
    monkeypatch.setenv("POP_MOTOR_TOKEN", token)
    monkeypatch.setenv("POP_SEM_IMAGENS", "1")
    spec = importlib.util.spec_from_file_location("space_app", RAIZ / "space" / "app.py")
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return TestClient(modulo.app)


@pytest.fixture
def cliente(monkeypatch):
    c = app_do_space(monkeypatch, token="teste")
    c.headers["X-Pop-Token"] = "teste"
    return c


def test_saude(cliente):
    r = cliente.get("/saude")
    assert r.status_code == 200 and r.json()["ok"] is True
    assert r.json()["diagramas"] == "bpmn"


def test_renderiza_os_tres_formatos(cliente):
    r = cliente.post("/renderizar", json={
        "word": carregar("pop_linac_manutencao.json"),
        "excel": carregar("planilha_linac_manutencao.json"),
        "ppt": carregar("ppt_layouts_novos.json"),
        "cor": "#0F766E", "logo_b64": "data:image/webp;base64," + imagem_b64("WEBP"),
        "estilo": "minimalista"})
    assert r.status_code == 200, r.text
    corpo = r.json()
    assert corpo["erros"] == {}
    formatos = {a["formato"]: a for a in corpo["arquivos"]}
    assert set(formatos) == {"word", "excel", "ppt"}
    for a in formatos.values():
        dados = base64.b64decode(a["base64"])
        assert dados[:2] == b"PK" and len(dados) == a["tamanho"]
        assert a["nome"].isascii()
    doc = Document(io.BytesIO(base64.b64decode(formatos["word"]["base64"])))
    assert doc.sections[0].header.tables


def test_um_formato_com_problema_nao_derruba_os_outros(cliente):
    r = cliente.post("/renderizar", json={"word": carregar("pop_linac_manutencao.json"),
                                          "excel": {"abas": "isto não é uma lista"}})
    assert r.status_code == 200
    assert [a["formato"] for a in r.json()["arquivos"]] == ["word"]


def test_nada_para_gerar(cliente):
    assert cliente.post("/renderizar", json={}).status_code == 422


def test_logo_invalido(cliente):
    pedido = {"word": carregar("pop_linac_manutencao.json")}
    assert cliente.post("/renderizar", json={**pedido, "logo_b64": "@@@"}).status_code == 400
    texto = base64.b64encode(b"nao sou uma imagem").decode()
    assert cliente.post("/renderizar", json={**pedido, "logo_b64": texto}).status_code == 400
    assert cliente.post("/renderizar", json={**pedido, "logo_b64": imagem_b64("JPEG")}).status_code == 200


def test_token(monkeypatch):
    cliente = app_do_space(monkeypatch, token="segredo")
    pedido = {"word": carregar("pop_linac_manutencao.json")}
    assert cliente.post("/renderizar", json=pedido).status_code == 401
    assert cliente.post("/renderizar", json=pedido, headers={"Authorization": "Bearer errado"}).status_code == 401
    assert cliente.post("/renderizar", json=pedido, headers={"Authorization": "Bearer segredo"}).status_code == 200
    cabecalhos = {"Authorization": "Bearer hf_conta", "X-Pop-Token": "segredo"}
    assert cliente.post("/renderizar", json=pedido, headers=cabecalhos).status_code == 200
    assert cliente.post("/renderizar", json=pedido, headers={"X-Pop-Token": "segredo2"}).status_code == 401
    assert cliente.get("/saude").status_code == 200


def test_sem_segredo_o_motor_recusa(monkeypatch):
    cliente = app_do_space(monkeypatch, token="")
    pedido = {"word": carregar("pop_linac_manutencao.json")}
    assert cliente.post("/renderizar", json=pedido).status_code == 403
    assert cliente.post("/renderizar", json=pedido, headers={"X-Pop-Token": ""}).status_code == 403
    assert cliente.get("/saude").status_code == 200


def test_ilustracoes_do_site_entram_nos_slides_de_imagem(cliente):
    """As ilustrações geradas no site vão no slide pelo número, com a legenda de IA;
    a panorâmica é recortada para preencher a faixa (recorte editável no PowerPoint)."""
    from pptx import Presentation

    ppt = carregar("ppt_linac_manutencao.json")
    for i, s in enumerate(ppt["slides"], 1):
        s["numero"] = i
    imagem = [s for s in ppt["slides"] if s["tipo"].startswith("imagem_web_")]
    lateral = next(s for s in imagem if s["tipo"] != "imagem_web_panoramica")
    panoramica = next(s for s in imagem if s["tipo"] == "imagem_web_panoramica")
    r = cliente.post("/renderizar", json={"ppt": ppt, "imagens": {
        str(lateral["numero"]): imagem_b64("PNG"), str(panoramica["numero"]): imagem_b64("JPEG"),
        "999": imagem_b64("PNG"), "x": "não é imagem"}})
    assert r.status_code == 200, r.text
    deck = Presentation(io.BytesIO(base64.b64decode(r.json()["arquivos"][0]["base64"])))

    def figuras(numero):
        return [sh for sh in deck.slides[numero - 1].shapes if sh.shape_type == 13]

    def textos(numero):
        return " ".join(sh.text_frame.text for sh in deck.slides[numero - 1].shapes if sh.has_text_frame)

    assert len(figuras(lateral["numero"])) == 1
    assert "Imagem gerada por IA" in textos(lateral["numero"])
    pano = figuras(panoramica["numero"])
    assert len(pano) == 1 and "Imagem gerada por IA" in textos(panoramica["numero"])
    assert pano[0].crop_top > 0.1 and abs(pano[0].crop_top - pano[0].crop_bottom) < 1e-6
    outro = next(s for s in imagem if s["numero"] not in (lateral["numero"], panoramica["numero"]))
    assert "Imagem gerada por IA" not in textos(outro["numero"])


def test_renderiza_documento_de_apoio(cliente):
    r = cliente.post("/renderizar", json={"word": carregar("chk_carro_emergencia.json"), "cor": "#1F6F5C"})
    assert r.status_code == 200, r.text
    arquivo = r.json()["arquivos"][0]
    assert arquivo["formato"] == "word" and arquivo["nome"].startswith("CHK-ENF-001_")
    doc = Document(io.BytesIO(base64.b64decode(arquivo["base64"])))
    assert any("☐" in p.text for p in doc.paragraphs)
