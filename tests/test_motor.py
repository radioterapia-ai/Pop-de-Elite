"""Testes do motor (JSON → Word, Excel, PowerPoint) — rode com:  python -m pytest"""

import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from motor import gerar_excel, gerar_ppt, gerar_word, natureza_processo, validar_word
from motor.base import (DEFAULT_PRIMARY, nome_arquivo_seguro, nome_de_saida, normalizar_metadata,
                        parse_color_input, somar_anos)

EXEMPLOS = Path(__file__).resolve().parent.parent / "exemplos"


def carregar(nome):
    return json.loads((EXEMPLOS / nome).read_text(encoding="utf-8"))


def textos(container):
    """Todo o texto de um documento/cabeçalho/rodapé, incluindo tabelas aninhadas."""
    out = [p.text for p in container.paragraphs]

    def tabela(t):
        for row in t.rows:
            for cell in row.cells:
                out.extend(p.text for p in cell.paragraphs)
                for t2 in cell.tables:
                    tabela(t2)

    for t in container.tables:
        tabela(t)
    return "\n".join(out)


@pytest.fixture(autouse=True)
def sem_rede(monkeypatch):
    monkeypatch.setenv("POP_SEM_IMAGENS", "1")



@pytest.mark.parametrize("entrada, esperado", [
    ("#283264", "283264"),
    ("283264", "283264"),
    ("#abc", "AABBCC"),
    ("rgb(40, 50, 100)", "283264"),
    ("", DEFAULT_PRIMARY),
    (None, DEFAULT_PRIMARY),
])
def test_cores(entrada, esperado):
    assert parse_color_input(entrada) == esperado


@pytest.mark.parametrize("entrada, esperado", [
    ("POP-RTX-014_Manutenção preventiva", "POP-RTX-014_Manutencao_preventiva"),
    ("Coleta/transporte: amostras", "Coleta-transporte_amostras"),
    ("———", "documento"),
])
def test_nome_de_arquivo_ascii(entrada, esperado):
    assert nome_arquivo_seguro(entrada) == esperado


@pytest.mark.parametrize("codigo, descricao, ext, esperado", [
    ("PLAN-RTX-014", "Gestão da Qualidade.xlsx", "xlsx", "PLAN-RTX-014_Gestao_da_Qualidade.xlsx"),
    ("PPT-RTX-014", "PPT-RTX-014_Treinamento.pptx", "pptx", "PPT-RTX-014_Treinamento.pptx"),
    ("", "Treinamento", "pptx", "Treinamento.pptx"),
    (None, None, "xlsx", "Planilha.xlsx"),
])
def test_nome_de_saida(codigo, descricao, ext, esperado):
    assert nome_de_saida(codigo, descricao, ext, "Planilha" if ext == "xlsx" else "Apresentacao") == esperado


def test_validade_de_dois_anos_e_revisado_por():
    m = normalizar_metadata({"codigo": "POP-RTX-014", "data_elaboracao": "08/2026", "validade": "08/2030",
                             "validado_por": {"nome": "Fulana", "cargo": "Coordenação"}})
    assert m["validade"] == "08/2028"
    assert m["revisado_por"] == {"nome": "Fulana", "cargo": "Coordenação"}
    assert "validado_por" not in m
    assert m["tipo_documento"] == "POP" and m["classificacao"] == "interno"


def test_data_invalida_vira_mes_atual():
    m = normalizar_metadata({"data_elaboracao": "ontem"}, hoje=date(2026, 10, 1))
    assert m["data_elaboracao"] == "10/2026" and m["validade"] == "10/2028"
    assert somar_anos("29/02/2028", 2) == "28/02/2030"


def test_tipo_pelo_prefixo_do_codigo():
    assert normalizar_metadata({"codigo": "PROT-ONC-002"})["tipo_documento"] == "PROT"
    assert normalizar_metadata({"codigo": "XYZ-1"})["tipo_documento"] == "POP"



def test_word_linac(tmp_path):
    dados = carregar("pop_linac_manutencao.json")
    assert natureza_processo(dados["metadata"], dados["secoes"]) == "radiotecnico"
    caminho = Path(gerar_word(dados, pasta=tmp_path, cor="#283264"))
    assert caminho.name.startswith("POP-RTX-014_") and caminho.suffix == ".docx"
    assert caminho.name.isascii()

    doc = Document(caminho)
    sec = doc.sections[0]
    cabecalho, rodape, corpo = textos(sec.header), textos(sec.footer), textos(doc)
    assert "PROCEDIMENTO OPERACIONAL PADRÃO" in cabecalho.upper()
    for rotulo in ("Elaborado por", "Revisado por", "Aprovado por", "Vigência"):
        assert rotulo in rodape
    assert "Validado por" not in rodape
    assert "08/2028" in rodape
    assert "POP-RTX-014" in rodape and "Página" in rodape
    assert "pulseira" not in corpo.lower()
    assert "RISCOS RADIOLÓGICOS" in corpo.upper()


def test_word_em_ingles(tmp_path):
    dados = carregar("pop_linac_manutencao.json")
    dados["metadata"]["idioma"] = "en"
    doc = Document(gerar_word(dados, pasta=tmp_path))
    rodape = textos(doc.sections[0].footer)
    assert "Prepared by" in rodape and "Reviewed by" in rodape and "Page" in rodape
    assert "STANDARD OPERATING PROCEDURE" in textos(doc.sections[0].header).upper()


def test_validador_estrutural():
    ok, _, dados = validar_word(carregar("pop_linac_manutencao.json"))
    assert ok and dados["metadata"]["codigo"] == "POP-RTX-014"
    ok, msg, _ = validar_word({"metadata": {"titulo_processo": "X"}})
    assert not ok and "metadata.codigo vazio" in msg and "'secoes' ausente" in msg
    ok, msg, _ = validar_word("{isto não é json")
    assert not ok and "JSON inválido" in msg



def test_excel_linac(tmp_path):
    caminho = Path(gerar_excel(carregar("planilha_linac_manutencao.json"), pasta=tmp_path))
    assert caminho.suffix == ".xlsx" and caminho.name.isascii()
    wb = load_workbook(caminho)
    assert len(wb.sheetnames) >= 3
    assert len(set(wb.sheetnames)) == len(wb.sheetnames)


def test_ppt_layouts_novos(tmp_path):
    caminho = Path(gerar_ppt(carregar("ppt_layouts_novos.json"), pasta=tmp_path, estilo="cientifico"))
    assert caminho.suffix == ".pptx" and caminho.name.isascii()
    prs = Presentation(caminho)
    assert len(prs.slides) >= 6
    todo = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "!" in todo


def test_ppt_completo_30_slides(tmp_path):
    caminho = Path(gerar_ppt(carregar("ppt_linac_manutencao.json"), pasta=tmp_path))
    assert caminho.name == "PPT-RTX-014_Treinamento_Manutencao_Preventiva_Acelerador_RTX.pptx"
    prs = Presentation(caminho)
    assert len(prs.slides) == 30
    largura, altura = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, 1):
        for forma in slide.shapes:
            if forma.left is None or forma.width is None:
                continue
            assert forma.left >= -1 and forma.left + forma.width <= largura + 1, f"slide {i}: {forma.name} fora da largura"
            assert forma.top + forma.height <= altura + 1, f"slide {i}: {forma.name} fora da altura"


def test_ppt_fluxograma_nativo_legivel(tmp_path):
    dados = carregar("ppt_linac_manutencao.json")
    fluxo = next(s for s in dados["slides"] if s["tipo"] == "fluxograma")
    prs = Presentation(gerar_ppt(dict(dados, slides=[fluxo]), pasta=tmp_path))
    formas = [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text.strip()]

    def texto(sh):
        return " ".join(sh.text_frame.text.split())

    tarefas = [n["texto"] for n in fluxo["etapas"] if n["tipo"] == "tarefa"]
    caixas = [sh for sh in formas if texto(sh) in tarefas]
    assert sorted(texto(sh) for sh in caixas) == sorted(tarefas)
    tamanhos = {r.font.size for sh in caixas for p in sh.text_frame.paragraphs for r in p.runs}
    assert len(tamanhos) == 1 and tamanhos.pop().pt >= 13
    rotulos = [sh for sh in formas if texto(sh) in ("Manutenção concluída", "Equipamento liberado", "Equipamento bloqueado")]
    assert len(rotulos) == 3 and all(not sh.text_frame.word_wrap for sh in rotulos)


def test_word_e_ppt_misto_braquiterapia(tmp_path):
    dados = carregar("pop_braquiterapia_hdr.json")
    assert natureza_processo(dados["metadata"], dados["secoes"]) == "misto"
    doc = Document(gerar_word(dados, pasta=tmp_path))
    assert "RISCOS ASSISTENCIAIS E OPERACIONAIS" in textos(doc).upper()
    assert "cantSplit" not in doc.element.xml
    prs = Presentation(gerar_ppt(carregar("ppt_braquiterapia_hdr.json"), pasta=tmp_path))
    assert len(prs.slides) >= 30


def test_medida_de_texto_usa_as_metricas_da_calibri():
    from motor import diagramas
    assert diagramas.metrica_calibri(), "instale a Carlito (fonts-crosextra-carlito) ou a Calibri"
