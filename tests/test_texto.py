"""Testes do layout "texto institucional" dos documentos de apoio — rode com:  python -m pytest"""

import json
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from motor import gerar_word, validar_word
from motor.base import TIPOS_DOCUMENTO, TIPOS_TEXTO, normalizar_metadata

EXEMPLOS = Path(__file__).resolve().parent.parent / "exemplos"

APOIO = {
    "CONS": ("cons_tcle_endoscopia.json", "TERMO DE CONSENTIMENTO LIVRE E ESCLARECIDO", "paciente"),
    "TERM": ("term_normas_acompanhante.json", "TERMO DE CIÊNCIA", "paciente"),
    "CHK": ("chk_carro_emergencia.json", "CHECKLIST", "interno"),
    "FORM": ("form_notificacao_incidente.json", "FORMULÁRIO", "interno"),
    "INF": ("inf_preparo_sedacao.json", "INFORMATIVO", "paciente"),
    "NT": ("nt_ia_documentos.json", "NOTA TÉCNICA", "interno"),
    "COD": ("cod_conduta.json", "CÓDIGO DE CONDUTA", "interno"),
}


def carregar(nome):
    return json.loads((EXEMPLOS / nome).read_text(encoding="utf-8"))


def textos(container):
    """Todo o texto de um documento/cabeçalho/rodapé, incluindo tabelas aninhadas."""
    out = [p.text for p in container.paragraphs]

    def tabela(t):
        for row in t.rows:
            for cell in row.cells:
                out.extend(p.text for p in cell.paragraphs)
                for t2 in cell.tables:
                    tabela(t2)

    for t in container.tables:
        tabela(t)
    return "\n".join(out)


def minimo(**meta):
    """Documento de apoio mínimo, para testar uma regra de cada vez."""
    m = {"tipo_documento": "CHK", "codigo": "CHK-ENF-001", "titulo": "Conferência", "data_elaboracao": "10/2026"}
    m.update(meta)
    return {"metadata": m, "corpo": [{"tipo": "paragrafo", "texto": "Texto."}]}


def test_tipos_de_apoio_reconhecidos():
    assert set(TIPOS_TEXTO) == {"CONS", "TERM", "CHK", "FORM", "INF", "NT", "COD"}
    assert set(TIPOS_TEXTO) <= set(TIPOS_DOCUMENTO)
    assert "FTI" in TIPOS_DOCUMENTO and "FTI" not in TIPOS_TEXTO
    assert normalizar_metadata({"codigo": "CONS-CIR-001"})["tipo_documento"] == "CONS"
    assert normalizar_metadata({"codigo": "FORM-ENF-003"})["tipo_documento"] == "FORM"


def test_publico_e_classificacao_padrao():
    m = normalizar_metadata({"codigo": "CONS-CIR-001", "publico": "Paciente"})
    assert m["publico"] == "paciente" and m["classificacao"] == "publico"
    m = normalizar_metadata({"codigo": "CHK-ENF-001"})
    assert m["publico"] == "interno" and m["classificacao"] == "interno"
    m = normalizar_metadata({"codigo": "INF-ONC-001", "publico": "paciente", "classificacao": "interno"})
    assert m["classificacao"] == "interno"


def test_titulo_aceito_no_lugar_de_titulo_processo():
    m = normalizar_metadata({"codigo": "NT-QUA-001", "titulo": "Uso de IA"}, hoje=date(2026, 10, 4))
    assert m["titulo_processo"] == "Uso de IA" and m["validade"] == "10/2028"


@pytest.mark.parametrize("tipo", sorted(APOIO))
def test_exemplo_de_cada_tipo_de_apoio(tipo, tmp_path):
    arquivo, titulo, publico = APOIO[tipo]
    dados = carregar(arquivo)
    meta = dados["metadata"]
    assert meta["tipo_documento"] == tipo and meta["publico"] == publico

    caminho = Path(gerar_word(dados, pasta=tmp_path))
    assert caminho.name.startswith(meta["codigo"] + "_") and caminho.suffix == ".docx"
    assert caminho.name.isascii()

    doc = Document(caminho)
    sec = doc.sections[0]
    cabecalho, rodape, corpo = textos(sec.header), textos(sec.footer), textos(doc)
    assert titulo in cabecalho.upper()
    assert meta["titulo"].upper() in cabecalho.upper()
    assert "OBJETIVO / FINALIDADE" not in corpo
    assert meta["codigo"] in rodape and "Página" in rodape
    if publico == "interno":
        for rotulo in ("Elaborado por", "Revisado por", "Aprovado por", "Vigência"):
            assert rotulo in rodape
    else:
        assert "Elaborado por" not in rodape
        assert "Vigência" in rodape
    for assinatura in dados.get("assinaturas", []):
        assert assinatura["papel"] in corpo


def test_tcle_modelo_sem_dado_de_paciente_e_com_assinaturas(tmp_path):
    dados = carregar("cons_tcle_endoscopia.json")
    papeis = [a["papel"] for a in dados["assinaturas"]]
    assert any("Testemunha" in p and "opcional" in p for p in papeis)
    texto = json.dumps(dados, ensure_ascii=False)
    assert "Resolução CFM nº 2.454/2026" in texto
    doc = Document(gerar_word(dados, pasta=tmp_path))
    corpo = textos(doc)
    assert corpo.count("Assinatura") >= len(papeis)
    assert "Nome do(a) paciente" in corpo


def test_caixas_campos_e_item_critico(tmp_path):
    dados = minimo()
    dados["corpo"] = [
        {"tipo": "titulo", "texto": "Conferência diária"},
        {"tipo": "caixas", "itens": ["Lacre íntegro", {"texto": "Desfibrilador testado", "critico": True}]},
        {"tipo": "campos", "itens": [{"rotulo": "Responsável", "linhas": 1}, {"rotulo": "Observações", "linhas": 3}]},
        {"tipo": "lista", "itens": ["Item A", "Item B"]},
        {"tipo": "tabela", "colunas": ["Item", "Quantidade"], "linhas": [["Adrenalina", "10"]]},
        "parágrafo solto vira texto",
    ]
    doc = Document(gerar_word(dados, pasta=tmp_path))
    corpo = textos(doc)
    assert "☐" in corpo and "Lacre íntegro" in corpo
    critico = next(p for p in doc.paragraphs if "Desfibrilador testado" in p.text)
    assert any(r.bold for r in critico.runs if "Desfibrilador" in r.text)
    assert "Responsável:" in corpo and "Observações:" in corpo
    assert "Adrenalina" in corpo and "Item B" in corpo and "parágrafo solto vira texto" in corpo


def test_negrito_markdown_sem_asteriscos(tmp_path):
    dados = minimo()
    dados["corpo"] = [{"tipo": "paragrafo", "texto": "Leia com **atenção** antes de assinar."}]
    doc = Document(gerar_word(dados, pasta=tmp_path))
    p = next(p for p in doc.paragraphs if "atenção" in p.text)
    assert "**" not in p.text
    assert any(r.bold and r.text == "atenção" for r in p.runs)


@pytest.mark.parametrize("idioma, titulo, assinatura, pagina", [
    ("en", "INFORMED CONSENT FORM", "Signature", "Page"),
    ("es", "CONSENTIMIENTO INFORMADO", "Firma", "Página"),
])
def test_texto_em_ingles_e_espanhol(idioma, titulo, assinatura, pagina, tmp_path):
    dados = carregar("cons_tcle_endoscopia.json")
    dados["metadata"]["idioma"] = idioma
    doc = Document(gerar_word(dados, pasta=tmp_path))
    assert titulo in textos(doc.sections[0].header).upper()
    assert assinatura in textos(doc)
    assert pagina in textos(doc.sections[0].footer)


def test_validador_estrutural_aceita_texto_institucional():
    ok, _, dados = validar_word(carregar("cons_tcle_endoscopia.json"))
    assert ok and dados["metadata"]["codigo"] == "CONS-END-001"
    ok, msg, _ = validar_word({"metadata": {"titulo": "X", "codigo": "CHK-ENF-001"}, "corpo": []})
    assert not ok and "'corpo' vazio" in msg
